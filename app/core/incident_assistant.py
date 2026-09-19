"""The assistant on the Incident (Phase 6 chapter 3, v1.61.0).

Three parts, each with its own switch on the Incidents page. The Case Chat
is told each synced camera's start by the Incident clock, so it answers
with the time of day and reads every camera at the same moment. Propose
events reads each synced camera's Digest (or its transcript) and offers
Events that join the Chronology only when a person accepts them, on the
Speaker check's pattern: propose, never apply. The Incident memo is written
from the incident record, the cameras' Digests merged onto the Incident
clock by the app for the call and never stored, and on the Chronology's
Events, each sentence that rests on one carrying its number.

A Speaker's numbered label is one camera's alone (the maintainer,
2026-09-17): the record drops it, and the rules tell the assistant so.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import time
import uuid
from urllib.parse import urlparse

from django.db import models
from django.utils import timezone

from core import (
    assistant,
    audit,
    chronology,
    engine,
    incidents,
    prompts,
    settings_store,
)
from core.assistant import DONE, FAILED, QUEUED, RUNNING, PromptTemplate

log = logging.getLogger(__name__)

FEATURE_EVENTS = "incident_events"
FEATURE_MEMO = "incident_memo"
# A proposal within this many seconds of an Event or another proposal on the
# same camera is the same thing said twice, and is dropped.
NEAR_SECONDS = 5.0
# The most proposals kept from one camera in one run.
PROPOSALS_MOST = 30
# Phase 7 chapter 2: proposals kept a window (the first and the second look
# together), and how long a watch phrase heard again joins its first hit.
WINDOW_MOST = 12
WATCH_JOIN_SECONDS = 30.0
STAMP = re.compile(r"\[(\d{1,2}):(\d\d):(\d\d)\]")

# The run's states while it is going: nothing is accepted or dismissed then.
PROPOSING = (QUEUED, RUNNING)
# How much of the line a proposal rests on must be found in what the camera
# was given, after whitespace and case are folded.
RESTS_CHECK = 40

CLOCK_TIME_IN = re.compile(r"\[(\d{1,2}):(\d{2}):(\d{2})\]")
DIGEST_NUMBER = re.compile(r"^\s*\d+[.)]\s*")
EVENT_MARK = re.compile(r"\(Event (\d+)\)")


# The switches -------------------------------------------------------------------


def _assistant_on() -> bool:
    return bool(settings_store.get("assistant_available"))


def case_chat_on() -> bool:
    """Case chat knows the incidents: the placements block in every question."""
    return bool(
        incidents.on() and _assistant_on() and settings_store.get("incidents_case_chat")
    )


def proposals_on() -> bool:
    return bool(
        incidents.on() and _assistant_on() and settings_store.get("incidents_propose")
    )


def memo_on() -> bool:
    return bool(
        incidents.on() and _assistant_on() and settings_store.get("incidents_memo")
    )


# The rows -----------------------------------------------------------------------


class IncidentMemo(models.Model):
    """The one memo an Incident has, written across its synced cameras."""

    id = models.UUIDField(primary_key=True, default=uuid.uuid4, editable=False)
    incident = models.OneToOneField(
        incidents.Incident, on_delete=models.CASCADE, related_name="memo"
    )
    asked_by = models.ForeignKey(
        "core.User", on_delete=models.SET_NULL, null=True, blank=True
    )
    state = models.CharField(max_length=10, default=QUEUED)
    stage = models.CharField(max_length=60, blank=True, default="")
    reason_class = models.CharField(max_length=40, blank=True, default="")
    text = models.TextField(blank=True, default="")
    # "[hh:mm:ss]" as printed -> seconds on the Incident clock.
    citations = models.JSONField(default=dict, blank=True)
    # The Chronology's numbering when the memo was written: "4" -> Event id.
    event_numbers = models.JSONField(default=dict, blank=True)
    cut_short = models.BooleanField(default=False)
    model = models.CharField(max_length=120, blank=True, default="")
    template_version = models.IntegerField(default=1)
    ground_rules_version = models.IntegerField(default=1)
    # Which cameras it drew on, which were read from their transcript alone,
    # which were not read (the record was too long), and which were left out
    # (not synced, or no transcript), by name.
    cameras_used = models.JSONField(default=list, blank=True)
    cameras_transcript_only = models.JSONField(default=list, blank=True)
    cameras_not_read = models.JSONField(default=list, blank=True)
    cameras_left_out = models.JSONField(default=list, blank=True)
    events_count = models.IntegerField(default=0)
    events_signature = models.CharField(max_length=32, blank=True, default="")
    cameras_signature = models.CharField(max_length=32, blank=True, default="")
    record_lines = models.IntegerField(default=0)
    created = models.DateTimeField(auto_now_add=True)
    written_at = models.DateTimeField(null=True, blank=True)

    def __str__(self) -> str:
        return f"memo of {self.incident_id}"


# What the cameras give ----------------------------------------------------------


def synced_cameras(incident) -> list:
    """The cameras settled on the clock, in the order they start."""
    return sorted(
        (
            one
            for one in incident.cameras.select_related("recording")
            if one.is_synced()
        ),
        key=lambda one: (one.starts_at, one.added),
    )


def _clock(incident, at: float) -> str:
    """A time on the Incident as the assistant reads and writes it: hh:mm:ss
    of the day when the Incident has a clock, else hh:mm:ss into it."""
    if incident.has_clock():
        return incidents.hms(incident.clock_zero + at)
    return incidents.hms(at)


def _clock_words(incident, at: float) -> str:
    """The same, said: "21:55:44 by the cameras' clock" or "0:47 into the incident"."""
    if incident.has_clock():
        return f"{_clock(incident, at)} by the cameras' clock"
    return f"{incidents.elapsed(at)} into the incident"


def events_signature(incident) -> str:
    parts = sorted(
        f"{one.pk}:{one.at:.2f}:{one.until or ''}:{one.text}:{one.note}:{one.to_check}"
        for one in incident.events.filter(proposed=False)
    )
    parts.append(f"about:{incident.about}")
    return hashlib.sha256("\n".join(parts).encode("utf-8")).hexdigest()[:16]


def cameras_signature(incident) -> str:
    parts = sorted(f"{one.pk}:{one.starts_at:.2f}" for one in synced_cameras(incident))
    return hashlib.sha256("\n".join(parts).encode("utf-8")).hexdigest()[:16]


# The Case Chat is told the placements -------------------------------------------


def placements_block(case, read: list) -> str:
    """One block per Incident among the recordings a question reads: each
    synced camera's start and end by the Incident clock, numbered as the
    question numbers the recordings, and the rule of arithmetic under it."""
    if not case_chat_on():
        return ""
    numbers = {recording.pk: number for number, recording in enumerate(read, 1)}
    blocks = []
    for incident in case.incidents.all():
        lines = []
        for camera in incident.cameras.select_related("recording"):
            number = numbers.get(camera.recording_id)
            if number is None:
                continue
            name = camera.camera_id()
            if camera.is_synced():
                lines.append(
                    f"Recording {number} ({name}) starts at "
                    f"{_clock(incident, camera.starts_at)} and runs to "
                    f"{_clock(incident, camera.ends_at())}."
                )
            elif camera.is_placed():
                lines.append(
                    f"Recording {number} ({name}) is in the incident but not yet "
                    f"synced; it runs from about {_clock(incident, camera.starts_at)}."
                )
        if not any("starts at" in line for line in lines):
            continue
        count = incident.cameras.count()
        if incident.has_clock():
            head = (
                f"Incident {incident.name}, {count} camera{'' if count == 1 else 's'} "
                f"on one clock, the time of day on {incident.clock_date}."
            )
            rule = (
                "A time in a recording plus that recording's start is the time of "
                "day; the same time of day on another camera is that time of day "
                "less the other camera's start. Where a camera's own stamp and the "
                "incident's clock differ, the incident's clock is right: the office "
                "synced the cameras."
            )
        else:
            head = (
                f"Incident {incident.name}, {count} camera{'' if count == 1 else 's'} "
                "on one clock; no camera has a clock, so the times below count "
                "from the incident's first camera, as hh:mm:ss into the incident."
            )
            rule = (
                "A time in a recording plus that recording's start is the time "
                "into the incident; the same moment on another camera is that "
                "less the other camera's start."
            )
        blocks.append("\n".join([head, *lines, rule]))
    return "\n\n".join(blocks)


# The incident record ------------------------------------------------------------


def _on_the_clock(incident, camera, text: str) -> tuple[float | None, str]:
    """Every [hh:mm:ss] of the recording rewritten to the Incident clock, and
    the first of them as seconds on the Incident."""
    first: list[float] = []

    def swap(match) -> str:
        hours, minutes, seconds = (int(part) for part in match.groups())
        at = camera.starts_at + hours * 3600 + minutes * 60 + seconds
        if not first:
            first.append(at)
        return f"[{_clock(incident, at)}]"

    rewritten = CLOCK_TIME_IN.sub(swap, text)
    return (first[0] if first else None), rewritten


def record_of(incident) -> dict:
    """The incident record: every synced camera's Digest, or its transcript
    when it has no Digest, with each line's times rewritten to the Incident
    clock and the camera's name in front, merged in time order. Made for the
    call, never stored.

    Returns rows as (seconds, camera name, line), and the names by how each
    camera was read: `used` from its Digest, `transcript_only` from its
    words, `left_out` with no transcript.
    """
    rows: list[tuple[float, str, str]] = []
    used: list[str] = []
    transcript_only: list[str] = []
    left_out: list[str] = []
    for camera in synced_cameras(incident):
        name = camera.camera_id()
        transcript = getattr(camera.recording, "transcript", None)
        if transcript is None:
            left_out.append(name)
            continue
        digest = assistant.digest_text(transcript) if assistant.digests_on() else ""
        if digest:
            for raw in digest.splitlines():
                line = DIGEST_NUMBER.sub("", raw).strip()
                if not line:
                    continue
                at, line = _on_the_clock(incident, camera, line)
                rows.append((camera.starts_at if at is None else at, name, line))
            used.append(name)
            continue
        for line in prompts.lines_of(transcript):
            label = prompts.plain_speaker(line.speaker)
            # A numbered label is this camera's alone and is not carried.
            who = "" if not label or assistant._is_a_label(label) else f"{label}: "
            at = camera.starts_at + line.start
            rows.append((at, name, f"[{_clock(incident, at)}] {who}{line.text}"))
        transcript_only.append(name)
    rows.sort(key=lambda one: one[0])
    return {
        "rows": rows,
        "used": used,
        "transcript_only": transcript_only,
        "left_out": left_out,
    }


def _record_text(incident, record: dict, dropped: set[str]) -> str:
    """The record as the memo reads it; a dropped camera is one line."""
    lines = [
        f"{name}: {line}" for _, name, line in record["rows"] if name not in dropped
    ]
    for camera in synced_cameras(incident):
        name = camera.camera_id()
        if name in dropped:
            lines.append(
                f"{name}: ran from {_clock(incident, camera.starts_at)} to "
                f"{_clock(incident, camera.ends_at())}; its words were not read "
                "(too long without a record of its own)."
            )
    return "\n".join(lines)


def _cameras_line(incident) -> str:
    cameras = synced_cameras(incident)
    said = "; ".join(
        f"{one.camera_id()}, {_clock(incident, one.starts_at)} to "
        f"{_clock(incident, one.ends_at())}, "
        f"{incidents.PLACED_WORDS.get(one.placed, '').lower()}"
        for one in cameras
    )
    if incident.has_clock():
        head = f"Incident {incident.name}, the time of day on {incident.clock_date}."
    else:
        head = (
            f"Incident {incident.name}; no camera has a clock, so every time is "
            "hh:mm:ss into the incident."
        )
    return f"{head} Cameras: {said}."


def _chronology_lines(incident) -> tuple[list[str], dict]:
    """The Chronology as the office wrote it, numbered in time order, and the
    numbering as Event ids for the memo's marks."""
    names = {str(one.pk): one.camera_id() for one in incident.cameras.all()}
    lines = []
    numbers: dict[str, str] = {}
    for number, event in enumerate(
        incident.events.filter(proposed=False).select_related("added_by"), 1
    ):
        numbers[str(number)] = str(event.pk)
        seen = ", ".join(names[one] for one in event.cameras if one in names)
        camera = names.get(str(event.camera_id), "a camera no longer in the incident")
        if event.source == chronology.WORDS:
            origin = f"from the words on {camera}"
        elif event.source == chronology.CAMERA:
            origin = f"from what {camera} showed"
        elif event.source == chronology.ASSISTANT:
            origin = f"proposed from {camera} and accepted by a person"
        elif event.source == chronology.REPORT:
            origin = "from the report, which says: " + (event.rests_on or "")
        else:
            origin = "added by a person"
        when = _clock(incident, event.at) + (
            f" to {_clock(incident, event.until)}" if event.until else ""
        )
        lines.append(
            f"Event {number}, {when}, {event.text}"
            + (f"; seen on {seen}" if seen else "")
            + f"; {origin}."
            + (" To check: the office has not settled this." if event.to_check else "")
            + (f" The office's note: {event.note}" if event.note else "")
        )
    return lines, numbers


def _record_audit(
    incident,
    feature: str,
    *,
    actor,
    templates: str,
    model: str,
    usage: dict,
    started: float,
    outcome: str,
    reason: str = "",
    **more,
) -> None:
    """The one audit row a call writes, metadata only, the Incident as the object."""
    audit.write(
        audit.Category.LLM,
        "AI assistant call",
        actor=actor,
        outcome=audit.Outcome.SUCCESS if outcome == "ok" else audit.Outcome.FAILURE,
        reason_class=reason,
        affected_user=(
            incident.case.owner
            if actor is not None and incident.case.owner_id != actor.pk
            else None
        ),
        object_type="incident",
        object_id=incident.pk,
        object_label=incident.name,
        feature=feature,
        model=model or engine.model_name(),
        endpoint_host=urlparse(engine.address()).hostname or "",
        templates=templates,
        input_tokens=usage.get("input_tokens", 0),
        output_tokens=usage.get("output_tokens", 0),
        duration_seconds=round(time.monotonic() - started, 1),
        **more,
    )


# Proposed Events ----------------------------------------------------------------


def proposer_offered() -> bool:
    """Whether Propose events is on the page: the assistant's judgement, or
    the watch phrases alone when the assistant is off (Phase 7 chapter 2)."""
    return incidents.on() and (proposals_on() or bool(settings_store.watch_phrases()))


def proposals_possible(incident) -> bool:
    """Propose events is offered while some synced camera has a transcript,
    and the assistant or the watch phrases are there to read it."""
    return proposer_offered() and any(
        hasattr(one.recording, "transcript") for one in synced_cameras(incident)
    )


def ask_for_proposals(incident, *, by, look_for: str = "") -> bool:
    """Propose events pressed: one run at a time per Incident. The Look for
    is this run's alone, cleared when it ends."""
    if not proposals_possible(incident):
        return False
    if incident.proposals_state in (QUEUED, RUNNING):
        return True
    incident.proposals_state = QUEUED
    incident.proposals_reason = ""
    incident.proposals_by = by
    incident.proposals_look_for = " ".join(str(look_for or "").split())[:300]
    incident.save(
        update_fields=[
            "proposals_state",
            "proposals_reason",
            "proposals_by",
            "proposals_look_for",
        ]
    )
    from core import tasks

    tasks.propose_incident_events.defer(incident_id=str(incident.pk))
    return True


def _folded(text: str) -> str:
    return " ".join(text.lower().split())


def _chunks(system: str, head: str, body_lines: list[str], answer_cap: int) -> list:
    """The camera's lines in as few calls as fit the window, whole lines only."""
    room = (
        assistant.window()
        - prompts.tokens(system)
        - prompts.tokens(head)
        - assistant.cap(answer_cap)
        - 200
    )
    chunks: list[list[str]] = []
    current: list[str] = []
    used = 0
    for line in body_lines:
        cost = prompts.tokens(line)
        if cost > room:
            # A line longer than the whole room cannot be read; it is left out.
            continue
        if current and used + cost > room:
            chunks.append(current)
            current, used = [], 0
        current.append(line)
        used += cost
    if current:
        chunks.append(current)
    return chunks


def _keep_proposal(item: dict, camera, given: str, taken: list[float]) -> dict | None:
    """One answer item, checked: a time inside the camera, a line of text, the
    words it rests on found in what the camera was given, and nothing already
    within a few seconds on this camera."""
    if not isinstance(item, dict):
        return None
    seconds = prompts.clock_seconds(str(item.get("time", "")).strip())
    if seconds is None or seconds < 0 or seconds > camera.length():
        return None
    end = prompts.clock_seconds(str(item.get("end", "") or "").strip())
    if end is not None and (end <= seconds or end > camera.length() + 1):
        end = None
    text = " ".join(str(item.get("text", "")).split())[: chronology.TEXT_MOST]
    rests_on = " ".join(str(item.get("rests_on", "")).split())[: chronology.TEXT_MOST]
    why = " ".join(str(item.get("why", "") or "").split())[: chronology.WHY_MOST]
    if not text or not rests_on:
        return None
    if _folded(rests_on)[:RESTS_CHECK] not in given:
        return None
    at = camera.starts_at + seconds
    if any(abs(at - other) < NEAR_SECONDS for other in taken):
        return None
    return {
        "at": at,
        "until": None if end is None else camera.starts_at + end,
        "text": text,
        "rests_on": rests_on,
        "why": why,
    }


def _windows(lines: list[str], seconds: int) -> list[list[str]]:
    """A camera's lines in stretches of the window's length by the time each
    line carries; a line with no time stays with the stretch before it."""
    groups: list[list[str]] = []
    current: list[str] = []
    index: int | None = None
    for line in lines:
        found = STAMP.search(line)
        here = index
        if found is not None:
            at = (
                int(found.group(1)) * 3600
                + int(found.group(2)) * 60
                + int(found.group(3))
            )
            here = at // max(seconds, 1)
        if current and here != index:
            groups.append(current)
            current = []
        index = here
        current.append(line)
    if current:
        groups.append(current)
    return groups


# The watch phrases, the floor under the judgement (Phase 7 chapter 2) --------------


def _plain(text: str) -> str:
    return " ".join(text.replace("\u2019", "'").replace("\u2018", "'").lower().split())


def phrase_pattern(phrase: str) -> re.Pattern:
    """Whole words, any case, the plural allowed: "gun" finds "gun" and
    "guns" and never "begun"; a phrase's words may be split by any spacing."""
    words = [re.escape(word) for word in _plain(phrase).split()]
    return re.compile(r"(?<![a-z0-9])" + r"\s+".join(words) + r"(?:e?s)?(?![a-z0-9])")


def _watch_lines(camera) -> list[tuple[float, str, bool]]:
    """What the search reads on a camera: the transcript's lines at their
    times, and the record's picture lines where the camera has one."""
    transcript = camera.recording.transcript
    lines = [
        (segment.start, " ".join(segment.text.split()), False)
        for segment in transcript.segments.order_by("start")
        if segment.text.strip()
    ]
    if assistant.digests_on():
        for line in assistant.digest_text(transcript).splitlines():
            if "(seen)" not in line:
                continue
            found = STAMP.search(line)
            if found is None:
                continue
            at = (
                int(found.group(1)) * 3600
                + int(found.group(2)) * 60
                + int(found.group(3))
            )
            seen = " ".join(line.split("(seen)", 1)[1].split())
            if seen:
                lines.append((float(at), seen, True))
    lines.sort(key=lambda one: one[0])
    return lines


def watch_hits(camera, phrases: list[str]) -> list[dict]:
    """Every line on the camera that carries a watch phrase, as the fields of
    a proposal: a phrase heard again within the join window joins its first
    hit rather than making another."""
    if not phrases:
        return []
    patterns = [(phrase, phrase_pattern(phrase)) for phrase in phrases]
    hits: list[dict] = []
    latest: dict[str, dict] = {}
    for start, text, seen in _watch_lines(camera):
        plain = _plain(text)
        for phrase, pattern in patterns:
            if not pattern.search(plain):
                continue
            last = latest.get(phrase.lower())
            if last is not None and start - last["_start"] <= WATCH_JOIN_SECONDS:
                last["_count"] += 1
                last["_start"] = start
                continue
            hit = {
                "at": camera.starts_at + start,
                "until": None,
                "text": "",
                "why": f'The office watches for "{phrase}".',
                "rests_on": text[: chronology.TEXT_MOST],
                "_phrase": phrase,
                "_seen": seen,
                "_start": start,
                "_count": 1,
            }
            hits.append(hit)
            latest[phrase.lower()] = hit
    for hit in hits:
        said = f" (said {hit['_count']} times)" if hit["_count"] > 1 else ""
        how = ", seen" if hit["_seen"] else ""
        line = hit["rests_on"]
        room = chronology.TEXT_MOST - len(said) - len(how) - len(hit["_phrase"]) - 18
        hit["text"] = (
            f'Watch phrase "{hit["_phrase"]}"{how}: {line[: max(room, 20)]}{said}'
        )
        for key in ("_phrase", "_seen", "_start", "_count"):
            del hit[key]
    hits.sort(key=lambda one: one["at"])
    return hits


def _taken(camera, *groups) -> list[float]:
    return [one.at for group in groups for one in group if one.camera_id == camera.pk]


def _known(incident, standing: list, made: list) -> list[str]:
    return [
        f"{_clock(incident, one.at)} {one.text}"
        + (f" (the office's note: {one.note})" if one.note else "")
        for one in sorted([*standing, *made], key=lambda one: one.at)
    ]


def propose(incident_id, attempt: int = 1) -> None:
    """The run (Phase 7 chapter 2): first the watch phrases, searched by the
    app itself on every synced camera; then the engine's judgement, each
    camera read in windows with a second look at each; the answers checked;
    the pending proposals replaced with this run's."""
    incident = (
        incidents.Incident.objects.filter(pk=incident_id)
        .select_related("case", "proposals_by")
        .first()
    )
    if incident is None:
        return
    if incident.proposals_state not in (QUEUED, RUNNING):
        return
    started = time.monotonic()
    ground = PromptTemplate.named(PromptTemplate.GROUND_RULES)
    template = PromptTemplate.named(PromptTemplate.INCIDENT_EVENTS)
    templates_line = (
        f"ground-rules v{ground.version}; Proposed events v{template.version}"
    )
    incident.proposals_state = RUNNING
    incident.save(update_fields=["proposals_state"])
    look_for = incident.proposals_look_for
    usage = {"input_tokens": 0, "output_tokens": 0}
    model = ""
    calls = 0
    cut = 0
    windows = 0
    second_looks = 0
    watch_found = 0
    made: list = []
    cameras_read = 0
    phrases = settings_store.watch_phrases()

    def finish(state: str, reason: str = "") -> None:
        incident.proposals_state = state
        incident.proposals_reason = reason
        incident.proposals_at = timezone.now()
        incident.proposals_found = len(made)
        incident.proposals_cameras = cameras_read
        incident.proposals_cut = cut
        incident.proposals_watch = watch_found
        incident.proposals_look_for = ""
        incident.save()
        if calls:
            _record_audit(
                incident,
                FEATURE_EVENTS,
                actor=incident.proposals_by,
                templates=templates_line,
                model=model,
                usage=usage,
                started=started,
                outcome="ok" if state == DONE else reason,
                reason="" if state == DONE else reason,
                calls=calls,
                cameras=cameras_read,
                found=len(made),
                cut_short=cut,
                windows=windows,
                second_looks=second_looks,
                watch_hits=watch_found,
                look_for=bool(look_for),
            )

    try:
        cameras = [
            one
            for one in synced_cameras(incident)
            if hasattr(one.recording, "transcript")
        ]
        if not cameras:
            raise engine.Problem(engine.ERROR, "no synced camera has a transcript")
        if not proposals_on() and not phrases:
            raise engine.Problem(engine.ERROR, "proposals are off")
        # This run's proposals replace the pending ones; a dismissed one stays
        # dismissed so it is not offered again.
        incident.events.filter(proposed=True, dismissed=False).delete()
        standing = list(incident.events.filter(proposed=False))
        dismissed = list(incident.events.filter(proposed=True, dismissed=True))

        # The floor: the watch phrases, found by the app itself and saved as
        # they are found, so the page shows them before the engine answers.
        for camera in cameras:
            taken = _taken(camera, standing, dismissed, made)
            for fields in watch_hits(camera, phrases):
                if any(abs(fields["at"] - other) < NEAR_SECONDS for other in taken):
                    continue
                event = chronology.Event.objects.create(
                    incident=incident,
                    source=chronology.WATCH,
                    camera=camera,
                    cameras=chronology.running_cameras(incident, fields["at"]),
                    proposed=True,
                    added_by=incident.proposals_by,
                    **fields,
                )
                made.append(event)
                taken.append(fields["at"])
                watch_found += 1

        if not proposals_on():
            cameras_read = len(cameras)
            finish(DONE)
            _events_proposed_row(incident, made, cameras_read, watch_found, windows)
            return

        # The judgement: each camera in windows, a second look at each.
        problem = assistant._unreachable()
        if problem:
            raise problem
        system = prompts.system_message(
            ground.text, template.text, prompts.INCIDENT_EVENTS_FORMAT
        )
        context = settings_store.incident_events_context()
        answer_cap = settings_store.incident_events_answer_cap()
        window_seconds = settings_store.incident_events_window()
        second = settings_store.second_look()
        schema = prompts.incident_events_schema()
        for camera in cameras:
            transcript = camera.recording.transcript
            digest = assistant.digest_text(transcript) if assistant.digests_on() else ""
            if digest:
                body_lines = [line for line in digest.splitlines() if line.strip()]
                nature = (
                    "the record of this camera (a condensation of its words "
                    "and its picture)"
                )
            else:
                body_lines = prompts.render(prompts.lines_of(transcript)).splitlines()
                nature = "the transcript of this camera"
            if not body_lines:
                continue
            cameras_read += 1
            for window in _windows(body_lines, window_seconds):
                windows += 1
                kept_here = 0
                mine: list[str] = []
                for look in ("first", "second"):
                    if kept_here >= WINDOW_MOST:
                        break
                    if look == "second" and not second:
                        break
                    taken = _taken(camera, standing, dismissed, made)
                    head = prompts.incident_events_input(
                        camera.camera_id(),
                        _clock_words(incident, camera.starts_at),
                        _clock_words(incident, camera.ends_at()),
                        nature,
                        _known(incident, standing, made),
                        context=context,
                        look_for=look_for,
                    )
                    for chunk in _chunks(system, head, window, answer_cap):
                        given = _folded("\n".join(chunk))
                        user = head + "\n\n" + "\n".join(chunk)
                        if look == "second":
                            second_looks += 1
                            user += (
                                "\n\n"
                                + prompts.INCIDENT_EVENTS_SECOND_LOOK
                                + "\n\nYou proposed:\n"
                                + ("\n".join(mine) or "none")
                            )
                        answer = engine.complete(
                            assistant._messages(system, user),
                            max_completion_tokens=assistant.cap(answer_cap),
                            thinking=False,
                            timeout=assistant.time_limit(FEATURE_EVENTS),
                            schema=schema,
                            **assistant.SUGGESTION_SAMPLING,
                        )
                        calls += 1
                        if answer.get("finish_reason") == "length":
                            cut += 1
                        model = answer.get("model", "") or model
                        usage["input_tokens"] += answer.get("input_tokens", 0) or 0
                        usage["output_tokens"] += answer.get("output_tokens", 0) or 0
                        try:
                            try:
                                parsed = json.loads(answer["text"])
                            except ValueError:
                                parsed = json.loads(
                                    prompts.salvage_json(answer["text"])
                                )
                            raw = parsed.get("events", [])
                            if not isinstance(raw, list):
                                raise ValueError("not a list")
                        except (ValueError, AttributeError):
                            # One unreadable answer is a lost look, not a lost run.
                            log.warning(
                                "proposed events for %s: an answer was unreadable",
                                incident.pk,
                            )
                            continue
                        for item in raw:
                            if kept_here >= WINDOW_MOST:
                                break
                            fields = _keep_proposal(item, camera, given, taken)
                            if fields is None:
                                continue
                            event = chronology.Event.objects.create(
                                incident=incident,
                                source=chronology.ASSISTANT,
                                camera=camera,
                                cameras=chronology.running_cameras(
                                    incident, fields["at"]
                                ),
                                proposed=True,
                                added_by=incident.proposals_by,
                                **fields,
                            )
                            made.append(event)
                            taken.append(fields["at"])
                            kept_here += 1
                            mine.append(f"{item.get('time', '')} {fields['text']}")
        finish(DONE)
        _events_proposed_row(incident, made, cameras_read, watch_found, windows)
    except engine.Problem as problem:
        # What the search and the earlier cameras proposed is kept: every one
        # passed the checks.
        finish(FAILED, problem.reason)


def _events_proposed_row(incident, made, cameras_read, watch_found, windows) -> None:
    audit.write(
        incidents.CATEGORY,
        "events proposed",
        actor=incident.proposals_by,
        affected_user=incident.case.owner,
        object_type="incident",
        object_id=incident.pk,
        object_label=incident.name,
        found=len(made),
        cameras=cameras_read,
        watch_hits=watch_found,
        windows=windows,
    )


def proposals_json(incident) -> dict:
    """The run's state for the page: the button's line and whether to poll."""
    state = incident.proposals_state
    pending = incident.events.filter(proposed=True, dismissed=False).count()
    if state == QUEUED:
        words = "Waiting for the engine..."
    elif state == RUNNING:
        words = "Reading the cameras..."
    elif state == DONE:
        found = incident.proposals_found
        when = (
            timezone.localtime(incident.proposals_at).strftime("%H:%M")
            if incident.proposals_at
            else ""
        )
        words = (
            f"Proposed {found} event{'' if found == 1 else 's'} at {when}"
            if found
            else f"Nothing to propose ({when})"
        )
        if incident.proposals_watch:
            words += f", {incident.proposals_watch} from the watch phrases"
        if incident.proposals_cut:
            cut = incident.proposals_cut
            words += (
                f"; {cut} answer{'' if cut == 1 else 's'} cut short, "
                "raise Proposed events answer cap"
            )
    elif state == FAILED:
        words = assistant.what_to_say(incident.proposals_reason)
    else:
        words = ""
    return {
        "on": proposer_offered(),
        "possible": proposals_possible(incident),
        "state": state,
        "words": words,
        "pending": pending,
        "busy": state in (QUEUED, RUNNING),
    }


def accept(event, *, by, request=None):
    """A proposal joins the Chronology as an Event with source assistant."""
    if not event.proposed:
        return event
    event.proposed = False
    event.dismissed = False
    event.added_by = by
    event.added = timezone.now()
    if not event.cameras:
        event.cameras = chronology.running_cameras(event.incident, event.at)
    event.save()
    audit.write(
        incidents.CATEGORY,
        "event added",
        actor=by,
        affected_user=event.incident.case.owner,
        object_type="incident",
        object_id=event.incident_id,
        object_label=event.incident.name,
        request=request,
        source=chronology.ASSISTANT,
    )
    from core import cases

    cases.note_activity(event.incident.case, by=by)
    return event


def dismiss(event, *, by, request=None) -> None:
    """A proposal put away: kept as dismissed so a later run does not offer it again."""
    if not event.proposed:
        return
    event.dismissed = True
    event.save(update_fields=["dismissed"])
    audit.write(
        incidents.CATEGORY,
        "event dismissed",
        actor=by,
        affected_user=event.incident.case.owner,
        object_type="incident",
        object_id=event.incident_id,
        object_label=event.incident.name,
        request=request,
    )


def accept_all(incident, *, by, request=None) -> int:
    count = 0
    for event in list(incident.events.filter(proposed=True, dismissed=False)):
        accept(event, by=by, request=request)
        count += 1
    return count


# The Incident memo --------------------------------------------------------------


def memo_possible(incident) -> tuple[bool, str]:
    """Whether a memo can be written, and why not in words."""
    if not memo_on():
        return False, "The incident memo is off."
    cameras = synced_cameras(incident)
    if not cameras:
        return False, "No camera is synced yet; sync one and the memo can be written."
    if not any(hasattr(one.recording, "transcript") for one in cameras):
        return False, "No synced camera has a transcript yet."
    return True, ""


def memo_of(incident) -> IncidentMemo | None:
    return IncidentMemo.objects.filter(incident=incident).first()


def ask_for_memo(incident, *, by) -> IncidentMemo | None:
    """Write the memo, or Regenerate: the one row, reset and queued."""
    possible, _ = memo_possible(incident)
    if not possible:
        return None
    memo = memo_of(incident)
    if memo is not None and memo.state in (QUEUED, RUNNING):
        return memo
    if memo is None:
        memo = IncidentMemo(incident=incident)
    memo.asked_by = by
    memo.state = QUEUED
    memo.stage = ""
    memo.reason_class = ""
    memo.text = ""
    memo.citations = {}
    memo.event_numbers = {}
    memo.cut_short = False
    memo.written_at = None
    memo.save()
    from core import tasks

    tasks.write_incident_memo.defer(memo_id=str(memo.pk))
    return memo


def cancel_memo(incident) -> None:
    """A memo asked for and not wanted: the row goes; the task finds nothing."""
    memo = memo_of(incident)
    if memo is not None and memo.state in (QUEUED, RUNNING):
        memo.delete()


def memo_citations(incident, text: str) -> dict:
    """Every [hh:mm:ss] in the memo that falls inside the Incident's span, as
    seconds on the Incident; the rest stays plain text."""
    low, high = incidents.span_of(incident)
    if low is None:
        return {}
    found: dict = {}
    for match in CLOCK_TIME_IN.finditer(text):
        hours, minutes, seconds = (int(part) for part in match.groups())
        of_day = hours * 3600 + minutes * 60 + seconds
        if incident.has_clock():
            at = (of_day - incident.clock_zero) % incidents.DAY
        else:
            at = float(of_day)
        if low - 1 <= at <= high + 1:
            found[match.group(0)] = round(at, 2)
    return found


def write_memo(memo_id, attempt: int = 1) -> None:
    """The call: the record and the Chronology in, the memo out, one audit row."""
    memo = (
        IncidentMemo.objects.filter(pk=memo_id)
        .select_related("incident", "incident__case", "asked_by")
        .first()
    )
    if memo is None:
        if attempt == 1:
            from core import tasks

            tasks.write_incident_memo.configure(
                schedule_in={"seconds": assistant.QUEUE_GRACE_SECONDS}
            ).defer(memo_id=str(memo_id), attempt=attempt + 1)
        return
    if memo.state not in (QUEUED, RUNNING):
        return
    incident = memo.incident
    started = time.monotonic()
    ground = PromptTemplate.named(PromptTemplate.GROUND_RULES)
    template = PromptTemplate.named(PromptTemplate.INCIDENT_MEMO)
    templates_line = (
        f"ground-rules v{ground.version}; Incident memo v{template.version}"
    )
    memo.state = RUNNING
    memo.stage = "Reading the cameras"
    memo.template_version = template.version
    memo.ground_rules_version = ground.version
    memo.save(
        update_fields=["state", "stage", "template_version", "ground_rules_version"]
    )
    usage: dict = {"input_tokens": 0, "output_tokens": 0}
    try:
        problem = assistant._unreachable()
        if problem:
            raise problem
        possible, _ = memo_possible(incident)
        if not possible:
            raise engine.Problem(engine.ERROR, "the memo cannot be written")
        record = record_of(incident)
        if not record["rows"]:
            raise engine.Problem(engine.ERROR, "the cameras gave nothing to write from")
        event_lines, numbers = _chronology_lines(incident)
        system = prompts.system_message(
            ground.text,
            template.text,
            prompts.INCIDENT_MEMO_FORMAT
            + "\n\n"
            + prompts.NARRATIVE_RULES
            + "\n\n"
            + prompts.INCIDENT_RULES,
        )
        wanted = settings_store.incident_memo_answer_cap()
        dropped: set[str] = set()

        def user_text() -> str:
            return prompts.incident_memo_input(
                _cameras_line(incident),
                event_lines,
                _record_text(incident, record, dropped),
                about=incident.about,
            )

        # A transcript-only camera drops to a line when the whole does not
        # fit, the longest first; a Digest is never dropped.
        by_size = sorted(
            record["transcript_only"],
            key=lambda name: -sum(1 for _, who, _ in record["rows"] if who == name),
        )
        user = user_text()
        while not prompts.fits(
            system, user, answer_cap=assistant.cap(wanted), window=assistant.window()
        ):
            if not by_size:
                raise engine.Problem(engine.TOO_LONG, "the incident is too long")
            dropped.add(by_size.pop(0))
            user = user_text()
        memo.stage = "Writing the memo"
        memo.save(update_fields=["stage"])
        answer = engine.complete(
            assistant._messages(system, user),
            max_completion_tokens=assistant.cap(wanted),
            thinking=assistant.thinking(),
            timeout=assistant.time_limit(FEATURE_MEMO),
            **assistant.SAMPLING,
        )
        usage = answer
        if assistant.thought_it_away(answer):
            raise assistant.ThoughtItAway()
        if not IncidentMemo.objects.filter(pk=memo.pk).exists():
            # Cancelled while the engine wrote: nothing is kept.
            return
        memo.text = answer["text"].strip()
        memo.citations = memo_citations(incident, memo.text)
        memo.event_numbers = numbers
        memo.cut_short = answer.get("finish_reason") == "length"
        memo.model = answer.get("model", "")
        memo.cameras_used = record["used"]
        memo.cameras_transcript_only = [
            name for name in record["transcript_only"] if name not in dropped
        ]
        memo.cameras_not_read = sorted(dropped)
        memo.cameras_left_out = record["left_out"] + [
            one.camera_id()
            for one in incident.cameras.select_related("recording")
            if not one.is_synced()
        ]
        memo.events_count = len(event_lines)
        memo.events_signature = events_signature(incident)
        memo.cameras_signature = cameras_signature(incident)
        memo.record_lines = sum(1 for _, who, _ in record["rows"] if who not in dropped)
        memo.stage = ""
        memo.state = DONE
        memo.reason_class = ""
        memo.written_at = timezone.now()
        memo.save()
        _record_audit(
            incident,
            FEATURE_MEMO,
            actor=memo.asked_by,
            templates=templates_line,
            model=memo.model,
            usage=usage,
            started=started,
            outcome="ok",
            cameras=len(memo.cameras_used) + len(memo.cameras_transcript_only),
            events=memo.events_count,
            cut_short=memo.cut_short,
        )
    except engine.Problem as problem:
        if not IncidentMemo.objects.filter(pk=memo.pk).exists():
            return
        memo.state = FAILED
        memo.stage = ""
        memo.reason_class = (
            assistant.THOUGHT_AWAY
            if isinstance(problem, assistant.ThoughtItAway)
            else problem.reason
        )
        memo.save(update_fields=["state", "stage", "reason_class"])
        _record_audit(
            incident,
            FEATURE_MEMO,
            actor=memo.asked_by,
            templates=templates_line,
            model="",
            usage=usage,
            started=started,
            outcome=problem.reason,
            reason=problem.reason,
        )


def _count_words(count: int, word: str) -> str:
    return f"{count} {word}{'' if count == 1 else 's'}"


def written_words(memo: IncidentMemo) -> str:
    """ "Written 17 Sep 2026 from the transcripts and the vision of 7 cameras,
    on the chronology's 12 events. One camera (X) is not synced and was left out."""
    if memo.state != DONE or not memo.written_at:
        return ""
    when = timezone.localtime(memo.written_at).strftime("%d %b %Y")
    seen = len(memo.cameras_used)
    heard = len(memo.cameras_transcript_only)
    if seen and heard:
        source = (
            f"the transcripts and the vision of {_count_words(seen, 'camera')} and "
            f"the words alone of {heard}"
        )
    elif seen:
        source = f"the transcripts and the vision of {_count_words(seen, 'camera')}"
    else:
        source = f"the words of {_count_words(heard, 'camera')}"
    said = (
        f"Written {when} from {source}, on the chronology's "
        f"{_count_words(memo.events_count, 'event')}."
    )
    if memo.cameras_not_read:
        said += (
            f" {_count_words(len(memo.cameras_not_read), 'camera')} "
            f"({', '.join(memo.cameras_not_read)}) could not be read in full: the "
            "incident was too long without a record of it."
        )
    if memo.cameras_left_out:
        many = len(memo.cameras_left_out) > 1
        said += (
            f" {_count_words(len(memo.cameras_left_out), 'camera')} "
            f"({', '.join(memo.cameras_left_out)}) {'are' if many else 'is'} not "
            f"synced or {'have' if many else 'has'} no transcript and "
            f"{'were' if many else 'was'} left out."
        )
    return said


def stale_words(memo: IncidentMemo, incident) -> str:
    """What changed since the memo was written, or nothing."""
    if memo.state != DONE:
        return ""
    changes = []
    if memo.events_signature != events_signature(incident):
        now = incident.events.filter(proposed=False).count()
        if now > memo.events_count:
            changes.append(_count_words(now - memo.events_count, "event") + " added")
        elif now < memo.events_count:
            changes.append(_count_words(memo.events_count - now, "event") + " removed")
        else:
            changes.append("an event changed")
    if memo.cameras_signature != cameras_signature(incident):
        changes.append("a camera was synced or joined")
    if not changes:
        return ""
    return (
        "The chronology has changed since this memo was written: "
        + ", ".join(changes)
        + ". Regenerate to write it on them."
    )


def memo_json(incident) -> dict:
    """The Memo tab's state."""
    on = memo_on()
    possible, why_not = memo_possible(incident) if on else (False, "")
    memo = memo_of(incident) if on else None
    synced = synced_cameras(incident)
    not_synced = [
        one.camera_id()
        for one in incident.cameras.select_related("recording")
        if not one.is_synced()
    ]
    count = incident.events.filter(proposed=False).count()
    before = (
        f"{_count_words(len(synced), 'camera')} synced, "
        f"{_count_words(count, 'event')} on the chronology"
        + (
            f"; {', '.join(not_synced)} {'are' if len(not_synced) > 1 else 'is'} "
            "not synced and will be left out"
            if not_synced
            else ""
        )
        + "."
    )
    out = {
        "on": on,
        "possible": possible,
        "why_not": why_not,
        "before": before,
        "state": "",
    }
    if memo is None:
        return out
    out.update(
        {
            "state": memo.state,
            "stage": memo.stage,
            "reason_words": (
                assistant.what_to_say(memo.reason_class) if memo.state == FAILED else ""
            ),
            "text": memo.text if memo.state == DONE else "",
            "citations": memo.citations if memo.state == DONE else {},
            "event_numbers": memo.event_numbers if memo.state == DONE else {},
            "cut_short": memo.cut_short,
            "written_words": written_words(memo),
            "stale_words": stale_words(memo, incident),
            "notice": assistant.notice(memo.model, memo.written_at)
            if memo.state == DONE
            else "",
            "written_at": memo.written_at.isoformat() if memo.written_at else "",
            "events_count": memo.events_count,
            "record_lines": memo.record_lines,
            "cameras_used": memo.cameras_used,
            "cameras_transcript_only": memo.cameras_transcript_only,
            "busy": memo.state in (QUEUED, RUNNING),
        }
    )
    return out


def memo_line(incident) -> str:
    """The case page's word on the memo: written, stale, or none yet."""
    if not memo_on():
        return ""
    memo = memo_of(incident)
    if memo is None or memo.state != DONE or not memo.written_at:
        return "no memo yet"
    when = timezone.localtime(memo.written_at).strftime("%d %b")
    newer = incident.events.filter(proposed=False).count() - memo.events_count
    stale = stale_words(memo, incident)
    if newer > 0:
        return f"memo written {when}, {_count_words(newer, 'event')} newer"
    if stale:
        return f"memo written {when}, changed since"
    return f"memo written {when}"


# The Word export ----------------------------------------------------------------


def memo_word(memo: IncidentMemo, picture: bytes | None, exported_by: str) -> bytes:
    """The memo in the app's export shape, with the Chronology as its last pages."""
    from datetime import datetime

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    from core import exports

    incident = memo.incident
    document = exports._open_record_document()
    exports._office_head(document)
    heading = document.add_paragraph(f"Incident memo: {incident.name}")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(20)
    said = document.add_paragraph(
        "Written across the incident's cameras, on its chronology"
    )
    said.alignment = WD_ALIGN_PARAGRAPH.CENTER
    said.runs[0].font.size = Pt(13)
    document.add_paragraph()
    written = memo.written_at or memo.created
    read = memo.cameras_used + memo.cameras_transcript_only
    exports._facts(
        document,
        [
            ("Case", incident.case.name),
            ("Incident", incident.name),
            (
                "Clock",
                (
                    f"the time of day on {incident.clock_date}, by the cameras' clocks"
                    if incident.has_clock()
                    else "no camera clock; times count from the first camera"
                ),
            ),
            ("Cameras drawn on", ", ".join(read) or "none"),
            (
                "Left out",
                ", ".join(memo.cameras_not_read + memo.cameras_left_out) or "none",
            ),
            ("Events", str(memo.events_count)),
            ("Template", f"Incident memo v{memo.template_version}"),
            (
                "Written",
                f"{written:{exports.DAY_AND_TIME}} by "
                f"{memo.model or 'the AI assistant'}",
            ),
            ("Exported", f"{datetime.now():{exports.DAY_AND_TIME}} by {exported_by}"),
        ],
        Inches,
    )
    document.add_paragraph()
    notice = assistant.notice(memo.model, written)
    if notice:
        line = document.add_paragraph(notice)
        line.runs[0].italic = True
    stale = stale_words(memo, incident)
    if stale:
        line = document.add_paragraph(
            stale.replace(" Regenerate to write it on them.", "")
        )
        line.runs[0].italic = True
    document.add_paragraph()
    pieces = re.compile(r"(\[\d{1,2}:\d{2}:\d{2}\])")

    def paragraph(text: str, style=None) -> None:
        made = (
            document.add_paragraph(style=style) if style else document.add_paragraph()
        )
        for piece in pieces.split(text):
            if not piece:
                continue
            run = made.add_run(piece)
            if piece in memo.citations:
                run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    for raw in memo.text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        stripped = line.lstrip("#* ").rstrip()
        if stripped.endswith(":") and len(stripped) <= 60:
            document.add_heading(stripped[:-1], level=2)
            continue
        if stripped.startswith(("- ", "* ", "• ")):
            paragraph(stripped[2:], style="List Bullet")
            continue
        paragraph(line)
    if memo.cut_short:
        note = document.add_paragraph("The memo was cut short.")
        note.runs[0].italic = True
    document.add_page_break()
    chronology.pages(document, incident, picture, exported_by)
    import io

    holder = io.BytesIO()
    document.save(holder)
    return holder.getvalue()


def memo_export_name(incident) -> str:
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in incident.name)
    return f"Incident memo - {safe.strip() or 'incident'}.docx"


def record_memo_export(request, incident) -> None:
    audit.write(
        audit.Category.EXPORTS,
        "incident memo exported",
        actor=request.user,
        request=request,
        affected_user=incident.case.owner,
        object_type="incident",
        object_id=incident.pk,
        object_label=incident.name,
    )
