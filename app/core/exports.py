"""What a person takes out of a Workspace.

Three shapes: a Word document built for printing and for citing by page and
line, a plain-text file, and Captions. Nothing here is ever stored: an export
is built from the live Transcript when it is asked for and handed straight to
the browser.

Every export except Captions carries exactly one notice, the Translation
notice on a translated Transcript and the Transcription notice on every other,
and never both. Captions carry none, because an SRT file has nowhere to put a
notice that would not appear on screen.

Layouts, file names, and formats are code, not settings. The only
admin-editable text on any export is the two notices.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from datetime import datetime
from urllib.parse import quote

from django.contrib.auth.decorators import login_required
from django.http import HttpRequest, HttpResponse, HttpResponseNotFound
from django.shortcuts import redirect
from django.urls import reverse

from core import settings_store
from core.jobs import Transcript
from core.recordings import Recording

log = logging.getLogger("transcribe.exports")

CORRECTED_LEGEND_TEXT = "Lines marked (corrected) were corrected by staff."
CORRECTED_LEGEND_WORD = "Segments marked * were corrected by staff."

# What both parties of a call heard: the recorded announcement before it
# connects is on both channels, so both Sides transcribe it. The app prints it
# once, named for both, and says here that it did. It is never the app's to
# decide quietly what a transcript shows.
BOTH_SIDES_LEGEND = (
    "{count} passage{s} at the same moment on both sides, which a phone "
    "system's recorded announcement is, {is_or_are} printed once and named "
    "Side 1 and Side 2. Both copies are kept."
)


def both_sides_legend(transcript) -> str:
    count = transcript.shared_segments
    if not count:
        return ""
    one = count == 1
    return BOTH_SIDES_LEGEND.format(
        count=count, s="" if one else "s", is_or_are="is" if one else "are"
    )


# Dates read the same way everywhere an export prints one.
DAY = "%d %B %Y"
DAY_AND_TIME = "%d %B %Y %H:%M"


# The words --------------------------------------------------------------------


def clock(seconds: float) -> str:
    """[hh:mm:ss], hours always shown, matching the way Citations are written."""
    whole = max(0, int(seconds or 0))
    return f"{whole // 3600:02d}:{(whole % 3600) // 60:02d}:{whole % 60:02d}"


LANGUAGE_NAMES = {
    "en": "English",
    "es": "Spanish",
    "fr": "French",
    "de": "German",
    "it": "Italian",
    "pt": "Portuguese",
    "nl": "Dutch",
    "pl": "Polish",
    "ru": "Russian",
    "uk": "Ukrainian",
    "ar": "Arabic",
    "fa": "Persian",
    "he": "Hebrew",
    "tr": "Turkish",
    "hi": "Hindi",
    "ur": "Urdu",
    "pa": "Punjabi",
    "bn": "Bengali",
    "vi": "Vietnamese",
    "th": "Thai",
    "km": "Khmer",
    "lo": "Lao",
    "zh": "Chinese",
    "ja": "Japanese",
    "ko": "Korean",
    "tl": "Tagalog",
    "so": "Somali",
    "am": "Amharic",
    "sw": "Swahili",
    "ht": "Haitian Creole",
    "ro": "Romanian",
    "sr": "Serbian",
    "hr": "Croatian",
    "el": "Greek",
}


def language_name(code: str) -> str:
    """A language a person can read, falling back to whatever the engine said."""
    return LANGUAGE_NAMES.get((code or "").lower(), code or "an unknown language")


def title_of(recording: Recording) -> str:
    """The Recording title, falling back to the file name without its extension."""
    if recording.title:
        return recording.title
    name = recording.original_filename or "recording"
    return name.rsplit(".", 1)[0] if "." in name else name


def languages_heard(transcript: Transcript) -> str:
    """The languages of a mixed Recording, read out: "Spanish and English"."""
    counts = (transcript.detection or {}).get("combined") or {}
    names = [
        language_name(code) for code in sorted(counts, key=counts.get, reverse=True)
    ]
    if not names:
        return "more than one language"
    if len(names) == 1:
        return names[0]
    return ", ".join(names[:-1]) + " and " + names[-1]


def kind_line(transcript: Transcript) -> str:
    """What this document is, in one line, on the cover and at the head."""
    if transcript.task_run == "translate":
        if transcript.language_mixed:
            return f"Translated to English ({languages_heard(transcript)} detected)"
        return f"Translated to English from {language_name(transcript.language)}"
    if transcript.language and transcript.language.lower() != "en":
        return f"Transcript in {language_name(transcript.language)}"
    return "Transcript"


def notice_for(transcript: Transcript) -> str:
    """The one notice, with its placeholders filled from the Provenance."""
    model = model_of(transcript)
    if transcript.task_run == "translate":
        spoken = (
            languages_heard(transcript)
            if transcript.language_mixed
            else language_name(transcript.language)
        )
        return _fill(
            settings_store.get("translation_notice"), language=spoken, model=model
        )
    return _fill(settings_store.get("transcription_notice"), model=model)


def _fill(notice: str, **filling) -> str:
    """Put the Provenance into an Admin's wording, whatever they wrote.

    An Admin writing a notice can leave a placeholder out or put one in that
    the app does not fill, and neither should raise on an export: a notice
    with an unknown placeholder prints as it was written.
    """
    for name, value in filling.items():
        notice = notice.replace("{" + name + "}", str(value))
    return notice


def model_of(transcript: Transcript) -> str:
    for run in (transcript.provenance or {}).values():
        model = (run.get("settings_used") or {}).get("model")
        if model:
            return model
    return ""


def length_of(recording: Recording) -> str:
    """A length a person reads, not a number of seconds."""
    seconds = int(recording.duration_seconds or 0)
    hours, rest = divmod(seconds, 3600)
    minutes = rest // 60
    if hours and minutes:
        return f"{hours} h {minutes} min"
    if hours:
        return f"{hours} h"
    if minutes:
        return f"{minutes} min"
    return f"{seconds} sec"


class Appearance:
    """One Speaker, as the Appearances table and the Speakers line read them."""

    def __init__(
        self,
        name: str,
        label: str,
        side: str,
        count: int,
        talking: float,
        first: float,
        role: str = "",
    ):
        self.name = name
        self.label = label
        self.side = side
        self.count = count
        self.talking = talking
        self.first = first
        # Carried from Phase 1 and empty until a Recording is in a Case and
        # the Person holds a Role. The column below is drawn only when
        # something fills it, so Phase 2 fills a column rather than adding one.
        self.role = role

    @property
    def speaking_time(self) -> str:
        whole = int(self.talking)
        return f"{whole // 60}:{whole % 60:02d}"


def appearances(segments) -> list[Appearance]:
    """Who spoke, in the order they were first heard.

    Speaking time is the sum of the Segments a Speaker holds, and first heard
    is the start of the earliest: both are read off the Segments themselves,
    which is the only record of who spoke when after a rename or a merge.
    """
    found: dict[str, Appearance] = {}
    for segment in segments:
        if not segment.speaker:
            continue
        one = found.get(segment.speaker)
        if one is None:
            found[segment.speaker] = Appearance(
                name=segment.speaker,
                label=segment.speaker_label,
                side=segment.side.name if segment.side else "",
                count=1,
                talking=max(0.0, segment.end - segment.start),
                first=segment.start,
            )
            continue
        one.count += 1
        one.talking += max(0.0, segment.end - segment.start)
        if not one.label:
            one.label = segment.speaker_label
    return sorted(found.values(), key=lambda one: one.first)


# Plain text -------------------------------------------------------------------


def plain_text(recording: Recording) -> str:
    """The four-line head, a blank line, then one line per Segment.

    This is also the shape the AI assistant reads and the shape a Clip's
    excerpt uses, so it stays one thing rather than three that drift apart.
    """
    transcript = recording.transcript
    segments = list(
        transcript.segments.filter(same_as_other_side=False).select_related("side")
    )
    who = appearances(segments)

    head = [
        title_of(recording),
        (
            f"{kind_line(transcript)}. {recording.original_filename}, "
            f"{length_of(recording)}, uploaded {recording.created:{DAY}} by "
            f"{recording.user.username}. Processed {transcript.created:{DAY}} "
            f"with Whisper {model_of(transcript)}."
        ),
    ]
    if who:
        head.append(
            "Speakers: "
            + "; ".join(
                f"{one.name} ({one.label})" if one.label else one.name for one in who
            )
        )

    notice = notice_for(transcript)
    if any(segment.corrected for segment in segments):
        notice += " " + CORRECTED_LEGEND_TEXT
    shared = both_sides_legend(transcript)
    if shared:
        notice += " " + shared
    head.append(notice)
    head.append("")

    lines = []
    for segment in segments:
        name = segment.speaker
        if name and segment.corrected:
            name = f"{name} (corrected)"
        elif not name and segment.corrected:
            name = "corrected"
        start = f"[{clock(segment.start)}]"
        lines.append(
            f"{start} {name}: {segment.text}" if name else f"{start} {segment.text}"
        )

    # Windows line endings: these are read on office workstations.
    return "\r\n".join(head + lines) + "\r\n"


# Captions ---------------------------------------------------------------------


def srt(recording: Recording, offset: float = 0.0) -> str:
    """One cue per Segment, `Speaker: text`, and no notice."""
    cues = []
    wanted = recording.transcript.segments.filter(same_as_other_side=False)
    for number, segment in enumerate(wanted, start=1):
        text = f"{segment.speaker}: {segment.text}" if segment.speaker else segment.text
        cues.append(
            f"{number}\r\n"
            f"{srt_time(segment.start - offset)} --> "
            f"{srt_time(segment.end - offset)}\r\n"
            f"{text}\r\n"
        )
    return "\r\n".join(cues)


def srt_time(seconds: float) -> str:
    seconds = max(0.0, seconds)
    whole = int(seconds)
    milliseconds = int(round((seconds - whole) * 1000))
    if milliseconds == 1000:
        whole, milliseconds = whole + 1, 0
    return (
        f"{whole // 3600:02d}:{(whole % 3600) // 60:02d}:"
        f"{whole % 60:02d},{milliseconds:03d}"
    )


# Word -------------------------------------------------------------------------


def word(recording: Recording, exported_by: str) -> bytes:
    """Layout "Record": a cover page, the line-numbered talk, then the record.

    US Letter, one-inch margins, the talk in a fixed-width face so that a
    printed page cites the same way whoever prints it.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    transcript = recording.transcript
    segments = list(
        transcript.segments.filter(same_as_other_side=False).select_related("side")
    )
    who = appearances(segments)
    # Inside a Case the Role column fills from the People; blank otherwise.
    if recording.case_id:
        from core import people

        for one in who:
            one.role = people.role_of(recording, one.name)
    corrections = sum(1 for segment in segments if segment.corrected)
    title = title_of(recording)
    kind = kind_line(transcript)

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)

    # The cover -----------------------------------------------------------------
    heading = document.add_paragraph(title)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(20)

    said = document.add_paragraph(kind)
    said.alignment = WD_ALIGN_PARAGRAPH.CENTER
    said.runs[0].font.size = Pt(13)
    document.add_paragraph()

    # The Case, when there is one. The upload date stays the only date on the
    # cover; a Case has no date of its own that belongs here.
    in_a_case = (
        [("Case", recording.case.name)] if getattr(recording, "case", None) else []
    )

    _facts(
        document,
        [
            ("Recording", recording.original_filename),
            *in_a_case,
            ("Length", length_of(recording)),
            ("Uploaded", f"{recording.created:{DAY}} by {recording.user.username}"),
            (
                "Speakers",
                "; ".join(one.name for one in who) if who else "not separated",
            ),
            ("Language", language_name(transcript.language)),
            (
                "Processed",
                f"{transcript.created:{DAY}} with Whisper {model_of(transcript)}",
            ),
            ("SHA-256", recording.sha256),
        ],
        Inches,
    )

    document.add_paragraph()
    notice = document.add_paragraph(notice_for(transcript))
    notice.runs[0].italic = True
    if corrections:
        legend = document.add_paragraph(CORRECTED_LEGEND_WORD)
        legend.runs[0].italic = True
    shared = both_sides_legend(transcript)
    if shared:
        both = document.add_paragraph(shared)
        both.runs[0].italic = True

    if who:
        document.add_paragraph()
        appearing = document.add_paragraph("Appearances")
        appearing.runs[0].bold = True
        _appearances_table(document, who, Pt)

    # The talk ------------------------------------------------------------------
    talk = document.add_section(WD_SECTION.NEW_PAGE)
    _number_the_lines(talk)
    _running_head(talk, title, kind, recording.sha256, Pt, RGBColor)

    for segment in segments:
        mark = "*" if segment.corrected else " "
        name = f"{segment.speaker.upper()}: " if segment.speaker else ""
        line = document.add_paragraph()
        line.paragraph_format.space_after = Pt(6)
        run = line.add_run(f"[{clock(segment.start)}]{mark} {name}{segment.text}")
        run.font.name = "Consolas"
        run.font.size = Pt(10)

    if not segments:
        document.add_paragraph("This transcript has no segments.")

    # The processing record -----------------------------------------------------
    record = document.add_section(WD_SECTION.NEW_PAGE)
    _running_head(record, title, kind, recording.sha256, Pt, RGBColor)

    label = document.add_paragraph("Processing record")
    label.runs[0].bold = True
    label.runs[0].font.size = Pt(14)
    _facts(
        document,
        _provenance(recording, transcript, segments, corrections, exported_by),
        Inches,
    )

    holder = io.BytesIO()
    document.save(holder)
    return holder.getvalue()


def _facts(document, rows, Inches) -> None:
    """A two-column table of label and value: the cover facts and the record."""
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[0].width = Inches(1.8)
        cells[1].width = Inches(4.7)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True


def _appearances_table(document, who, Pt) -> None:
    """Name, Role, label, Side, segments, speaking time, first heard.

    The Role column is carried and hidden while empty, which is every Phase 1
    export: nothing fills a Role until a Recording is in a Case. It is written
    here rather than added later so that Phase 2 fills a column instead of
    changing the shape of the table.
    """
    roles = any(one.role for one in who)
    columns = ["Name"]
    if roles:
        columns.append("Role")
    columns += ["Label", "Side", "Segments", "Speaking time", "First heard"]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, heading in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)

    for one in who:
        cells = table.add_row().cells
        values = [one.name]
        if roles:
            values.append(one.role)
        values += [
            one.label,
            one.side,
            str(one.count),
            one.speaking_time,
            clock(one.first),
        ]
        for index, value in enumerate(values):
            cells[index].text = value
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(9)


def _provenance(recording, transcript, segments, corrections, exported_by):
    """How this Transcript came to be, as the rows the last pages print.

    The raw ffprobe output stays in the Details panel and is not printed.
    """
    runs = list((transcript.provenance or {}).values())
    first = runs[0] if runs else {}
    used = first.get("settings_used") or {}
    service = first.get("service") or {}
    vad = used.get("vad") or {}
    sides = list(recording.sides.all())

    rows = [
        ("Original name", recording.original_filename),
        ("Size", f"{recording.size_bytes / 1024 / 1024:.1f} MB"),
        ("SHA-256", recording.sha256),
        (
            "Uploaded by",
            f"{recording.user.username} on {recording.created:{DAY_AND_TIME}}",
        ),
        ("Sides", ", ".join(str(one) for one in sides) if sides else "1"),
        (
            "Two-channel call",
            "yes, one side per party" if recording.is_two_channel_call else "no",
        ),
        (
            "Preprocessing",
            f"{recording.preprocessing.title()}, loudness normalised "
            "(linear, -16 LUFS)",
        ),
        ("Model", used.get("model", "")),
        ("Model revision", used.get("model_revision", "")),
        (
            "Task",
            f"{used.get('task_run', '')} ({used.get('task_reason', '')})"
            if used.get("task_run")
            else "",
        ),
        (
            "Language",
            language_name(transcript.language)
            + (
                f", detected with {transcript.language_probability:.0%} confidence"
                if transcript.language_probability
                else ""
            ),
        ),
        (
            "Word timing",
            "yes"
            if transcript.word_timestamps
            else f"no ({transcript.word_timestamps_reason or 'not available'})",
        ),
        ("Diarization", "yes" if used.get("diarize") else "no"),
        ("Compute type", used.get("compute_type", "")),
        ("Batch size", str(used.get("batch_size", ""))),
        (
            "Voice activity",
            f"{vad.get('method', '')} onset {vad.get('onset', '')}, offset "
            f"{vad.get('offset', '')}, {vad.get('chunk_seconds', '')} s chunks"
            if vad
            else "",
        ),
        (
            "Vocabulary",
            f"{used.get('vocabulary_terms_used', 0)} of "
            f"{used.get('vocabulary_terms_given', 0)} terms used"
            if used.get("vocabulary_terms_given")
            else "none given",
        ),
        ("Service version", service.get("version", "")),
        (
            "Service pins",
            ", ".join(
                f"{name} {value}"
                for name, value in sorted((service.get("pins") or {}).items())
            ),
        ),
        ("Processed", f"{transcript.created:{DAY_AND_TIME}}"),
        ("Segments", str(len(segments))),
        ("Corrections", str(corrections)),
        ("Exported", f"{datetime.now():{DAY_AND_TIME}} by {exported_by}"),
    ]
    return [(label, value) for label, value in rows if str(value).strip()]


def _number_the_lines(section) -> None:
    """Line numbers down the margin, running on through the talk.

    They run through rather than restarting on each page, so a citation to a
    line means the same thing whatever a printer does with the page breaks.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = OxmlElement("w:lnNumType")
    numbering.set(qn("w:countBy"), "1")
    numbering.set(qn("w:restart"), "continuous")
    numbering.set(qn("w:distance"), "360")
    section._sectPr.append(numbering)


def _running_head(section, title, kind, sha256, Pt, RGBColor) -> None:
    """The title and kind above; the fingerprint, page, and date below.

    The exporter's name is in the Processing record only, never on every page.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    above = section.header.paragraphs[0]
    above.text = f"{title} - {kind}"
    above.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in above.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    below = section.footer.paragraphs[0]
    below.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_of_pages(below, sha256, Pt, RGBColor)


def _page_of_pages(paragraph, sha256, Pt, RGBColor) -> None:
    """ "<fingerprint>  Page X of Y  Gideon Transcribe, exported <date>".

    X and Y are Word fields rather than numbers, because only Word knows how
    many pages a transcript came to once it laid the text out.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def quiet(run):
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        return run

    def field(name):
        holder = quiet(paragraph.add_run())
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = f" {name} "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        holder._r.append(begin)
        holder._r.append(instruction)
        holder._r.append(end)

    quiet(paragraph.add_run(f"{sha256[:12]}    Page "))
    field("PAGE")
    quiet(paragraph.add_run(" of "))
    field("NUMPAGES")
    quiet(paragraph.add_run(f"    Gideon Transcribe, exported {datetime.now():{DAY}}"))


# File names -------------------------------------------------------------------

# Windows will not take these as a file name, whatever follows the dot.
RESERVED = (
    {"CON", "PRN", "AUX", "NUL"}
    | {f"COM{number}" for number in range(1, 10)}
    | {f"LPT{number}" for number in range(1, 10)}
)

# Letters, digits, spaces, and these survive; everything else becomes a hyphen.
NOT_ALLOWED = re.compile(r"[^A-Za-z0-9 \-_.,()']")


def safe_name(title: str) -> str:
    """A file name a Windows machine will take, from a title a person typed."""
    name = NOT_ALLOWED.sub("-", title or "")
    name = re.sub(r"\s+", " ", name).strip()[:80].strip().rstrip(".")
    if not name:
        return "recording"
    if name.split(".")[0].strip().upper() in RESERVED:
        name = f"{name} (file)"
    return name


def export_name(recording: Recording, ending: str) -> str:
    """`<title> - transcript.docx` and its neighbours."""
    return f"{safe_name(title_of(recording))} - {ending}"


def zip_name(kind: str = "Transcripts") -> str:
    return f"{kind} {datetime.now():%Y-%m-%d %H%M}.zip"


def without_clashes(taken: set[str], name: str) -> str:
    """Two Recordings can carry the same title; two zip entries cannot."""
    if name not in taken:
        taken.add(name)
        return name
    stem, dot, ending = name.rpartition(".")
    number = 2
    while True:
        tried = f"{stem} ({number}){dot}{ending}" if dot else f"{name} ({number})"
        if tried not in taken:
            taken.add(tried)
            return tried
        number += 1


# Many at once -----------------------------------------------------------------


def transcripts_zip(recordings) -> tuple[bytes, list]:
    """One plain-text file per Done Recording, and which ones went in.

    Recordings not yet Done are left out. The caller writes the audit rows: a
    zip writes one row per file inside it.
    """
    holder = io.BytesIO()
    taken: set[str] = set()
    included = []
    with zipfile.ZipFile(holder, "w", zipfile.ZIP_DEFLATED) as bundle:
        for recording in recordings:
            if not hasattr(recording, "transcript"):
                continue
            name = without_clashes(taken, export_name(recording, "transcript.txt"))
            bundle.writestr(name, plain_text(recording).encode("utf-8"))
            included.append(recording)
    return holder.getvalue(), included


def everything_zip(recordings, exported_by: str) -> tuple[bytes, list, list]:
    """For every Done Recording, its Word document and its plain-text file.

    In Phase 1 the Word document is the Word Transcript; the combined document
    gains the Summaries and Chats when those are built. Every Ready Clip goes
    in beside its Recording's two files, with its excerpt and captions. The
    shape inside is flat, because a Workspace is a holding area rather than a
    filing system and folders would only add a level to click through.
    """
    from core import clip_work

    holder = io.BytesIO()
    taken: set[str] = set()
    made = []
    clips = []
    with zipfile.ZipFile(holder, "w", zipfile.ZIP_DEFLATED) as bundle:
        for recording in recordings:
            if not hasattr(recording, "transcript"):
                continue
            document = without_clashes(taken, export_name(recording, "transcript.docx"))
            bundle.writestr(document, word(recording, exported_by))
            made.append((recording, "transcript word"))

            text = without_clashes(taken, export_name(recording, "transcript.txt"))
            bundle.writestr(text, plain_text(recording).encode("utf-8"))
            made.append((recording, "transcript text"))

            # Every Ready Clip of this Recording, beside its two files.
            for clip in _clips_of(recording):
                clip_work.add_to_zip(bundle, clip, taken)
                clips.append(clip)
    return holder.getvalue(), made, clips


def _clips_of(recording):
    """The Ready Clips of one Recording, or none while Clips are switched off."""
    from core import settings_store

    if not settings_store.get("clips_available"):
        return []
    return [
        one
        for one in recording.clips.all()
        if one.state == "ready" and one.path.exists()
    ]


# Handing them over ------------------------------------------------------------


def _may_open(request, recording_id):
    """The Recording, if this person may take it out.

    An Admin may export another person's material, and the row says whose.
    """
    recording = Recording.objects.filter(pk=recording_id).select_related("user").first()
    if recording is None:
        return None

    # Nothing in a Case that is in the Recycle bin is exported by anybody.
    from core import cases

    if not cases.reachable(recording):
        return None
    if recording.user_id == request.user.pk or request.user.is_admin:
        return recording
    return None


def record_export(request, recording, kind) -> None:
    """One "export made" row per file, naming the kind and never any text."""
    from core import audit

    audit.write(
        audit.Category.EXPORTS,
        "export made",
        actor=request.user,
        request=request,
        affected_user=(
            recording.user if recording.user_id != request.user.pk else None
        ),
        object_type="recording",
        object_id=recording.pk,
        object_label=recording.original_filename,
        kind=kind,
    )


def _record_clip(request, clip) -> None:
    """One "Clip downloaded" row per Clip inside any zip."""
    from core import audit, clip_work

    audit.write(
        audit.Category.CLIPS,
        "Clip downloaded",
        actor=request.user,
        request=request,
        object_type="clip",
        object_id=clip.pk,
        object_label=f"{clip.start:.1f}-{clip.end:.1f}",
        burn_captions=clip.burn_captions,
        include_excerpt=clip.include_excerpt,
    )
    clip_work.mark_downloaded(clip, clip.recording.user_id == request.user.pk)


def hand_over(body, filename: str, content_type: str) -> HttpResponse:
    """A download, named so that the browser keeps the name the app chose."""
    answer = HttpResponse(body, content_type=content_type)
    answer["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return answer


WORD_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@login_required
def export(request: HttpRequest, recording_id, shape: str) -> HttpResponse:
    """The three exports the viewer offers: Word, text, and Captions."""
    recording = _may_open(request, recording_id)
    if recording is None:
        return HttpResponseNotFound("There is no such recording.")
    if not hasattr(recording, "transcript"):
        return redirect(reverse("viewer", args=[recording.pk]))

    from core import cases

    # Exporting is use of the Case the Recording is in.
    cases.used(recording, by=request.user)

    if shape == "word":
        body = word(recording, request.user.username)
        record_export(request, recording, "transcript word")
        return hand_over(body, export_name(recording, "transcript.docx"), WORD_TYPE)

    if shape == "text":
        body = plain_text(recording).encode("utf-8")
        record_export(request, recording, "transcript text")
        return hand_over(
            body,
            export_name(recording, "transcript.txt"),
            "text/plain; charset=utf-8",
        )

    body = srt(recording).encode("utf-8")
    record_export(request, recording, "captions")
    return hand_over(
        body, export_name(recording, "captions.srt"), "application/x-subrip"
    )


@login_required
def batch_download(request: HttpRequest, batch_id) -> HttpResponse:
    """The Batch page's one button: every Done Recording of this Batch, as text.

    Recordings still in the queue are simply left out, which is why the button
    says how many are not ready. So is a Recording in a Case that is out of
    reach, because a hidden Case is hidden from every list, page, and zip.
    """
    from core import cases
    from core.recordings import Batch

    batch = Batch.objects.filter(pk=batch_id).select_related("user").first()
    if batch is None or (
        batch.user_id != request.user.pk and not request.user.is_admin
    ):
        return HttpResponseNotFound("There is no such batch.")

    wanted = batch.recordings.select_related("transcript", "user").order_by("created")
    if not cases.folder_management_on():
        wanted = wanted.filter(case__isnull=True)

    body, included = transcripts_zip(wanted)
    for recording in included:
        record_export(request, recording, "transcript text")

    return hand_over(body, zip_name(), "application/zip")


@login_required
def workspace_download(request: HttpRequest, shape: str) -> HttpResponse:
    """The sign-out dialog's downloads: every Done Recording in the Workspace.

    Not offered on the Batch page, which has its own download of its own
    Recordings only. The Workspace is what has no Case: what is in a Case is
    kept, so it is not among the things a person is being offered before they
    lose them.
    """
    recordings = (
        Recording.objects.filter(user=request.user, case__isnull=True)
        .select_related("transcript", "user")
        .order_by("created")
    )

    if shape == "everything":
        body, made, clips = everything_zip(recordings, request.user.username)
        for recording, kind in made:
            record_export(request, recording, kind)
        for clip in clips:
            _record_clip(request, clip)
        return hand_over(body, zip_name("Everything"), "application/zip")

    body, included = transcripts_zip(recordings)
    for recording in included:
        record_export(request, recording, "transcript text")
    return hand_over(body, zip_name(), "application/zip")


# The Summary and Chat exports -----------------------------------------------------
#
# Each on its own, in layout Record's typography: the cover facts with no
# Appearances table and no Processing record, the header block, the AI notice
# above the Transcript's own notice, then the body with Citations printed as
# [hh:mm:ss] in grey. Summaries and Chat answers are English whatever the
# Transcript's language.


def _open_record_document():
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
    return document


def _cover(document, recording, transcript, kind: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    heading = document.add_paragraph(title_of(recording))
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(20)
    said = document.add_paragraph(kind)
    said.alignment = WD_ALIGN_PARAGRAPH.CENTER
    said.runs[0].font.size = Pt(13)
    document.add_paragraph()
    in_a_case = (
        [("Case", recording.case.name)] if getattr(recording, "case", None) else []
    )
    _facts(
        document,
        [
            ("Recording", recording.original_filename),
            *in_a_case,
            ("Length", length_of(recording)),
            ("Uploaded", f"{recording.created:{DAY}} by {recording.user.username}"),
            ("Language", language_name(transcript.language)),
            (
                "Processed",
                f"{transcript.created:{DAY}} with Whisper {model_of(transcript)}",
            ),
            ("SHA-256", recording.sha256),
        ],
        Inches,
    )
    document.add_paragraph()


def _notices(document, transcript, ai_notice: str) -> None:
    for text in (ai_notice, notice_for(transcript)):
        if text:
            line = document.add_paragraph(text)
            line.runs[0].italic = True


CITED = re.compile(r"(\[\d{1,2}:\d{2}:\d{2}\])")


def _paragraph_with_citations(document, text: str, citations: dict, style=None):
    """One paragraph, the matched times in grey, everything else as it came."""
    from docx.shared import RGBColor

    paragraph = (
        document.add_paragraph(style=style) if style else document.add_paragraph()
    )
    for piece in CITED.split(text):
        if not piece:
            continue
        run = paragraph.add_run(piece)
        if piece in citations:
            run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    return paragraph


def _summary_body(document, summary) -> None:
    """The template's parts as Heading 2, read off the lines that end in a colon."""
    for raw in summary.text.splitlines():
        line = raw.rstrip()
        if not line:
            continue
        stripped = line.lstrip("#* ").rstrip()
        if stripped.endswith(":") and len(stripped) <= 60:
            document.add_heading(stripped[:-1], level=2)
            continue
        if stripped.startswith(("- ", "* ", "• ")):
            _paragraph_with_citations(
                document, stripped[2:], summary.citations, style="List Bullet"
            )
            continue
        _paragraph_with_citations(document, line, summary.citations)
    if summary.cut_short:
        note = document.add_paragraph("The answer was cut short.")
        note.runs[0].italic = True


def summary_word(summary, exported_by: str) -> bytes:
    from docx.shared import Inches

    from core import assistant

    recording = summary.recording
    transcript = recording.transcript
    document = _open_record_document()
    _cover(document, recording, transcript, f"Summary ({summary.template_name})")
    written = summary.written_at or summary.created
    _facts(
        document,
        [
            ("Template", f"{summary.template_name} v{summary.template_version}"),
            ("Focus", summary.focus or "none"),
            ("Length", summary.length.capitalize()),
            (
                "Written",
                f"{written:{DAY_AND_TIME}} by {summary.model or 'the AI assistant'}",
            ),
            ("Exported", f"{datetime.now():{DAY_AND_TIME}} by {exported_by}"),
        ],
        Inches,
    )
    document.add_paragraph()
    _notices(document, transcript, assistant.notice(summary.model, written))
    document.add_paragraph()
    _summary_body(document, summary)
    holder = io.BytesIO()
    document.save(holder)
    return holder.getvalue()


def chat_word(chat, exported_by: str) -> bytes:
    from docx.shared import Inches

    from core import assistant

    recording = chat.recording
    transcript = recording.transcript
    turns = list(chat.turns.filter(state="done"))
    document = _open_record_document()
    _cover(document, recording, transcript, "Chat")
    first_when = turns[0].answered_at if turns else chat.created
    first_model = next((one.model for one in turns if one.model), "")
    _facts(
        document,
        [
            ("Started", f"{chat.created:{DAY_AND_TIME}}"),
            ("Questions", str(len(turns))),
            ("Model", first_model or "the AI assistant"),
            ("Exported", f"{datetime.now():{DAY_AND_TIME}} by {exported_by}"),
        ],
        Inches,
    )
    document.add_paragraph()
    _notices(document, transcript, assistant.notice(first_model, first_when))
    for turn in turns:
        document.add_paragraph()
        document.add_heading(f"Question {turn.number}.", level=2)
        asked = document.add_paragraph(turn.question)
        asked.runs[0].italic = True
        for paragraph in turn.answer.split("\n"):
            if paragraph.strip():
                _paragraph_with_citations(document, paragraph, turn.citations)
        if turn.cut_short:
            note = document.add_paragraph("The answer was cut short.")
            note.runs[0].italic = True
    holder = io.BytesIO()
    document.save(holder)
    return holder.getvalue()


def summary_name(summary) -> str:
    when = summary.written_at or summary.created
    return (
        f"{safe_name(title_of(summary.recording))} - summary "
        f"({safe_name(summary.template_name)}) {when:%Y-%m-%d %H%M}.docx"
    )


def chat_name(chat) -> str:
    return (
        f"{safe_name(title_of(chat.recording))} - chat "
        f"{chat.created:%Y-%m-%d %H%M}.docx"
    )
