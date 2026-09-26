"""The Timeline report (Phase 8 chapter 14).

The incident's chronology as a document a person hands over: the cover, the
cameras, the band of their spans, then the events down the pages in spells,
each entry with a still from its camera at its moment, the words heard on
that camera around it, what the camera showed, the note and the why. Made
from the chronology as it stands, never kept; the stills are cut for the
document and deleted with it. The chronology is where the events are
decided; this report only prints them.
"""

from __future__ import annotations

import contextlib
import io
import re
import shutil
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from django.utils import timezone

from core import assistant, chronology, incidents, prompts, settings_store

# The still's width on the page; its height follows the frame.
STILL_WIDTH_INCHES = 2.4
# How many spoken lines an entry quotes at most, the nearest to the moment.
MOST_WORDS = 8
# How many camera lines an entry carries at most.
MOST_CAMERA_LINES = 3
# Stills are cut a few at a time so a long chronology stays inside a request.
CUTTERS = 4
MUTED = (0x5B, 0x66, 0x73)

DIGEST_LINE = re.compile(
    r"^\s*\d*[.)]?\s*\[(\d{1,2}):(\d{2}):(\d{2})\](?:-\[\d{1,2}:\d{2}:\d{2}\])?\s*"
    r"\((said|seen|both)\)\s*(.*)$"
)


# The switches -----------------------------------------------------------------------


def on() -> bool:
    return incidents.on() and bool(settings_store.get("incidents_report"))


def stills_on() -> bool:
    return on() and bool(settings_store.get("incidents_report_stills"))


def json_for(incident) -> dict:
    """What the incident page is told: the export offered, and whether there
    is anything to print yet."""
    return {
        "on": on(),
        "events": incident.events.filter(proposed=False).count(),
    }


def export_name(incident) -> str:
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in incident.name)
    return f"Timeline report - {safe.strip() or 'incident'}.docx"


# What each entry draws on ------------------------------------------------------


def camera_of(event, cameras: dict):
    """The event's own camera when it is synced, else the first synced camera
    the event is seen on, else None: the rule the page uses to bring a camera
    to the front (chapter 12)."""
    own = cameras.get(str(event.camera_id))
    if own is not None:
        return own
    for one in event.cameras:
        if one in cameras:
            return cameras[one]
    return None


def moment_on(camera, at: float) -> float | None:
    """The event's moment as seconds into the camera's recording, or None
    when the camera was not running then."""
    into = at - camera.starts_at
    length = float(camera.recording.duration_seconds or 0.0)
    if into < 0 or (length and into > length):
        return None
    return into


def words_near(incident, camera, at: float, window: int) -> list[str]:
    """The lines heard on the camera within the window either side of the
    moment, the nearest first chosen and then put in time order, each as its
    time of day and its words as the transcript labels them."""
    if window <= 0:
        return []
    transcript = getattr(camera.recording, "transcript", None)
    if transcript is None:
        return []
    near = []
    for line in prompts.lines_of(transcript):
        when = camera.starts_at + line.start
        if abs(when - at) <= window:
            near.append((abs(when - at), when, line))
    near.sort(key=lambda one: one[0])
    chosen = sorted(near[:MOST_WORDS], key=lambda one: one[1])
    out = []
    for _, when, line in chosen:
        label = prompts.plain_speaker(line.speaker)
        who = f"{label}: " if label else ""
        out.append(f"{incidents.time_of_day(incident, when)}  {who}{line.text}")
    return out


def camera_lines_near(camera, at: float, window: int) -> list[str]:
    """What the camera's Digest says it showed within the window: the lines
    marked (seen) or (both), their span dropped, at most a few."""
    if window <= 0:
        return []
    transcript = getattr(camera.recording, "transcript", None)
    if transcript is None or not assistant.digests_on():
        return []
    out = []
    for raw in assistant.digest_text(transcript).splitlines():
        match = DIGEST_LINE.match(raw)
        if not match:
            continue
        hours, minutes, seconds, mark, text = match.groups()
        if mark == "said":
            continue
        own = int(hours) * 3600 + int(minutes) * 60 + int(seconds)
        if abs(camera.starts_at + own - at) <= window:
            out.append(text.strip())
        if len(out) >= MOST_CAMERA_LINES:
            break
    return out


def cut_stills(
    plan: list[tuple[int, object, float]], folder: Path, height: int
) -> dict:
    """One JPEG per planned entry, cut a few at a time; an entry whose frame
    cannot be cut goes without, and the document still prints."""
    from core import media

    def one(item):
        index, camera, into = item
        source = camera.recording.playback_path()
        if source is None or not Path(source).exists():
            return index, None
        target = folder / f"still-{index}"
        target.mkdir(parents=True, exist_ok=True)
        try:
            written = media.grab_frames(Path(source), target, [into], height=height)
        except Exception:  # noqa: BLE001 - a still is never worth failing the report
            return index, None
        return index, (written[0] if written else None)

    if not plan:
        return {}
    with ThreadPoolExecutor(max_workers=CUTTERS) as pool:
        return dict(pool.map(one, plan))


# The document -----------------------------------------------------------------------


def word(incident, exported_by: str) -> bytes:
    """The Timeline report as a Word document, in the app's export shape."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    from core import exports
    from core.incident_assistant import synced_cameras

    rows = chronology._rows(incident)
    cameras = {str(one.pk): one for one in synced_cameras(incident)}
    window = int(settings_store.get("incidents_report_words_seconds") or 0)
    height = int(settings_store.get("incidents_report_still_height") or 360)
    most_stills = int(settings_store.get("incidents_report_stills_most") or 0)

    document = exports._open_record_document()
    exports._office_head(document)
    heading = document.add_paragraph(f"Timeline report: {incident.name}")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(20)
    said = document.add_paragraph(
        "What happened, on the cameras' clock, camera by camera"
    )
    said.alignment = WD_ALIGN_PARAGRAPH.CENTER
    said.runs[0].font.size = Pt(13)
    document.add_paragraph()
    all_cameras = list(incident.cameras.select_related("recording"))
    placed_words, _ = incidents.placed_words(incident)
    exports._facts(
        document,
        [
            ("Case", incident.case.name),
            ("Incident", incident.name),
            (
                "Clock",
                (
                    f"the time of day on {incident.clock_date}, by the cameras' clocks"
                    if incident.has_clock()
                    else "no camera clock; times count from the first camera"
                ),
            ),
            ("Span", incidents.span_words(incident)),
            ("Cameras", f"{len(all_cameras)}, {placed_words}"),
            ("Events", str(len(rows))),
            (
                "Made from",
                "the chronology as it stood at "
                f"{timezone.localtime():{exports.DAY_AND_TIME}}",
            ),
            ("Exported", f"{datetime.now():{exports.DAY_AND_TIME}} by {exported_by}"),
        ],
        Inches,
    )
    if incident.about:
        document.add_paragraph()
        for line in incident.about.splitlines():
            if line.strip():
                document.add_paragraph(line.strip())
    document.add_paragraph()

    # The cameras, and the band of their spans.
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, title in zip(
        table.rows[0].cells, ("Camera", "Recording", "Starts", "Ends"), strict=True
    ):
        cell.text = title
    for camera in sorted(
        all_cameras, key=lambda one: (one.starts_at is None, one.starts_at or 0.0)
    ):
        cells = table.add_row().cells
        cells[0].text = camera.camera_id()
        cells[1].text = camera.recording.title
        if camera.is_placed():
            cells[2].text = incidents.time_of_day(incident, camera.starts_at)
            ends = camera.ends_at()
            cells[3].text = incidents.time_of_day(incident, ends) if ends else ""
    document.add_paragraph()
    _band(document, incident, Inches, Pt, RGBColor)
    document.add_paragraph()
    how = document.add_paragraph()
    how_run = how.add_run(
        "Each entry is one event of the chronology: its number, its time of day "
        "and the camera it came from; the still is that camera's picture at the "
        "moment; the words are what was heard on that camera around it; a line "
        "beginning Camera is what the picture showed. Events run in time order in "
        "spells, a new spell wherever ten minutes or more pass without one."
    )
    how_run.italic = True
    how_run.font.size = Pt(9)
    how_run.font.color.rgb = RGBColor(*MUTED)

    if not rows:
        empty = document.add_paragraph("No events on the chronology yet.")
        empty.runs[0].italic = True
        exports.stamp_pages(
            document,
            f"{incident.case.name}, {incident.name}",
            "Timeline report",
            incident.name,
        )
        return _bytes(document)

    # The stills, cut before the entries are laid out.
    events = {str(one.pk): one for one in incident.events.filter(proposed=False)}
    plan: list[tuple[int, object, float]] = []
    chosen_camera: dict[int, object] = {}
    for index, row in enumerate(rows):
        event = events.get(row["id"])
        camera = camera_of(event, cameras) if event is not None else None
        if camera is None:
            continue
        chosen_camera[index] = camera
        into = moment_on(camera, row["seconds"])
        if into is not None and stills_on() and len(plan) < most_stills:
            plan.append((index, camera, into))
    folder = Path(tempfile.mkdtemp(prefix="timeline-report-"))
    try:
        stills = cut_stills(plan, folder, height) if plan else {}
        for group in chronology.spells(rows):
            _spell(
                document,
                incident,
                group,
                chosen_camera,
                stills,
                window,
                Inches,
                Pt,
                RGBColor,
            )
    finally:
        shutil.rmtree(folder, ignore_errors=True)

    document.add_paragraph()
    to_check = sum(1 for row in rows if row["to_check"])
    if to_check:
        line = document.add_paragraph(
            f"{to_check} event{'' if to_check == 1 else 's'} marked to check: "
            "the office has not settled the point."
        )
        line.runs[0].italic = True
    without = len(chosen_camera) - len([one for one in stills.values() if one])
    if stills_on() and without:
        line = document.add_paragraph(
            f"{without} event{' prints' if without == 1 else 's print'} "
            "without a still: past the report's limit, or the camera had no "
            "picture at that moment."
        )
        line.runs[0].italic = True
    for legend in (chronology.WORDS_LEGEND, chronology.QUOTE_LEGEND):
        line = document.add_paragraph(legend)
        line.runs[0].italic = True
    if any(row["is_note"] for row in rows):
        line = document.add_paragraph(chronology.NOTE_LEGEND)
        line.runs[0].italic = True
    if any(row["assistant"] for row in rows):
        notice = assistant.notice("", timezone.now())
        if notice:
            line = document.add_paragraph(notice)
            line.runs[0].italic = True
    line = document.add_paragraph(
        "A still is one frame of the camera's playback copy at the event's moment, "
        "cut for this document. The words quoted are the transcript's, as it "
        "labels the speakers; a camera line is a model's description of the "
        "picture, not the words. Every time is the cameras' clock."
    )
    line.runs[0].italic = True
    line.runs[0].font.size = Pt(9)
    exports.stamp_pages(
        document,
        f"{incident.case.name}, {incident.name}",
        "Timeline report",
        incident.name,
    )
    return _bytes(document)


def _band(document, incident, Inches, Pt, RGBColor) -> None:
    band = chronology.band_png(incident)
    if not band:
        return
    with contextlib.suppress(Exception):
        document.add_picture(io.BytesIO(band), width=Inches(6.5))
    legend = document.add_paragraph()
    colours = chronology.camera_colours(incident)
    cameras = sorted(
        (one for one in incident.cameras.all() if one.is_placed()),
        key=lambda one: (one.starts_at, one.added),
    )
    for index, camera in enumerate(cameras):
        run = legend.add_run(("   " if index else "") + "■ ")
        run.font.color.rgb = RGBColor.from_string(colours[str(camera.pk)].lstrip("#"))
        run.font.size = Pt(9)
        legend.add_run(camera.camera_id()).font.size = Pt(9)
    tail = legend.add_run("   Each bar is one camera's recording on the clock.")
    tail.font.size = Pt(8)
    tail.italic = True


def _spell(
    document, incident, group, chosen_camera, stills, window, Inches, Pt, RGBColor
) -> None:
    heading = document.add_heading(level=2)
    head = heading.add_run(chronology.spell_words(group))
    head.font.name = "Consolas"
    count = heading.add_run(f"   {len(group)} event{'' if len(group) == 1 else 's'}")
    count.font.size = Pt(10)
    count.font.color.rgb = RGBColor(*MUTED)
    for row in group:
        _entry(
            document, incident, row, chosen_camera, stills, window, Inches, Pt, RGBColor
        )


def _entry(
    document, incident, row, chosen_camera, stills, window, Inches, Pt, RGBColor
) -> None:
    index = row["number"] - 1
    camera = chosen_camera.get(index)
    still = stills.get(index)
    table = document.add_table(rows=1, cols=2 if still else 1)
    table.autofit = False
    cells = table.rows[0].cells
    if still:
        cells[0].width = Inches(STILL_WIDTH_INCHES + 0.1)
        picture = cells[0].paragraphs[0]
        with contextlib.suppress(Exception):
            picture.add_run().add_picture(str(still), width=Inches(STILL_WIDTH_INCHES))
        caption = cells[0].add_paragraph()
        caption_run = caption.add_run(f"{camera.camera_id()} at {row['time']}")
        caption_run.font.size = Pt(8)
        caption_run.font.color.rgb = RGBColor(*MUTED)
        text = cells[1]
        text.width = Inches(6.5 - STILL_WIDTH_INCHES - 0.1)
    else:
        cells[0].width = Inches(6.5)
        text = cells[0]

    head = text.paragraphs[0]
    number = head.add_run(f"{row['number']}  ")
    number.bold = True
    if row["colour"]:
        number.font.color.rgb = RGBColor.from_string(row["colour"].lstrip("#"))
    when = head.add_run(row["time"] + (f" to {row['end']}" if row["end"] else ""))
    when.font.name = "Consolas"
    when.font.size = Pt(10)
    if camera is not None:
        where = head.add_run(f"   {camera.camera_id()}")
        where.font.size = Pt(9)
        where.font.color.rgb = RGBColor(*MUTED)
    line = text.add_paragraph()
    words = line.add_run(row["line"])
    words.bold = True
    if row["to_check"]:
        mark = line.add_run("  To check")
        mark.font.size = Pt(9)
        mark.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)
    if row["detail"]:
        detail = text.add_paragraph()
        detail.add_run(row["detail"]).font.size = Pt(10)

    if camera is not None:
        heard = words_near(incident, camera, row["seconds"], window)
        if heard:
            lead = text.add_paragraph()
            lead_run = lead.add_run(f"Said on {camera.camera_id()}")
            lead_run.font.size = Pt(8)
            lead_run.font.color.rgb = RGBColor(*MUTED)
            for one in heard:
                quoted = text.add_paragraph()
                time_part, _, rest = one.partition("  ")
                mono = quoted.add_run(time_part + "  ")
                mono.font.name = "Consolas"
                mono.font.size = Pt(8)
                mono.font.color.rgb = RGBColor(*MUTED)
                quoted.add_run(rest).font.size = Pt(9)
        for one in camera_lines_near(camera, row["seconds"], window):
            seen = text.add_paragraph()
            seen_lead = seen.add_run("Camera: ")
            seen_lead.font.size = Pt(9)
            seen_lead.bold = True
            seen_lead.font.color.rgb = RGBColor(*MUTED)
            seen_run = seen.add_run(one)
            seen_run.italic = True
            seen_run.font.size = Pt(9)

    under = []
    if row["note"]:
        under.append("Note: " + row["note"])
    if row["why"]:
        under.append("Why it matters: " + row["why"])
    under.append(
        row["source"] + (f"; seen on {row['seen_on']}" if row["seen_on"] else "")
    )
    small = text.add_paragraph()
    run = small.add_run("  ".join(under))
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(*MUTED)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _bytes(document) -> bytes:
    holder = io.BytesIO()
    document.save(holder)
    return holder.getvalue()
