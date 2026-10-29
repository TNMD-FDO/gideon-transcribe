"""Notes on a line of a transcript, and the case's Notes tab (Phase 8 chapter 2).

A note is a person's own line under a Segment, under the rules an Event's
note already has: up to 2,000 characters, kept with who wrote it and when it
last changed, written by a person and never by the assistant. This module
holds what the recording page, the exports, the case page, Gideon and Search
share about them; the Event's own note stays in chronology.py.

Chapter 11: a note on a line of a recording that is a synced camera of an
Incident is also an Event on that Incident's Chronology, at the line's moment
(a note event). This module is the one writer of those rows: they are made,
moved and removed here as the note and the camera come and go, and never
carry an audit row of their own.

v1.99.0: a note where nothing was said. A thing the camera showed in a
silence has no line to take a note, so a note written there is kept at the
moment itself (a MomentNote): one kind of note to the person, anchored by
the app on the line being spoken when there is one and at the moment when
there is not. It shows wherever a line's note shows, is an Event on the
Chronology of a synced camera at its exact moment, and is never a line: no
plain export, caption or reading of the transcript carries it. It belongs
to the Recording and not to the Transcript, so Process again leaves it be.
"""

from __future__ import annotations

import contextlib
from datetime import UTC, datetime
from pathlib import Path

from django.db import models
from django.urls import reverse
from django.utils import timezone

from core import audit, cases, clock_check, exports, incidents
from core.jobs import Segment

NOTE_MOST = 2000
TAB_MOST = 200
NEVER = datetime(1970, 1, 1, tzinfo=UTC)
# A line counts as being spoken at a moment when the moment falls inside it
# or within this many seconds of its start or its end (v1.99.0).
SPOKEN_WITHIN = 3.0
# A silence this long or longer is said where the lines are shown.
SILENCE_SAID = 30.0
# Two notes at a moment closer than this are one note.
SAME_MOMENT = 1.0
# What the assistant is told of a note at a moment, so it looks for no line
# to match it. A person is never shown it (v1.99.1): the note has a row, a
# time and a colour of its own, and its words are enough.
NOTHING_SAID = "nothing was said here"


class MomentNote(models.Model):
    """A note kept at a moment of a Recording where nothing was being said."""

    recording = models.ForeignKey(
        "core.Recording", on_delete=models.CASCADE, related_name="moment_notes"
    )
    # Seconds from the recording's start.
    at = models.FloatField()
    note = models.TextField(max_length=NOTE_MOST)
    note_by = models.ForeignKey(
        "core.User", on_delete=models.SET_NULL, null=True, blank=True, related_name="+"
    )
    note_changed = models.DateTimeField(null=True, blank=True)
    # To check (Phase 10 chapter 1, v1.164.0), as on a line's note.
    note_to_check = models.BooleanField(default=False)
    # As the Segment's (v1.168.0).
    note_from = models.JSONField(default=dict, blank=True)
    created = models.DateTimeField(auto_now_add=True)

    class Meta:
        ordering = ["at", "id"]

    def __str__(self) -> str:
        return f"note at {self.at:.1f}s of {self.recording_id}"

    # Read as a line is read, by what shows a note event: its moment and no words.
    @property
    def start(self) -> float:
        return self.at

    @property
    def text(self) -> str:
        return ""


def clean(text) -> str:
    """The note as it is kept: lines tidied, capped at the chapter's length."""
    lines = [" ".join(line.split()) for line in str(text or "").splitlines()]
    return "\n".join(lines).strip()[:NOTE_MOST]


def date_of(when) -> str:
    """The date beside the writer, in the app's short form ("19 Sep 2026")."""
    if when is None:
        return ""
    return f"{timezone.localtime(when):%d %b %Y}"


def origin_words(anchor) -> str:
    """Where the note came from, in a few words, or nothing (v1.168.0)."""
    if not anchor.note:
        return ""
    return (anchor.note_from or {}).get("words", "")


def _origin_fields(anchor) -> dict:
    """What a note event carries from the note's origin: the why and what
    it rested on, so the chronology's hover and the exports have them."""
    from core import chronology

    origin = anchor.note_from or {}
    return {
        "why": (origin.get("why") or "")[: chronology.WHY_MOST],
        "rests_on": (origin.get("rests_on") or "")[: chronology.TEXT_MOST],
    }


def line_json(segment) -> dict:
    """The note's fields as the page and the tab read them."""
    return {
        "note": segment.note,
        "note_from": origin_words(segment),
        "note_by": (
            segment.note_by.shown_name if segment.note and segment.note_by else ""
        ),
        "note_on": date_of(segment.note_changed) if segment.note else "",
    }


def set_note(segment, text, *, by, request=None, where: str = "") -> dict:
    """Write, change or clear the note on a line; one audit row, never the words.

    `where` names the view it was done from when that is not the line itself
    ("chronology", chapter 11); the row carries it and nothing else changes.
    """
    wanted = clean(text)
    had = segment.note
    if wanted == had:
        return line_json(segment)
    segment.note = wanted
    # The writer is whoever last wrote it; a cleared note has no writer.
    segment.note_by = by if wanted else None
    segment.note_changed = timezone.now() if wanted else None
    if not wanted:
        segment.note_from = {}
    segment.save(update_fields=["note", "note_by", "note_changed", "note_from"])
    reflect_line(segment)
    recording = segment.transcript.recording
    audit.write(
        audit.Category.EDITS,
        "Note removed" if not wanted else ("Note changed" if had else "Note added"),
        actor=by,
        request=request,
        affected_user=cases.affected_by(recording, by),
        object_type="segment",
        object_id=segment.pk,
        # The recording and the time (v1.123.0), as every other row reads.
        object_label=(
            f"{recording.title or recording.original_filename} "
            f"at {exports.clock(segment.start)}"
        )[:400],
        **({"where": where} if where else {}),
    )
    if recording.case_id:
        cases.note_activity(recording.case, by=by)
    return line_json(segment)


# The note at a moment (v1.99.0) --------------------------------------------------


def spoken_at(transcript, at: float):
    """The line being spoken at a moment: the one the moment falls inside,
    else the nearest whose start or end is within SPOKEN_WITHIN seconds; None
    when nothing is being said."""
    if transcript is None:
        return None
    near = list(
        transcript.segments.filter(
            same_as_other_side=False,
            start__lte=at + SPOKEN_WITHIN,
            end__gte=at - SPOKEN_WITHIN,
        ).order_by("start", "id")
    )
    inside = [one for one in near if one.start <= at <= one.end]
    if inside:
        return inside[-1]
    if not near:
        return None
    return min(near, key=lambda one: min(abs(one.start - at), abs(one.end - at)))


def silence_at(transcript, at: float) -> dict | None:
    """The silence a moment falls in, when it is SILENCE_SAID long or longer:
    from the end of the last line before it to the start of the next."""
    if transcript is None:
        return None
    lines = transcript.segments.filter(same_as_other_side=False)
    if lines.filter(start__lte=at, end__gte=at).exists():
        return None
    before = lines.filter(end__lt=at).order_by("-end").first()
    after = lines.filter(start__gt=at).order_by("start").first()
    low = before.end if before is not None else 0.0
    length = transcript.recording.duration_seconds or at
    high = after.start if after is not None else max(float(length), at)
    if high - low < SILENCE_SAID:
        return None
    return {"from": low, "to": high, "seconds": round(high - low)}


def moment_json(one) -> dict:
    """A note at a moment as the pages read it."""
    return {
        "id": one.pk,
        "at": one.at,
        "clock": exports.clock(one.at),
        "note": one.note,
        "note_from": origin_words(one),
        "note_by": one.note_by.shown_name if one.note_by else "",
        "note_on": date_of(one.note_changed),
    }


def moments_of(recording, low: float | None = None, high: float | None = None):
    """The recording's notes at moments, in time order, within a span when given."""
    found = MomentNote.objects.filter(recording=recording).select_related("note_by")
    if low is not None:
        found = found.filter(at__gte=low)
    if high is not None:
        found = found.filter(at__lte=high)
    return found.order_by("at", "id")


def _moment_row(one, event: str, *, by, request, where: str) -> None:
    recording = one.recording
    audit.write(
        audit.Category.EDITS,
        event,
        actor=by,
        request=request,
        affected_user=cases.affected_by(recording, by),
        object_type="moment",
        object_id=recording.pk,
        object_label=f"{one.at:.1f}",
        **({"where": where} if where else {}),
    )
    if recording.case_id:
        cases.note_activity(recording.case, by=by)


def note_at(recording, at: float, text, *, by, request=None, where: str = "") -> dict:
    """Write a note at a moment where nothing is being said: a new one, or the
    one already within a second of it. One audit row, never the words.
    Nothing is kept for an empty note."""
    wanted = clean(text)
    at = max(0.0, float(at))
    one = (
        MomentNote.objects.filter(
            recording=recording, at__gte=at - SAME_MOMENT, at__lte=at + SAME_MOMENT
        )
        .order_by("at")
        .first()
    )
    if one is not None:
        return set_moment_note(one, wanted, by=by, request=request, where=where)
    if not wanted:
        return {
            "id": None,
            "at": at,
            "clock": exports.clock(at),
            "note": "",
            "note_by": "",
            "note_on": "",
        }
    one = MomentNote.objects.create(
        recording=recording, at=at, note=wanted, note_by=by, note_changed=timezone.now()
    )
    reflect_moment(one)
    _moment_row(one, "Note added", by=by, request=request, where=where)
    return moment_json(one)


def set_moment_note(one, text, *, by, request=None, where: str = "") -> dict:
    """Change the note at a moment, or remove it when the words are cleared."""
    wanted = clean(text)
    if wanted == one.note:
        return moment_json(one)
    if not wanted:
        told = {
            **moment_json(one),
            "id": None,
            "note": "",
            "note_by": "",
            "note_on": "",
        }
        _moment_row(one, "Note removed", by=by, request=request, where=where)
        # Its events on the chronologies go with it.
        one.delete()
        return told
    one.note = wanted
    one.note_by = by
    one.note_changed = timezone.now()
    one.save(update_fields=["note", "note_by", "note_changed"])
    reflect_moment(one)
    _moment_row(one, "Note changed", by=by, request=request, where=where)
    return moment_json(one)


def reflect_moment(one) -> None:
    """The note events of one note at a moment, made right: one on every
    Incident the recording is a synced camera of, at the exact moment on the
    cameras' clock; none on an Incident where it is not synced."""
    from core.chronology import NOTE, Event

    for camera in incidents.IncidentCamera.objects.filter(
        recording=one.recording
    ).select_related("incident"):
        if not camera.is_synced():
            Event.objects.filter(incident=camera.incident, moment_note=one).delete()
            continue
        at = camera.starts_at + one.at
        event = Event.objects.filter(incident=camera.incident, moment_note=one).first()
        if event is None:
            Event.objects.create(
                incident=camera.incident,
                moment_note=one,
                at=at,
                text=one.note,
                source=NOTE,
                camera=camera,
                cameras=[str(camera.pk)],
                added_by=one.note_by,
                **_origin_fields(one),
            )
            continue
        event.at = at
        event.text = one.note
        event.source = NOTE
        event.camera = camera
        for key, value in _origin_fields(one).items():
            setattr(event, key, value)
        event.save(update_fields=["at", "text", "source", "camera", "why", "rests_on"])


# The note on the chronology (chapter 11) -------------------------------------------


def reflect_line(segment) -> None:
    """The note events of one line, made right: one on every Incident the
    recording is a synced camera of while the line is shown and noted, none
    otherwise. A sync never touches the Event's own changed marks."""
    from core.chronology import NOTE, Event

    recording = segment.transcript.recording
    wanted = bool(segment.note) and not segment.same_as_other_side
    for camera in incidents.IncidentCamera.objects.filter(
        recording=recording
    ).select_related("incident"):
        if not (wanted and camera.is_synced()):
            Event.objects.filter(incident=camera.incident, segment=segment).delete()
            continue
        at = camera.starts_at + segment.start
        event = Event.objects.filter(incident=camera.incident, segment=segment).first()
        if event is None:
            Event.objects.create(
                incident=camera.incident,
                segment=segment,
                at=at,
                text=segment.note,
                source=NOTE,
                camera=camera,
                # Its own camera alone (v1.85.2): a note is on one line of one
                # camera, and "seen on" every camera running at that moment
                # read as if the note were on each. The person may add more.
                cameras=[str(camera.pk)],
                added_by=segment.note_by,
                **_origin_fields(segment),
            )
            continue
        event.at = at
        event.text = segment.note
        event.source = NOTE
        event.camera = camera
        for key, value in _origin_fields(segment).items():
            setattr(event, key, value)
        event.save(update_fields=["at", "text", "source", "camera", "why", "rests_on"])


def reflect_camera(camera) -> None:
    """Every note event of one camera, made right: called when the camera is
    placed, re-placed or guessed, so a note written before the recording
    joined appears the moment the camera is synced, and moves with it."""
    from core.chronology import Event

    transcript = getattr(camera.recording, "transcript", None)
    if transcript is None:
        for one in MomentNote.objects.filter(recording=camera.recording):
            reflect_moment(one)
        return
    noted = list(transcript.segments.exclude(note="").filter(same_as_other_side=False))
    for segment in noted:
        reflect_line(segment)
    for one in MomentNote.objects.filter(recording=camera.recording):
        reflect_moment(one)
    # A row whose line is no longer noted, or is hidden, is an orphan.
    Event.objects.filter(
        incident=camera.incident,
        segment__transcript=transcript,
    ).exclude(segment__in=noted).delete()


def drop_note_events(recording, incident=None) -> None:
    """A camera leaving an Incident, or a recording leaving its case, takes
    its note events with it; the notes stay on the lines."""
    from core.chronology import Event

    found = Event.objects.filter(
        models.Q(segment__transcript__recording=recording)
        | models.Q(moment_note__recording=recording)
    )
    if incident is not None:
        found = found.filter(incident=incident)
    found.delete()


# Processing again -----------------------------------------------------------------


def remember(recording) -> list[dict]:
    """The notes of the Transcript about to be replaced, by their moments.

    Their note events (chapter 11) are let go of their lines first, so the
    Transcript's deletion does not take them: the rows, with their marks and
    their clips, wait for carry() to give them the new lines."""
    from core.chronology import Event

    transcript = getattr(recording, "transcript", None)
    if transcript is None:
        return []
    kept = [
        {
            "start": one.start,
            "end": one.end,
            "note": one.note,
            "note_by_id": one.note_by_id,
            "note_changed": one.note_changed,
            "events": [
                str(pk)
                for pk in Event.objects.filter(segment=one).values_list("pk", flat=True)
            ],
        }
        for one in transcript.segments.exclude(note="").order_by("start", "id")
    ]
    Event.objects.filter(segment__transcript=transcript).update(segment=None)
    return kept


def carry(kept: list[dict], transcript) -> int:
    """Put each remembered note on the new line that spans its moment, else the
    line with the nearest start. Two notes landing on one line are joined.
    Returns how many were carried; one audit row each, never the words."""
    from core.chronology import Event

    carried = 0
    lines = list(
        transcript.segments.filter(same_as_other_side=False).order_by("start", "id")
    )
    for old in kept if lines else []:
        at = old["start"]
        home = next((one for one in lines if one.start <= at < one.end), None)
        if home is None:
            home = min(lines, key=lambda one: abs(one.start - at))
        joined = bool(home.note)
        home.note = clean(
            old["note"] if not home.note else home.note + "\n" + old["note"]
        )
        home.note_by_id = old["note_by_id"]
        home.note_changed = old["note_changed"] or timezone.now()
        home.save(update_fields=["note", "note_by", "note_changed"])
        # The note's events go to the new line (chapter 11): when two notes
        # join on one line, the line keeps the first's rows and the second's
        # go, as one line holds one note event per Incident.
        waiting = Event.objects.filter(pk__in=old.get("events", []))
        if joined:
            waiting.delete()
        else:
            waiting.update(segment=home)
        reflect_line(home)
        audit.write(
            audit.Category.EDITS,
            "Note carried",
            system="process again",
            affected_user=transcript.recording.user,
            object_type="segment",
            object_id=home.pk,
            object_label=f"{home.start:.1f}-{home.end:.1f}",
            was=f"{old['start']:.1f}-{old['end']:.1f}",
        )
        carried += 1
    # A note event with no line is an orphan (a note no line could take, or
    # a Transcript with no lines): gone.
    Event.objects.filter(
        source="note",
        segment__isnull=True,
        moment_note__isnull=True,
        incident__cameras__recording=transcript.recording,
    ).delete()
    return carried


# The case's notes ---------------------------------------------------------------------


def _page_url(recording, key: str) -> str:
    """The Notes page (Phase 9 chapter 4) opened on this note, or nothing
    outside a Case."""
    if not recording.case_id:
        return ""
    return f"{reverse('notes-page', args=[recording.case_id])}?note={key}"


def _line_rows(case, links: dict | None = None) -> list[dict]:
    # The All cameras links come from one query for the case (Phase 9
    # chapter 1), not one per row.
    links = incidents.links_in(case) if links is None else links
    rows = []
    found = (
        Segment.objects.filter(transcript__recording__case=case)
        .exclude(note="")
        .filter(same_as_other_side=False)
        .select_related("transcript__recording", "note_by")
    )
    for one in found:
        recording = one.transcript.recording
        viewer = reverse("viewer", args=[recording.pk])
        rows.append(
            {
                "kind": "line",
                "id": one.pk,
                "text": one.note,
                "by": one.note_by.shown_name if one.note_by else "",
                "by_id": one.note_by_id,
                "on": date_of(one.note_changed),
                "changed": one.note_changed,
                "where": recording.title,
                "recording_id": recording.pk,
                "type": recording.recording_type,
                "created": recording.created,
                "at": one.start,
                "when": exports.clock(one.start),
                "url": f"{viewer}?t={one.start:.1f}&note=1",
                "page_url": _page_url(recording, f"line:{one.pk}"),
                "all_cameras": incidents.url_at(links.get(recording.pk), one.start),
                "rests_on": one.text,
                "who": one.speaker,
                "from": origin_words(one),
                # The clock verdict of the line (ADR 0023), for the export's pill.
                "verdict": clock_check.verdict_json(one),
            }
        )
    return rows


def _moment_rows(case, links: dict | None = None) -> list[dict]:
    """The notes at moments where nothing was said (v1.99.0), listed with the
    notes on lines."""
    links = incidents.links_in(case) if links is None else links
    rows = []
    found = MomentNote.objects.filter(recording__case=case).select_related(
        "recording", "note_by"
    )
    for one in found:
        recording = one.recording
        viewer = reverse("viewer", args=[recording.pk])
        rows.append(
            {
                "kind": "moment",
                "from": origin_words(one),
                "id": one.pk,
                "text": one.note,
                "by": one.note_by.shown_name if one.note_by else "",
                "by_id": one.note_by_id,
                "on": date_of(one.note_changed),
                "changed": one.note_changed,
                "where": recording.title,
                "recording_id": recording.pk,
                "type": recording.recording_type,
                "created": recording.created,
                "at": one.at,
                "when": exports.clock(one.at),
                "url": f"{viewer}?t={one.at:.1f}&note=1",
                "page_url": _page_url(recording, f"moment:{one.pk}"),
                "all_cameras": incidents.url_at(links.get(recording.pk), one.at),
                "rests_on": "",
                "who": "",
            }
        )
    return rows


def counts_in(case) -> dict:
    """How many notes each recording in a Case has and who wrote them, for the
    Recordings tab's Notes column (Phase 9 chapter 1): by recording id,
    {"count": n, "writers": [names, the busiest first]}. Three queries for the
    whole case, whatever its size."""
    from core.models import User

    tally: dict = {}

    def add(recording_id, writer_id):
        one = tally.setdefault(recording_id, {"count": 0, "by": {}})
        one["count"] += 1
        if writer_id is not None:
            one["by"][writer_id] = one["by"].get(writer_id, 0) + 1

    lines = (
        Segment.objects.filter(transcript__recording__case=case)
        .exclude(note="")
        .filter(same_as_other_side=False)
        .values_list("transcript__recording_id", "note_by_id")
    )
    for recording_id, writer_id in lines:
        add(recording_id, writer_id)
    for recording_id, writer_id in MomentNote.objects.filter(
        recording__case=case
    ).values_list("recording_id", "note_by_id"):
        add(recording_id, writer_id)
    writer_ids = {w for one in tally.values() for w in one["by"]}
    names = {}
    if writer_ids:
        for person in User.objects.filter(pk__in=writer_ids):
            names[person.pk] = person.shown_name
    counts = {}
    for recording_id, one in tally.items():
        ordered = sorted(
            one["by"].items(), key=lambda item: (-item[1], names.get(item[0], ""))
        )
        counts[recording_id] = {
            "count": one["count"],
            "writers": [names.get(w, "") for w, _ in ordered if names.get(w)],
        }
    return counts


def count_words(found: dict | None) -> str:
    """The Notes column's words: "3 notes, A. Okafor and D. Chen", two writers
    named and then "and N more"; nothing when there is no note."""
    if not found or not found["count"]:
        return ""
    count = found["count"]
    said = f"{count} note{'s' if count != 1 else ''}"
    writers = found["writers"]
    if len(writers) == 1:
        said += f", {writers[0]}"
    elif len(writers) == 2:
        said += f", {writers[0]} and {writers[1]}"
    elif len(writers) > 2:
        said += f", {writers[0]}, {writers[1]} and {len(writers) - 2} more"
    return said


def _event_rows(case) -> list[dict]:
    if not incidents.on():
        return []
    from core.chronology import Event

    rows = []
    # A note event is a line's note, listed once as that (chapter 11).
    found = (
        Event.objects.filter(
            incident__case=case,
            proposed=False,
            segment__isnull=True,
            moment_note__isnull=True,
        )
        .exclude(note="")
        .select_related("incident", "note_by")
    )
    # The camera the event sits on and the moment by its clock (v1.166.0), so
    # the Notes tab's time can play it in the Preview: the chronology's own
    # reading, one query for the case.
    from core import notes_page

    placed = {
        row["id"]: row
        for row in notes_page._event_rows(case, list(case.recordings.all()))
    }
    for event in found:
        incident = event.incident
        camera = placed.get(str(event.pk))
        rows.append(
            {
                "kind": "event",
                "recording_id": camera["recording"] if camera else "",
                "at": camera["at"] if camera else None,
                "text": event.note,
                "by": event.note_by.shown_name if event.note_by else "",
                "on": date_of(event.note_changed),
                "changed": event.note_changed,
                "where": incident.name,
                "when": incidents.time_of_day(incident, event.at),
                "url": f"{incident.url()}?t={event.at:.2f}",
                "all_cameras": "",
                "rests_on": event.text,
                "who": "",
            }
        )
    return rows


def of_case(case, kind: str = "") -> dict:
    """Every note in the case, newest first, for the Notes tab and its export."""
    # The notes at moments are listed with the notes on lines (v1.99.0).
    links = incidents.links_in(case)
    moments = _moment_rows(case, links)
    lines = _line_rows(case, links) + moments
    events = _event_rows(case)
    rows = {"lines": lines, "events": events}.get(kind, lines + events)
    # Newest first; a note without a date (none is written that way) goes last.
    rows.sort(key=lambda one: one["changed"] or NEVER, reverse=True)
    return {
        "rows": rows[:TAB_MOST],
        "all": rows,
        "more": max(0, len(rows) - TAB_MOST),
        "lines": len(lines),
        "events": len(events),
        "moments": len(moments),
        "total": len(lines) + len(events),
        "kind": kind if kind in ("lines", "events") else "",
    }


def count(case) -> int:
    """The Notes tab's count: on lines and on events."""
    lines = (
        Segment.objects.filter(transcript__recording__case=case)
        .exclude(note="")
        .filter(same_as_other_side=False)
        .count()
    ) + MomentNote.objects.filter(recording__case=case).count()
    if not incidents.on():
        return lines
    from core.chronology import Event

    return (
        lines
        + Event.objects.filter(
            incident__case=case,
            proposed=False,
            segment__isnull=True,
            moment_note__isnull=True,
        )
        .exclude(note="")
        .count()
    )


def word(case, exported_by: str) -> bytes:
    """Download notes: every note in the case as a Word document, for the
    office's own reading. One block per note (v1.156.0; the drawing
    phase-9-notes-export-and-names.html, figure 1): the note first in the
    largest type, where it rests under it, and for a note on a camera
    recording a still from the playback copy at the moment, taken now and
    kept nowhere. Grouped under the recording, newest recording first, and
    an incident's event notes under the incident."""
    import io
    import tempfile

    from django.conf import settings
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    listed = of_case(case)
    document = exports._open_record_document()
    exports._office_head(document)
    heading = document.add_paragraph(f"Notes: {case.name}")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(20)
    total = listed["total"]
    made = f"{timezone.localtime():%d %B %Y %H:%M}"
    line = document.add_paragraph(
        f"{total} note{'' if total == 1 else 's'}: {listed['lines']} on lines of "
        f"transcripts ({listed['moments']} of them at a moment where nothing was "
        f"said), {listed['events']} on events. Made {made} by {exported_by}. "
        "The office's own words, for the office's own reading."
    )
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line.runs[0].italic = True

    rows = listed["all"]
    groups = _grouped(rows)
    Path(settings.SCRATCH_DIR).mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=str(settings.SCRATCH_DIR)) as folder:
        stills, left_out = _stills_for(rows, Path(folder))
        if left_out:
            note = document.add_paragraph(
                f"{len(stills)} stills are in this document; {left_out} more "
                f"camera notes are without one, since the document carries at "
                f"most {STILLS_MOST}."
            )
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            note.runs[0].italic = True
            note.runs[0].font.size = Pt(9)
        for group in groups:
            head = document.add_paragraph()
            head.paragraph_format.space_before = Pt(14)
            head.paragraph_format.space_after = Pt(4)
            title = head.add_run(group["title"])
            title.bold = True
            title.font.size = Pt(12)
            if group["sub"]:
                sub = head.add_run(f"   {group['sub']}")
                sub.font.size = Pt(9)
                sub.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            for one in group["rows"]:
                _note_block(document, one, stills.get(_key(one)), Inches, Pt, RGBColor)
    # The running head and the page marking every Word export carries (v1.121.3).
    exports.stamp_pages(document, case.name, "Notes", case.name)
    holder = io.BytesIO()
    document.save(holder)
    return holder.getvalue()


# The document carries at most this many stills, so a download never runs away.
STILLS_MOST = 300
STILL_HEIGHT = 360
STILL_WIDTH_INCHES = 2.0


def _key(row: dict) -> str:
    return f"{row['kind']}:{row.get('id', '')}"


def _grouped(rows: list[dict]) -> list[dict]:
    """The notes under their recording, newest recording first as the Notes
    page lists them, then the incidents' event notes under their incident."""
    by_recording: dict = {}
    by_incident: dict = {}
    for one in rows:
        if one["kind"] == "event":
            by_incident.setdefault(one["where"], []).append(one)
        else:
            by_recording.setdefault(one.get("recording_id"), []).append(one)
    groups = []
    for _recording_id, mine in sorted(
        by_recording.items(),
        key=lambda item: item[1][0].get("created") or 0,
        reverse=True,
    ):
        first = mine[0]
        made = first.get("created")
        groups.append(
            {
                "title": first["where"],
                "sub": " · ".join(
                    part
                    for part in (
                        first.get("type") or "",
                        f"{timezone.localtime(made):%d %B %Y}"
                        if hasattr(made, "year")
                        else "",
                        f"{len(mine)} note{'' if len(mine) == 1 else 's'}",
                    )
                    if part
                ),
                "rows": sorted(mine, key=lambda one: float(one.get("at") or 0)),
            }
        )
    for name, mine in by_incident.items():
        groups.append(
            {
                "title": f"Incident: {name}",
                "sub": f"Chronology · {len(mine)} note{'' if len(mine) == 1 else 's'}",
                "rows": mine,
            }
        )
    return groups


def _stills_for(rows: list[dict], folder: Path) -> tuple[dict, int]:
    """A still per note on a camera recording, at the note's moment, up to
    STILLS_MOST; the count left without one. A frame that cannot be cut
    leaves its note without a picture, and the document still prints."""
    from core import media
    from core.recordings import Recording

    wanted = [one for one in rows if one["kind"] in ("line", "moment")]
    recordings = {
        one.pk: one
        for one in Recording.objects.filter(
            pk__in={one["recording_id"] for one in wanted}
        )
    }
    stills: dict = {}
    taken = 0
    left_out = 0
    for one in wanted:
        recording = recordings.get(one["recording_id"])
        source = recording.playback_path() if recording is not None else None
        if source is None or Path(source).suffix != ".mp4":
            continue
        if taken >= STILLS_MOST:
            left_out += 1
            continue
        target = folder / f"still-{len(stills)}"
        target.mkdir(parents=True, exist_ok=True)
        try:
            written = media.grab_frames(
                Path(source), target, [float(one["at"])], height=STILL_HEIGHT
            )
        except Exception:  # noqa: BLE001 - a still is never worth failing the export
            continue
        if written:
            stills[_key(one)] = written[0]
            taken += 1
    return stills, left_out


def _note_block(document, one: dict, still, Inches, Pt, RGBColor) -> None:
    """One note: the still at the left when there is one, the note first in
    the largest type, where it rests, the line quoted, the writer and date."""
    muted = RGBColor(0x66, 0x66, 0x66)
    table = document.add_table(rows=1, cols=2 if still else 1)
    table.autofit = False
    cells = table.rows[0].cells
    if still:
        cells[0].width = Inches(STILL_WIDTH_INCHES + 0.1)
        picture = cells[0].paragraphs[0]
        with contextlib.suppress(Exception):  # the words matter more than the picture
            picture.add_run().add_picture(str(still), width=Inches(STILL_WIDTH_INCHES))
        caption = cells[0].add_paragraph()
        caption_run = caption.add_run(f"still at {one['when']}")
        caption_run.font.size = Pt(8)
        caption_run.font.color.rgb = muted
        body = cells[1]
        body.width = Inches(6.5 - STILL_WIDTH_INCHES - 0.1)
    else:
        cells[0].width = Inches(6.5)
        body = cells[0]
    text = body.paragraphs[0]
    said = text.add_run(one["text"])
    said.font.size = Pt(12.5)
    where = body.add_paragraph()
    when = where.add_run(one["when"])
    when.bold = True
    when.font.size = Pt(9.5)
    if one["kind"] == "event":
        rest = " on the incident clock · on the event"
    elif one["kind"] == "moment":
        rest = " · at a moment where nothing was said"
    elif one.get("who"):
        rest = f" · on {one['who']}'s line"
    else:
        rest = " · on a line"
    tail = where.add_run(rest)
    tail.font.size = Pt(9.5)
    tail.font.color.rgb = muted
    verdict = one.get("verdict") or {}
    if verdict.get("retimed"):
        pill = where.add_run(f"   clock re-read · was {exports.clock(verdict['was'])}")
        pill.font.size = Pt(8)
        pill.font.color.rgb = muted
    if one.get("rests_on"):
        quoted = body.add_paragraph()
        quote = quoted.add_run(f"“{one['rests_on']}”")
        quote.italic = True
        quote.font.size = Pt(9.5)
    by = body.add_paragraph()
    who = by.add_run(f"{one['by']}, {one['on']}" if one.get("by") else one["on"])
    who.font.size = Pt(8.5)
    who.font.color.rgb = muted
    if one.get("from"):
        came = by.add_run(f" · {one['from']}")
        came.font.size = Pt(8.5)
        came.font.color.rgb = muted
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def export_name(case) -> str:
    # The name, the kind and the time (v1.123.0), as the transcripts zip is named.
    when = timezone.localtime().strftime("%Y-%m-%d %H%M")
    return exports.without_clashes(
        set(), f"{exports.safe_name(case.name)} notes {when}.docx"
    )


# Told to Gideon ------------------------------------------------------------------

RULE = (
    "The office's notes are the office's own words: respect them, never "
    "contradict or rewrite them, and never write a note yourself."
)


def any_on(recordings) -> bool:
    """Whether any of these recordings' lines carries a note."""
    found = list(recordings)
    return (
        Segment.objects.filter(transcript__recording__in=found)
        .exclude(note="")
        .exists()
        or MomentNote.objects.filter(recording__in=found).exists()
    )


def recording_block(recording) -> str:
    """The notes on a recording's lines, as the case chat is told them; "" when none."""
    transcript = getattr(recording, "transcript", None)
    if transcript is None:
        return ""
    told = [
        (
            one.start,
            f"[{exports.clock(one.start)}] "
            + (f"{one.note_by.shown_name}: " if one.note_by else "")
            + one.note.replace("\n", " "),
        )
        for one in transcript.segments.exclude(note="")
        .filter(same_as_other_side=False)
        .select_related("note_by")
    ]
    # A note at a moment (v1.99.0), told as that: nothing was said there.
    told += [
        (
            one.at,
            f"[{exports.clock(one.at)}] ({NOTHING_SAID}) "
            + (f"{one.note_by.shown_name}: " if one.note_by else "")
            + one.note.replace("\n", " "),
        )
        for one in moments_of(recording)
    ]
    lines = [words for _, words in sorted(told, key=lambda pair: pair[0])]
    if not lines:
        return ""
    return "The office's notes on this recording:\n" + "\n".join(lines)


def any_on_chronology(incident) -> bool:
    """Whether the Chronology carries a note event, or an Event with a note:
    when it does, the assistant reading it is given the RULE (chapter 11)."""
    from core.chronology import Event

    return (
        Event.objects.filter(incident=incident, proposed=False)
        .filter(
            models.Q(segment__isnull=False)
            | models.Q(moment_note__isnull=False)
            | ~models.Q(note="")
        )
        .exists()
    )
