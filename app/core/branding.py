"""The office's logo: uploaded on the Panel's Appearance page, shown on the
sign-in page and, when the Admin says so, at the head of every Word export.

The image lives in the database, one row, so it rides with the backup like
every other office fact the app keeps, needs no folder of its own on disk,
and is served by the app itself to a page nobody has signed in to yet. A PNG
or JPEG up to 2 MB, recognised by its first bytes rather than its name.
Nothing about it is in the repository: every office uploads its own.
"""

from __future__ import annotations

import hashlib
import io
import uuid

from django.db import models
from django.utils import timezone

from core import audit, settings_store

LARGEST = 2 * 1024 * 1024
KINDS = {
    b"\x89PNG\r\n\x1a\n": "image/png",
    b"\xff\xd8\xff": "image/jpeg",
}


class Branding(models.Model):
    """One row: the logo as uploaded. Absent while no logo has been uploaded."""

    id = models.UUIDField(primary_key=True, default=uuid.uuid4, editable=False)
    logo = models.BinaryField()
    content_type = models.CharField(max_length=40)
    uploaded_at = models.DateTimeField(default=timezone.now)
    uploaded_by = models.ForeignKey(
        "core.User", on_delete=models.SET_NULL, null=True, blank=True
    )

    @property
    def etag(self) -> str:
        return hashlib.sha256(bytes(self.logo)).hexdigest()[:24]


class Refused(ValueError):
    """An upload the app will not keep, with the sentence for the page."""


def kind_of(head: bytes) -> str | None:
    for signature, kind in KINDS.items():
        if head.startswith(signature):
            return kind
    return None


def current() -> Branding | None:
    return Branding.objects.order_by("-uploaded_at").first()


def logo_bytes() -> tuple[bytes, str] | None:
    """The logo and its type, for the sign-in page and the exports; None when none."""
    row = current()
    if row is None:
        return None
    return bytes(row.logo), row.content_type


def upload(data: bytes, *, by, request=None) -> Branding:
    """Keep a new logo in place of any old one, and write the audit row."""
    if not data:
        raise Refused("Choose a file first.")
    if len(data) > LARGEST:
        raise Refused("That file is larger than 2 MB. Make it smaller and try again.")
    kind = kind_of(data[:12])
    if kind is None:
        raise Refused("That file is not a PNG or JPEG image.")
    Branding.objects.all().delete()
    row = Branding.objects.create(logo=data, content_type=kind, uploaded_by=by)
    audit.write(
        audit.Category.ADMIN,
        "Logo uploaded",
        actor=by,
        request=request,
        object_type="branding",
        object_id=row.pk,
        object_label="office logo",
        kind=kind,
        bytes=len(data),
    )
    return row


def remove(*, by, request=None) -> bool:
    row = current()
    if row is None:
        return False
    Branding.objects.all().delete()
    audit.write(
        audit.Category.ADMIN,
        "Logo removed",
        actor=by,
        request=request,
        object_type="branding",
        object_id=row.pk,
        object_label="office logo",
    )
    return True


def office_name() -> str:
    return str(settings_store.get("office_name") or "").strip()


def on_exports() -> bool:
    return bool(settings_store.get("logo_on_exports"))


def stamp(document) -> bool:
    """The logo at the head of a Word export's cover, when the Admin said so.

    Centred, 1.5 inches wide, above the title. Nothing is added when there is
    no logo or the toggle is off, so the export reads as before.
    """
    if not on_exports():
        return False
    found = logo_bytes()
    if found is None:
        return False
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    data, _ = found
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(1.5))
    return True
