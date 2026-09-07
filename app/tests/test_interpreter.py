"""The Interpreter (Phase 3, chapter 3), first slice: text only.

A Session is a Live recording of its own style. Each Turn's audio arrives on
its own, is heard by the service (the fast lane when there is one) and
translated by the engine; the page polls the rows; at Stop the Transcript is
one Segment per Turn, with the words as heard and their translation.
"""

from __future__ import annotations

import json
from types import SimpleNamespace as NS

import pytest
from core import (
    engine,
    interpreter,
    live,
    media,
    settings_store,
    tasks,
    uploads,
    whisperx,
)
from core.audit import Row
from core.cases import Case
from core.interpreter import Turn
from core.jobs import Transcript
from core.models import LoginSession, User
from core.recordings import MediaState
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

PASSWORD = "a-long-enough-password"


@pytest.fixture(autouse=True)
def its_own_disk(tmp_path, settings):
    settings.DATA_DIR = tmp_path
    settings.SCRATCH_DIR = tmp_path / "scratch"
    settings.UPLOADS_DIR = tmp_path / "uploads"
    (tmp_path / "uploads").mkdir()
    return tmp_path


@pytest.fixture(autouse=True)
def room(monkeypatch):
    monkeypatch.setattr(uploads, "free_disk_bytes", lambda: 10**15)


@pytest.fixture
def person(db):
    return User.objects.create_local_admin("intake", PASSWORD)


@pytest.fixture
def on(db, monkeypatch):
    settings_store.set_to("folder_management", True)
    settings_store.set_to("live_recording", True)
    settings_store.set_to("dictation", True)
    settings_store.set_to("interpreter", True)
    monkeypatch.setattr(engine, "configured", lambda: True)


@pytest.fixture
def a_case(person):
    return Case.objects.create(owner=person, name="Ramirez")


@pytest.fixture
def quiet(monkeypatch):
    asked = {"hear": [], "translate": []}
    monkeypatch.setattr(tasks.hear_turn, "defer", lambda **f: asked["hear"].append(f))
    monkeypatch.setattr(
        tasks.translate_turn, "defer", lambda **f: asked["translate"].append(f)
    )
    return asked


def signed_in(client, who):
    client.force_login(who)
    LoginSession.objects.create(user=who, session_key=client.session.session_key)
    return client


# The setting and the page ----------------------------------------------------------


def test_the_interpreter_exists_only_under_its_setting_and_an_engine(
    person, client, monkeypatch
):
    settings_store.set_to("folder_management", True)
    settings_store.set_to("live_recording", True)
    settings_store.set_to("dictation", True)
    monkeypatch.setattr(engine, "configured", lambda: True)
    signed_in(client, person)
    # Off: no page, no card on the New recording page.
    assert client.get(reverse("interpret")).status_code == 404
    assert "Interpreter" not in client.get(reverse("record-new")).content.decode()
    settings_store.set_to("interpreter", True)
    page = client.get(reverse("interpret")).content.decode()
    assert "The visitor's language" in page and "Spanish (Español)" in page
    assert "not an interpreter" in page.lower()
    # Spanish alone is offered, so the shipped notice reads in both languages.
    assert interpreter.SHIPPED_NOTICE_ES in page
    assert (
        'href="/record/interpret"' in client.get(reverse("record-new")).content.decode()
    )
    # No engine: nothing to translate with, so off again.
    monkeypatch.setattr(engine, "configured", lambda: False)
    assert client.get(reverse("interpret")).status_code == 404


def test_the_languages_offered_are_the_admins_ticks(on):
    assert [one["code"] for one in interpreter.languages_offered()] == ["es"]
    settings_store.set_to("interpreter_languages", "es\nvi\nen\nxx\nes")
    offered = interpreter.languages_offered()
    assert [one["code"] for one in offered] == ["es", "vi", "xx"]
    assert offered[1]["name"] == "Vietnamese" and offered[2]["name"] == "xx"


# Starting ----------------------------------------------------


def test_a_session_starts_as_a_live_recording_of_its_own_style(
    on, person, a_case, client
):
    signed_in(client, person)
    answer = client.post(
        reverse("interpret-start"),
        json.dumps({"case": str(a_case.pk), "language": "es"}),
        "application/json",
    )
    assert answer.status_code == 200, answer.content
    said = answer.json()
    recording = live.start.__module__ and interpreter.Turn.objects.none()  # noqa: F841
    from core.recordings import Recording

    recording = Recording.objects.get(pk=said["id"])
    assert recording.is_live and interpreter.is_session(recording)
    assert recording.live["style"] == "interpreter"
    assert recording.live["visitor_language"] == "es"
    assert recording.recording_type == "Interview" and not recording.diarize
    assert recording.case_id == a_case.pk and not recording.is_dictation
    assert said["language_name"] == "Spanish" and said["after"] == f"/case/{a_case.pk}"
    assert said["notice"] == [interpreter.SHIPPED_NOTICE, interpreter.SHIPPED_NOTICE_ES]
    assert recording.title.startswith("Interpreter, Spanish")
    row = Row.objects.get(event="Live recording started")
    assert row.details["style"] == "interpreter"

    # Without a case it is kept on its own, as a dictation is; a language the
    # office does not offer is refused.
    answer = client.post(
        reverse("interpret-start"), json.dumps({"language": ""}), "application/json"
    )
    alone = Recording.objects.get(pk=answer.json()["id"])
    assert alone.is_dictation and alone.case is None
    assert alone.live["visitor_language"] == "" and answer.json()["after"] == "/"
    answer = client.post(
        reverse("interpret-start"), json.dumps({"language": "vi"}), "application/json"
    )
    assert (
        answer.status_code == 400
        and "not one the office offers" in answer.json()["why"]
    )


# Turns ----------------------------------------------------


def a_session(person, case=None, language="es"):
    return interpreter.start(person, case, language=language, title="", browser="")


def test_a_spoken_turn_is_heard_on_the_lane_and_translated_by_the_engine(
    on, person, quiet, monkeypatch, settings
):
    session = a_session(person)
    monkeypatch.setattr(
        media, "make_asr_audio", lambda src, dst, track, **k: dst.write_bytes(b"w")
    )
    settings.WHISPERX_FAST_URL = "http://whisperx-fast:8000"
    monkeypatch.setattr(whisperx, "is_alive", lambda lane="": True)
    submitted = []
    monkeypatch.setattr(
        whisperx,
        "submit",
        lambda audio, request, lane="": (
            submitted.append((request, lane))
            or whisperx.Submitted(id="t1", position=0, audio_minutes_ahead=0.0)
        ),
    )
    monkeypatch.setattr(
        whisperx, "jobs", lambda lane="": [{"id": "t1", "state": "done"}]
    )
    monkeypatch.setattr(
        whisperx,
        "result",
        lambda job_id, lane="": {
            "segments": [{"text": " ¿Dónde está la oficina?"}],
            "language": {"detected": "es"},
        },
    )
    monkeypatch.setattr(whisperx, "delete", lambda job_id, lane="": None)

    turn = interpreter.turn_arrived(
        session, start_at=3.0, end_at=6.5, audio=b"\x1aE\xdf\xa3" + b"x" * 40
    )
    assert turn.number == 1 and turn.state == Turn.HEARING and turn.audio_path.exists()
    assert quiet["hear"] == [{"turn_id": str(turn.pk)}]

    interpreter.hear(turn)
    turn.refresh_from_db()
    request, lane = submitted[0]
    # Nobody held a button, so the language is left to the service, and the
    # Turn goes ahead of everything at the front of the fast lane.
    assert request["language"] == "" and request["priority"] == 100
    assert request["task"] == "transcribe" and lane == whisperx.FAST
    assert turn.heard == "¿Dónde está la oficina?" and turn.language == "es"
    assert turn.side == "visitor" and turn.state == Turn.TRANSLATING
    assert quiet["translate"] == [{"turn_id": str(turn.pk)}]

    # The engine translates into the other side's language.
    asked = []
    monkeypatch.setattr(
        engine,
        "complete",
        lambda messages, **k: (
            asked.append((messages, k))
            or {"text": " Where is the office? ", "finish_reason": "stop"}
        ),
    )
    interpreter.translate(turn)
    turn.refresh_from_db()
    assert turn.translation == "Where is the office?" and turn.state == Turn.DONE
    messages, options = asked[0]
    assert "into English" in messages[0]["content"]
    assert messages[1]["content"] == "¿Dónde está la oficina?"
    assert options["thinking"] is False and options["temperature"] == 0.2

    # A Staff Turn goes the other way, into the Visitor's language.
    staff = interpreter.turn_arrived(
        session, start_at=8.0, end_at=8.0, side="staff", typed="Please wait a moment."
    )
    assert staff.typed and staff.language == "en" and staff.state == Turn.TRANSLATING
    monkeypatch.setattr(
        engine,
        "complete",
        lambda messages, **k: (
            asked.append((messages, k))
            or {"text": "Espere un momento, por favor.", "finish_reason": "stop"}
        ),
    )
    interpreter.translate(staff)
    staff.refresh_from_db()
    assert staff.translation == "Espere un momento, por favor."
    assert "into Spanish" in asked[-1][0][0]["content"]


def test_the_language_is_settled_by_the_first_visitor_turn_when_left_to_be_heard(
    on, person, quiet, monkeypatch, settings
):
    session = a_session(person, language="")
    monkeypatch.setattr(
        media, "make_asr_audio", lambda src, dst, track, **k: dst.write_bytes(b"w")
    )
    settings.WHISPERX_FAST_URL = ""
    monkeypatch.setattr(
        whisperx,
        "submit",
        lambda audio, request, lane="": whisperx.Submitted(
            id="t2", position=0, audio_minutes_ahead=0.0
        ),
    )
    monkeypatch.setattr(
        whisperx, "jobs", lambda lane="": [{"id": "t2", "state": "done"}]
    )
    monkeypatch.setattr(whisperx, "delete", lambda job_id, lane="": None)
    # Staff speak first: English, and nothing to translate into yet.
    monkeypatch.setattr(
        whisperx,
        "result",
        lambda job_id, lane="": {
            "segments": [{"text": "Good morning."}],
            "language": {"detected": "en"},
        },
    )
    first = interpreter.turn_arrived(session, start_at=0, end_at=2, audio=b"x" * 10)
    interpreter.hear(first)
    first.refresh_from_db()
    assert first.side == "staff" and first.state == Turn.TRANSLATING
    interpreter.translate(first)
    first.refresh_from_db()
    assert first.state == Turn.DONE and first.translation == ""
    session.refresh_from_db()
    assert interpreter.visitor_language(session) == ""
    # Then the visitor: the language is settled, and a row says so.
    monkeypatch.setattr(
        whisperx,
        "result",
        lambda job_id, lane="": {
            "segments": [{"text": "Buenos días."}],
            "language": {"detected": "es"},
        },
    )
    second = interpreter.turn_arrived(session, start_at=3, end_at=5, audio=b"x" * 10)
    interpreter.hear(second)
    session.refresh_from_db()
    assert interpreter.visitor_language(session) == "es"
    assert session.live["language_settled_by"] == "heard"
    settled = Row.objects.get(event="Language settled")
    assert settled.details == {"language": "es", "at_turn": 2}
    # The rows the page polls carry both, and the language.
    rows = interpreter.rows(session, since=0)
    assert [one["side"] for one in rows] == ["staff", "visitor"]
    assert rows[1]["state"] == "translating" and rows[1]["heard"] == "Buenos días."


def test_a_filler_whisper_made_up_is_nothing_heard_and_a_service_away_is_asked_again(
    on, person, quiet, monkeypatch, settings
):
    session = a_session(person)
    monkeypatch.setattr(
        media, "make_asr_audio", lambda src, dst, track, **k: dst.write_bytes(b"w")
    )
    settings.WHISPERX_FAST_URL = ""

    # The service is away: the Turn stays "hearing" and is asked again later.
    def away(audio, request, lane=""):
        raise whisperx.ServiceError("down", "service_unreachable")

    monkeypatch.setattr(whisperx, "submit", away)
    again = []
    monkeypatch.setattr(
        tasks.hear_turn,
        "configure",
        lambda **c: NS(defer=lambda **f: again.append({**f, **c})),
    )
    turn = interpreter.turn_arrived(session, start_at=0, end_at=2, audio=b"x" * 10)
    interpreter.hear(turn, attempt=1)
    turn.refresh_from_db()
    assert turn.state == Turn.HEARING
    assert again == [
        {
            "turn_id": str(turn.pk),
            "attempt": 2,
            "schedule_in": {"seconds": interpreter.HEAR_RETRY_SECONDS},
        }
    ]
    # The last try fails plainly.
    interpreter.hear(turn, attempt=interpreter.HEAR_TRIES)
    turn.refresh_from_db()
    assert turn.state == Turn.FAILED and turn.failure == "service_unreachable"

    # "Gracias." from a clip of nothing is nothing heard, not a Turn to translate.
    monkeypatch.setattr(
        whisperx,
        "submit",
        lambda audio, request, lane="": whisperx.Submitted(
            id="t3", position=0, audio_minutes_ahead=0.0
        ),
    )
    monkeypatch.setattr(
        whisperx, "jobs", lambda lane="": [{"id": "t3", "state": "done"}]
    )
    monkeypatch.setattr(whisperx, "delete", lambda job_id, lane="": None)
    monkeypatch.setattr(
        whisperx,
        "result",
        lambda job_id, lane="": {
            "segments": [{"text": " Gracias. "}],
            "language": {"detected": "es"},
        },
    )
    quiet["translate"].clear()
    filler = interpreter.turn_arrived(
        session, start_at=3, end_at=4, side="visitor", audio=b"x" * 10
    )
    interpreter.hear(filler)
    filler.refresh_from_db()
    assert filler.state == Turn.DONE and filler.heard == "" and not quiet["translate"]
    assert interpreter.is_filler("Thank you.") and interpreter.is_filler(
        "¡Muchas gracias!"
    )
    assert not interpreter.is_filler("Gracias por venir hoy.")


def test_the_page_posts_turns_and_polls_them(on, person, client, quiet):
    session = a_session(person)
    signed_in(client, person)
    answer = client.post(
        reverse("interpret-turn", args=[session.pk]),
        {
            "audio": SimpleUploadedFile("turn.webm", b"\x1aE\xdf\xa3" + b"x" * 20),
            "start": "1.5",
            "end": "4.0",
            "side": "visitor",
        },
    )
    assert answer.status_code == 200 and answer.json()["number"] == 1
    turn = Turn.objects.get(recording=session)
    assert turn.side == "visitor" and turn.start == 1.5 and turn.end == 4.0
    assert quiet["hear"] == [{"turn_id": str(turn.pk)}]
    # Typed, as JSON.
    answer = client.post(
        reverse("interpret-turn", args=[session.pk]),
        json.dumps(
            {"typed": "My name is Ana.", "side": "visitor", "start": 6, "end": 6}
        ),
        "application/json",
    )
    assert answer.status_code == 200 and answer.json()["number"] == 2
    typed = Turn.objects.get(recording=session, number=2)
    assert typed.typed and typed.language == "es" and typed.state == Turn.TRANSLATING
    # The poll.
    said = client.get(reverse("interpret-turns", args=[session.pk]) + "?since=0").json()
    assert [one["number"] for one in said["turns"]] == [1, 2]
    assert said["language_name"] == "Spanish"
    # Nobody else's, and nothing after the end.
    other = User.objects.create_local_admin("other", PASSWORD)
    signed_in(client, other)
    assert client.get(reverse("interpret-turns", args=[session.pk])).status_code == 404
    signed_in(client, person)
    live.ended(session, how="stop", seconds=30)
    answer = client.post(
        reverse("interpret-turn", args=[session.pk]),
        json.dumps({"typed": "late", "side": "staff"}),
        "application/json",
    )
    assert answer.status_code == 400


# At Stop ----------------------------------------------------


def test_the_transcript_is_one_segment_per_turn_in_both_languages(
    on, person, quiet, monkeypatch
):
    session = a_session(person)
    for number, (side, heard, translation, start) in enumerate(
        [
            ("visitor", "¿Dónde está la oficina?", "Where is the office?", 3.0),
            ("staff", "Down the hall.", "Al final del pasillo.", 8.0),
        ],
        start=1,
    ):
        Turn.objects.create(
            recording=session,
            number=number,
            side=side,
            start=start,
            end=start + 3,
            language="es" if side == "visitor" else "en",
            heard=heard,
            translation=translation,
            state=Turn.DONE,
        )
    live.ended(session, how="stop", seconds=12)
    session.refresh_from_db()
    session.media_state = MediaState.READY
    session.duration_seconds = 12.0
    session.save()
    from core.recordings import Side

    Side.objects.create(recording=session, number=1, kind=Side.WHOLE)

    # prepare_recording hands a Session to the Interpreter, never to the queue.
    made = []
    monkeypatch.setattr(tasks.hand_over_job, "defer", lambda **f: made.append(f))
    monkeypatch.setattr(tasks.make_playback_copy, "defer", lambda **f: None)
    from core import pipeline

    monkeypatch.setattr(pipeline, "prepare", lambda recording: recording)
    tasks.prepare_recording.func(str(session.pk))
    assert not made
    transcript = Transcript.objects.get(recording=session)
    rows = list(
        transcript.segments.order_by("start").values_list(
            "speaker", "text", "translation", "language"
        )
    )
    assert rows == [
        ("Visitor", "¿Dónde está la oficina?", "Where is the office?", "es"),
        ("Staff", "Down the hall.", "Al final del pasillo.", "en"),
    ]
    assert (
        transcript.language == "es"
        and transcript.provenance["interpreter"]["turns"] == 2
    )
    assert ("Turns", "2, each transcribed and translated as it ended") in (
        interpreter.provenance_rows(session)
    )
    assert "Interpreted" in dict(interpreter.provenance_rows(session))


def test_the_viewer_and_the_export_carry_the_translation(on, person, client, quiet):
    session = a_session(person)
    Turn.objects.create(
        recording=session,
        number=1,
        side="visitor",
        start=1.0,
        end=3.0,
        language="es",
        heard="Hola.",
        translation="Hello.",
        state=Turn.DONE,
    )
    session.media_state = MediaState.READY
    session.duration_seconds = 4.0
    session.save()
    interpreter.finish(session)
    signed_in(client, person)
    said = client.get(reverse("segments", args=[session.pk])).json()
    assert said["segments"][0]["translation"] == "Hello."
    assert said["segments"][0]["language"] == "es"
    from core import exports

    body = exports.word(session, person.username)
    assert body[:2] == b"PK"
    from io import BytesIO

    from docx import Document

    document = Document(BytesIO(body))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "VISITOR: Hola." in text and "Hello." in text
    # The Session's rows are on the processing record, a table at the end.
    cells = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Machine translation, not an interpreter" in cells
