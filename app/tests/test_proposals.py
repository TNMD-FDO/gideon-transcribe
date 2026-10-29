"""v1.168.0: a proposed event is a moment on a recording, and accepting it
writes a note there (Phase 10 chapter 4, the first release).

Checked here: a Proposal accepted where nothing is said is a note at the
moment, with its origin, counted on the dashboard until decided and listed
with its origin on the Notes tab, the recording page and Download notes; the
chronology's note event carries the why and the AI notice prints; Dismiss
all puts both kinds away; a cleared note forgets its origin; an old row on a
camera no longer synced is refused plainly; a recording's comparison can make
a note, and the memo's lines say where it came from.
"""

from __future__ import annotations

import io

import pytest
from core import chronology, dashboard, incident_assistant, incidents, notes, proposals
from core.audit import Row
from core.cases import Case
from core.chronology import Event
from core.jobs import Segment
from core.models import User
from core.notes import MomentNote
from core.proposals import Proposal
from core.recordings import Batch, MediaState, Recording
from tests.helpers import signed_in

PASSWORD = "a-long-enough-password"


@pytest.fixture(autouse=True)
def switched_on(db):
    from core import settings_store

    settings_store.set_to("folder_management", True)
    settings_store.set_to("incidents", True)
    settings_store.set_to("assistant_available", True)


@pytest.fixture
def person(db):
    return User.objects.create_local_admin("asker", PASSWORD)


@pytest.fixture
def a_case(person):
    return Case.objects.create(owner=person, name="Traffic stop")


def video(person, case, title, lines=()):
    from core.jobs import Transcript

    recording = Recording.objects.create(
        batch=Batch.objects.create(user=person),
        user=person,
        case=case,
        title=title,
        original_filename=f"{title}.mp4",
        media_state=MediaState.READY,
        duration_seconds=600.0,
        playback_ready=True,
        probe={},
    )
    recording.folder.mkdir(parents=True, exist_ok=True)
    (recording.folder / "playback.mp4").write_bytes(b"not really media")
    transcript = Transcript.objects.create(recording=recording, language="en")
    for start, text in lines:
        Segment.objects.create(
            transcript=transcript,
            start=start,
            end=start + 4,
            text=text,
            speaker="Speaker 1",
            speaker_label="SPEAKER_1",
        )
    return recording


def a_proposal(case, recording, incident=None, **fields):
    given = {
        "at": 120.0,
        "text": "A weapon was mentioned.",
        "why": "A weapon changes the stop.",
        "rests_on": "I got gun",
        "source": "assistant",
    }
    given.update(fields)
    return Proposal.objects.create(
        case=case, recording=recording, incident=incident, **given
    )


def test_accept_at_a_moment_writes_a_note_with_its_origin_everywhere(
    person, a_case, client
):
    cam = video(person, a_case, "BWC2-1", lines=((10.0, "Step out."),))
    made = incidents.make(a_case, "Stop", [cam], by=person)
    camera = made.cameras.get()
    incidents.place_by_hand(camera, 0.0, by=person)
    one = a_proposal(a_case, cam, made)
    # Counted while waiting, on the incident's strip and the dashboard.
    assert incident_assistant.proposals_json(made)["pending"] == 1
    assert dashboard.incident_counts(a_case)[made.pk]["proposals"] == 1
    rows = chronology.events_json(made)
    assert [(row["proposed"], row["at"]) for row in rows] == [(True, 120.0)]
    # Nothing is being said at 120 s: the note sits at the moment.
    incident_assistant.accept(one, by=person)
    one.refresh_from_db()
    note = MomentNote.objects.get()
    assert one.state == "accepted" and one.note_moment == note
    assert note.note == "A weapon was mentioned." and note.note_by == person
    assert note.note_from["words"] == "from a proposed event"
    assert note.note_from["why"] == "A weapon changes the stop."
    assert incident_assistant.proposals_json(made)["pending"] == 0
    assert dashboard.incident_counts(a_case).get(made.pk, {}).get("proposals", 0) == 0
    # The chronology's note event, with the why and the origin on its pill.
    event = Event.objects.get(source="note")
    assert event.moment_note == note and event.at == 120.0
    assert event.why == "A weapon changes the stop." and event.rests_on == "I got gun"
    row = chronology.events_json(made)[0]
    assert not row["proposed"] and row["at_a_moment"]
    assert (
        row["source_words"]
        == f"Note by {person.shown_name}, BWC2-1, from a proposed event"
    )
    assert row["why"] == "A weapon changes the stop."
    assert chronology._rows(made)[0]["assistant"] is True
    # The exports print the AI notice for it.
    from docx import Document as Docx

    exported = Docx(io.BytesIO(chronology.word(made, None, "asker")))
    text = "\n".join(p.text for p in exported.paragraphs) + "\n".join(
        cell.text
        for table in exported.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Why it matters: A weapon changes the stop." in text
    assert "AI" in text or "assistant" in text.lower()
    # The pages: the recording page's JSON, the Notes tab, Download notes.
    assert notes.moment_json(note)["note_from"] == "from a proposed event"
    listed = notes.of_case(a_case)["rows"]
    assert (
        listed[0]["from"] == "from a proposed event" and listed[0]["kind"] == "moment"
    )
    signed_in(client, person)
    page = client.get(f"/case/{a_case.pk}?tab=notes").content.decode()
    assert "from a proposed event" in page
    words = "\n".join(
        cell.text
        for table in Docx(io.BytesIO(notes.word(a_case, "asker"))).tables
        for row in table.rows
        for cell in row.cells
    )
    assert "from a proposed event" in words
    # The audit rows: the proposal's and the note's, never the sentence.
    accepted = Row.objects.get(event="proposal accepted")
    assert (
        accepted.details["source"] == "assistant"
        and accepted.details["how"] == "moment"
    )
    assert accepted.details["incident"] == "Stop"
    assert Row.objects.filter(event="Note added", details__where="proposal").exists()
    for row in Row.objects.all():
        assert "weapon" not in str(row.details).lower()
    # The second accepted at the same moment joins the note; a cleared note
    # forgets its origin.
    incident_assistant.accept(a_proposal(a_case, cam, made, text="Second."), by=person)
    note.refresh_from_db()
    assert note.note == "A weapon was mentioned.\n\nSecond."
    assert note.note_from["words"] == "from a proposed event"
    notes.set_moment_note(note, "", by=person)
    assert not MomentNote.objects.exists() and not Event.objects.exists()


def test_accept_on_a_line_the_proposal_rested_on_and_the_pages_say_so(
    person, a_case, client
):
    cam = video(person, a_case, "BWC2-1", lines=((10.0, "Step out."), (30.0, "Gun.")))
    made = incidents.make(a_case, "Stop", [cam], by=person)
    incidents.place_by_hand(made.cameras.get(), 0.0, by=person)
    one = a_proposal(a_case, cam, made, at=31.0, source="watch", rests_on="Gun.")
    signed_in(client, person)
    got = client.post(
        f"{made.url()}/act", {"action": "event_accept", "event": str(one.pk)}
    ).json()
    assert got["ok"]
    line = Segment.objects.get(start=30.0)
    assert line.note == "A weapon was mentioned." and line.note_by == person
    assert line.note_from["words"] == "from a proposed event (a watch phrase)"
    assert (
        notes.line_json(line)["note_from"] == "from a proposed event (a watch phrase)"
    )
    row = next(r for r in got["state"]["events"] if r["source"] == "note")
    assert row["source_words"].endswith(
        "BWC2-1, from a proposed event (a watch phrase)"
    )
    # The app's own search is not the assistant: no AI notice for it.
    assert chronology._rows(made)[0]["assistant"] is False
    # Dismiss all puts a Proposal row and an old Event row away together.
    a_proposal(a_case, cam, made, at=200.0)
    Event.objects.create(
        incident=made,
        at=100.0,
        text="Old",
        source="assistant",
        camera=made.cameras.get(),
        proposed=True,
        rests_on="a line",
    )
    got = client.post(f"{made.url()}/act", {"action": "event_dismiss_all"}).json()
    assert got["said"] == "2 proposed events dismissed."
    assert Proposal.objects.filter(state="dismissed").count() == 1
    assert Event.objects.filter(proposed=True, dismissed=True).count() == 1
    assert Row.objects.get(event="events dismissed").details["count"] == 2


def test_an_old_row_on_a_camera_no_longer_synced_is_refused_plainly(
    person, a_case, client
):
    cam = video(person, a_case, "BWC2-1", lines=((10.0, "Step out."),))
    made = incidents.make(a_case, "Stop", [cam], by=person)
    camera = made.cameras.get()
    old = Event.objects.create(
        incident=made,
        at=20.0,
        text="Old",
        source="assistant",
        camera=camera,
        proposed=True,
        rests_on="a line",
    )
    assert not camera.is_synced()
    signed_in(client, person)
    told = client.post(
        f"{made.url()}/act", {"action": "event_accept", "event": str(old.pk)}
    )
    assert told.status_code == 400 and "no longer synced" in told.json()["error"]
    assert Event.objects.filter(pk=old.pk, proposed=True).exists()
    # Accept all skips it and says so by the count.
    got = client.post(f"{made.url()}/act", {"action": "event_accept_all"}).json()
    assert got["said"] == "0 notes written."
    # A stray id is a 404, as before.
    assert (
        client.post(
            f"{made.url()}/act", {"action": "event_accept", "event": "nope"}
        ).status_code
        == 404
    )


def test_the_memo_lines_and_the_recording_comparison_make_a_note(person, a_case):
    cam = video(person, a_case, "BWC2-1", lines=((60.0, "That bag is not mine."),))
    made = incidents.make(a_case, "Stop", [cam], by=person)
    incidents.place_by_hand(made.cameras.get(), 0.0, by=person)
    # A note from a finding, written through the same door the comparison uses.
    origin = {
        "kind": "finding",
        "source": "report",
        "why": "Differs: heard on camera.",
        "rests_on": "[Report, page 1, paragraph 3] A small bag was found.",
        "words": proposals.origin_words("report", page=1, n=3),
    }
    kind, home = proposals.write_note(
        cam, 61.0, "The driver said the bag was not his", origin=origin, by=person
    )
    assert kind == "line" and home.start == 60.0
    assert home.note_from["words"] == "from the report, page 1, paragraph 3"
    lines, _ = incident_assistant._chronology_lines(made)
    assert len(lines) == 1
    assert (
        "accepted from the report, which says: [Report, page 1, paragraph 3]"
        in (lines[0])
    )
    assert chronology._rows(made)[0]["assistant"] is True
    # The same note, accepted from a proposal instead, says that.
    notes.set_note(home, "", by=person)
    incident_assistant.accept(
        a_proposal(a_case, cam, made, at=62.0, rests_on="not mine"), by=person
    )
    lines, _ = incident_assistant._chronology_lines(made)
    assert "accepted from a proposed event" in lines[0]
