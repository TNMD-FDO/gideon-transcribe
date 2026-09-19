"""v1.69.0: Notes (Phase 8, chapter 2).

The rules checked here: a note on a line is written, changed and removed by
whoever may correct the transcript, kept with the writer and the date, and
audited without a word of it; a transcript being replaced takes no note; the
plain exports never carry a note and the with-notes exports do, offered only
while a line has one; the Notes tab counts and lists every note in the case,
on lines and on events, newest first, with Download notes; Search's Notes
kind and Find read the notes on lines; Gideon on the case page and on the
incident page is told them as the office's own words; processing again
carries each note to the new line at its moment.
"""

from __future__ import annotations

import io
import json

import pytest
from core import (
    case_chat,
    case_search,
    chronology,
    incident_chat,
    incidents,
    notes,
    settings_store,
    tasks,
    viewer,
)
from core.audit import Row
from core.case_chat import CaseChat, CaseChatTurn
from core.cases import Case
from core.jobs import Segment, Transcript
from core.models import LoginSession, User
from core.recordings import Batch, MediaState, Recording
from tests.test_prepare import engine_answering

PASSWORD = "a-long-enough-password"


@pytest.fixture(autouse=True)
def its_own_disk(tmp_path, settings):
    settings.DATA_DIR = tmp_path
    settings.SCRATCH_DIR = tmp_path / "scratch"
    settings.UPLOADS_DIR = tmp_path / "uploads"


@pytest.fixture(autouse=True)
def switched_on(db, monkeypatch):
    settings_store.set_to("folder_management", True)
    settings_store.set_to("incidents", True)
    settings_store.set_to("assistant_available", True)
    settings_store.set_to("chat_available", True)
    settings_store.set_to("chat_across_cases", True)
    monkeypatch.setattr(tasks.answer_incident_turn, "defer", lambda **f: None)
    monkeypatch.setattr(tasks.answer_case_turn, "defer", lambda **f: None)


@pytest.fixture
def person(db):
    return User.objects.create_local_admin("asker", PASSWORD)


@pytest.fixture
def stranger(db):
    return User.objects.create(username="nobody", display_name="nobody")


@pytest.fixture
def a_case(person):
    return Case.objects.create(owner=person, name="Traffic stop")


def signed_in(client, who):
    client.force_login(who)
    LoginSession.objects.create(user=who, session_key=client.session.session_key)
    return client


def stamp(time: str, camera: str):
    return {
        "date": "06/07/2025",
        "time": time,
        "camera": camera,
        "at": 2.0,
        "checked": True,
    }


LINES = (
    (0.0, "Officer Hale", "This is Officer Hale."),
    (30.0, "Officer Hale", "I got, I got gun, I got gun."),
    (60.0, "Speaker 2", "That's not mine."),
)


def video(person, case, title, *, stamp=None, lines=LINES):
    recording = Recording.objects.create(
        batch=Batch.objects.create(user=person),
        user=person,
        case=case,
        title=title,
        original_filename=f"{title}.mp4",
        media_state=MediaState.READY,
        duration_seconds=600.0,
        playback_ready=True,
        stamp=stamp,
        probe={},
    )
    recording.folder.mkdir(parents=True, exist_ok=True)
    (recording.folder / "playback.mp4").write_bytes(b"not really media")
    if lines:
        transcript = Transcript.objects.create(recording=recording, language="en")
        for start, speaker, text in lines:
            Segment.objects.create(
                transcript=transcript,
                start=start,
                end=start + 4,
                text=text,
                speaker=speaker,
                speaker_label=speaker.upper().replace(" ", "_"),
            )
    return recording


def gun_line(recording):
    return recording.transcript.segments.get(start=30.0)


def post_note(client, recording, segment, text):
    return client.post(
        f"/recording/{recording.pk}/segment/{segment.pk}/note",
        json.dumps({"note": text}),
        content_type="application/json",
    )


def docx_text(body: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(body))
    return "\n".join(paragraph.text for paragraph in document.paragraphs) + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )


# A note on a line ---------------------------------------------------------------------


def test_a_note_is_written_changed_and_removed_without_a_word_logged(
    person, a_case, client
):
    recording = video(person, a_case, "first")
    line = gun_line(recording)
    signed_in(client, person)
    told = post_note(
        client, recording, line, "  Check the  body-worn footage here.\n\n"
    )
    assert told.status_code == 200
    got = told.json()
    assert got["saved"] and got["note"] == "Check the body-worn footage here."
    assert got["note_by"] == "asker" and got["note_on"]
    line.refresh_from_db()
    assert line.note_by == person and line.note_changed is not None
    # The page reads it with the line.
    segments = client.get(f"/recording/{recording.pk}/segments").json()["segments"]
    noted = [one for one in segments if one["note"]]
    assert len(noted) == 1 and noted[0]["note_by"] == "asker"
    # The Export menu offers the with-notes entries only now.
    page = client.get(f"/recording/{recording.pk}").content.decode()
    assert "Export to Word, with notes" in page and "?notes=1" in page
    # Changed, then removed; the writer is whoever last wrote it.
    assert (
        post_note(client, recording, line, "Page 4 of the report.").status_code == 200
    )
    got = post_note(client, recording, line, "").json()
    assert got["note"] == "" and got["note_by"] == "" and got["note_on"] == ""
    line.refresh_from_db()
    assert line.note_by is None and line.note_changed is None
    page = client.get(f"/recording/{recording.pk}").content.decode()
    assert "with notes" not in page
    rows = list(Row.objects.filter(category="edits").order_by("at"))
    assert [row.event for row in rows] == ["Note added", "Note changed", "Note removed"]
    for row in rows:
        assert row.object_label == "30.0-34.0"
        blob = json.dumps(row.details).lower()
        assert "footage" not in blob and "page 4" not in blob
    # The same words again write nothing.
    post_note(client, recording, line, "Again.")
    post_note(client, recording, line, "Again.")
    assert Row.objects.filter(category="edits").count() == 4


def test_only_who_may_correct_writes_a_note_and_not_while_replacing(
    person, stranger, a_case, client, monkeypatch
):
    recording = video(person, a_case, "first")
    line = gun_line(recording)
    signed_in(client, stranger)
    assert post_note(client, recording, line, "Mine?").status_code == 404
    client.logout()
    signed_in(client, person)
    monkeypatch.setattr(viewer, "being_replaced", lambda recording: True)
    told = post_note(client, recording, line, "Now?")
    assert told.status_code == 409 and "being replaced" in told.json()["error"]
    # The hidden second copy of a line takes no note.
    monkeypatch.setattr(viewer, "being_replaced", lambda recording: False)
    Segment.objects.filter(pk=line.pk).update(same_as_other_side=True)
    assert post_note(client, recording, line, "Now?").status_code == 404
    # Over the length, the note is cut at the chapter's 2,000 characters.
    Segment.objects.filter(pk=line.pk).update(same_as_other_side=False)
    got = post_note(client, recording, line, "x" * 2500).json()
    assert len(got["note"]) == 2000


# Where the notes print ----------------------------------------------------------------


def test_the_plain_exports_never_carry_a_note_and_the_with_notes_ones_do(
    person, a_case, client
):
    recording = video(person, a_case, "first")
    signed_in(client, person)
    post_note(client, recording, gun_line(recording), "Check the footage.")
    plain = client.get(f"/recording/{recording.pk}/export/text").content.decode()
    assert "Check the footage." not in plain and "Note" not in plain
    captions = client.get(f"/recording/{recording.pk}/export/captions").content.decode()
    assert "Check the footage." not in captions
    with_notes = client.get(
        f"/recording/{recording.pk}/export/text?notes=1"
    ).content.decode()
    assert "\r\n    Note (asker, " in with_notes
    assert "): Check the footage." in with_notes
    assert "Lines beginning Note are the office's own notes" in with_notes
    # The note sits on the line after its line.
    lines = with_notes.split("\r\n")
    at = next(n for n, one in enumerate(lines) if "I got gun" in one)
    assert lines[at + 1].startswith("    Note (asker")
    word = client.get(f"/recording/{recording.pk}/export/word")
    assert "Check the footage." not in docx_text(word.content)
    word = client.get(f"/recording/{recording.pk}/export/word?notes=1")
    text = docx_text(word.content)
    assert "Check the footage." in text and "With the office's notes: 1." in text
    assert "with%20notes" in word["Content-Disposition"]
    kinds = list(
        Row.objects.filter(category="exports")
        .order_by("at")
        .values_list("details__kind", flat=True)
    )
    assert kinds == [
        "transcript text",
        "captions",
        "transcript text with notes",
        "transcript word",
        "transcript word with notes",
    ]


# The Notes tab ------------------------------------------------------------------------


def test_the_notes_tab_lists_every_note_in_the_case_newest_first(
    person, a_case, client
):
    first = video(person, a_case, "first", stamp=stamp("21:56:19", "BWC2-1"))
    second = video(person, a_case, "second", stamp=stamp("22:01:00", "BWC2-2"))
    incident = incidents.make(a_case, "Stop", [first, second], by=person)
    signed_in(client, person)
    page = client.get(f"/case/{a_case.pk}?tab=notes").content.decode()
    assert 'Notes <span class="count">0</span>' in page and "None yet" in page
    assert "Download notes" not in page
    # A note on an event, then a note on a line: the line's is newer.
    cams = {one.camera_id(): one for one in incident.cameras.all()}
    chronology.add(
        incident,
        {
            "at": "30",
            "text": "Gun found",
            "camera": str(cams["BWC2-1"].pk),
            "note": "Page 4 of the report.",
        },
        by=person,
    )
    post_note(client, first, gun_line(first), "Check the footage.")
    page = client.get(f"/case/{a_case.pk}?tab=notes").content.decode()
    assert 'Notes <span class="count">2</span>' in page
    assert "2 notes in this case" in page and "1 on line of transcripts" in page
    assert page.index("Check the footage.") < page.index("Page 4 of the report.")
    assert "On lines (1)" in page and "On events (1)" in page
    assert f"/recording/{first.pk}?t=30.0&amp;note=1" in page
    assert f"{incident.url()}?t=30.00" in page and "on the event:" in page
    assert "all cameras" in page and "Officer Hale:" in page
    only = client.get(f"/case/{a_case.pk}?tab=notes&kind=events").content.decode()
    assert "Page 4" in only and "Check the footage." not in only
    # Download notes: both, for the office's own reading; one row, no words.
    got = client.get(f"/case/{a_case.pk}/notes.docx")
    assert got.status_code == 200 and "notes.docx" in got["Content-Disposition"]
    text = docx_text(got.content)
    assert "Notes: Traffic stop" in text and "2 notes" in text
    assert "Check the footage." in text and "Page 4 of the report." in text
    row = Row.objects.get(event="notes exported")
    assert row.object_label == "Traffic stop"
    assert "footage" not in json.dumps(row.details).lower()
    # The case page's other tabs count the same.
    assert (
        'Notes <span class="count">2</span>'
        in client.get(f"/case/{a_case.pk}").content.decode()
    )


# Search and Find ----------------------------------------------------------------------


def test_search_and_find_read_the_notes_on_lines(person, a_case, client):
    first = video(person, a_case, "first", stamp=stamp("21:56:19", "BWC2-1"))
    second = video(person, a_case, "second", stamp=stamp("22:01:00", "BWC2-2"))
    incident = incidents.make(a_case, "Stop", [first, second], by=person)
    signed_in(client, person)
    post_note(client, first, gun_line(first), "Check the footage for the gun.")
    got = case_search.search(a_case, "footage", "notes")
    assert got["total"] == 1 and got["groups"][0]["head"] == "first"
    hit = got["groups"][0]["hits"][0]
    assert hit["who"] == "Note, asker" and "<mark>footage</mark>" in str(hit["under"])
    assert hit["url"].endswith("?t=30.0&note=1") and hit["all_cameras"]
    # The Notes filter counts lines' notes with events' notes.
    counts = {
        one["key"]: one["count"] for one in case_search.search(a_case, "gun")["kinds"]
    }
    assert counts["notes"] == 1 and counts["words"] == 2
    found = case_search.find(incident, "footage")
    assert [one["kind"] for one in found["hits"]] == ["note"]
    assert found["hits"][0]["who"] == "Note, asker"
    assert found["hits"][0]["camera_id"] == "BWC2-1" and found["hits"][0]["at"] == 30.0


# Told to Gideon -----------------------------------------------------------------------


def test_gideon_is_told_the_notes_as_the_offices_own_words(
    person, a_case, client, monkeypatch
):
    first = video(person, a_case, "first", stamp=stamp("21:56:19", "BWC2-1"))
    second = video(person, a_case, "second", stamp=stamp("22:01:00", "BWC2-2"))
    incident = incidents.make(a_case, "Stop", [first, second], by=person)
    signed_in(client, person)
    # Without a note, nothing about notes is said.
    chat = CaseChat.objects.create(case=a_case, asked_by=person, name="Gun?")
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="Was a gun found?")
    asked = engine_answering(monkeypatch, ["No [Recording 1, 00:00:30]."])
    case_chat.answer_case_turn(turn.pk)
    assert "office's notes" not in asked[0]["messages"][-1]["content"]
    assert notes.RULE not in asked[0]["messages"][0]["content"]
    post_note(client, first, gun_line(first), "Check the footage.")
    turn = CaseChatTurn.objects.create(chat=chat, number=2, question="And?")
    asked = engine_answering(monkeypatch, ["Yes [Recording 1, 00:00:30]."])
    case_chat.answer_case_turn(turn.pk)
    user = asked[0]["messages"][-1]["content"]
    assert (
        "The office's notes on this recording:\n[00:00:30] asker: Check the footage."
        in user
    )
    assert user.index("I got gun") < user.index("The office's notes on this recording")
    assert notes.RULE in asked[0]["messages"][0]["content"]
    # The incident's Gideon: the note on the incident clock, with its camera.
    inc_chat = incident_chat.new_chat(incident, by=person)
    turn = CaseChatTurn.objects.create(chat=inc_chat, number=1, question="Gun?")
    asked = engine_answering(monkeypatch, ["Yes [21:56:47]."])
    incident_chat.answer_incident_turn(turn.pk)
    user = asked[0]["messages"][-1]["content"]
    assert (
        "The office's notes on the cameras' lines:\n[21:56:47] BWC2-1 asker: "
        "Check the footage." in user
    )
    assert user.index("office's notes on the cameras") < user.index("The question:")
    assert notes.RULE in asked[0]["messages"][0]["content"]
    for row in Row.objects.all():
        assert "footage" not in json.dumps(row.details).lower()


# Processing again ---------------------------------------------------------------------


def test_processing_again_carries_each_note_to_the_line_at_its_moment(person, a_case):
    recording = video(person, a_case, "first")
    line = gun_line(recording)
    notes.set_note(line, "Check the footage.", by=person)
    notes.set_note(recording.transcript.segments.get(start=60.0), "Whose?", by=person)
    kept = notes.remember(recording)
    assert [one["start"] for one in kept] == [30.0, 60.0]
    Transcript.objects.filter(recording=recording).delete()
    made = Transcript.objects.create(recording=recording, language="en")
    for start, end, text in (
        (0.0, 10.0, "Hello."),
        (28.0, 35.0, "Gun."),
        (40.0, 50.0, "Bye."),
    ):
        Segment.objects.create(transcript=made, start=start, end=end, text=text)
    assert notes.carry(kept, made) == 2
    assert made.segments.get(start=28.0).note == "Check the footage."
    assert made.segments.get(start=28.0).note_by == person
    # No line spans 60.0: the nearest start takes it.
    assert made.segments.get(start=40.0).note == "Whose?"
    rows = Row.objects.filter(event="Note carried")
    assert {row.details.get("was") for row in rows} == {"30.0-34.0", "60.0-64.0"}
    assert all("footage" not in json.dumps(row.details) for row in rows)
    # Nothing to carry writes nothing.
    assert notes.carry([], made) == 0
