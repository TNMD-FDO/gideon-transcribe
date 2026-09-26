"""v1.61.0: the assistant on the Incident (Phase 6, chapter 3).

The rules checked here: the Case Chat is told each synced camera's start by
the incident clock, and nothing when the switch is off; the incident
record merges the cameras' Digests onto the clock, reads a camera with no
Digest from its words with numbered labels dropped and names kept, and is
never stored; Propose events runs one call per camera, keeps only the
proposals that pass the checks, replaces the pending ones and offers a
dismissed one again never; Accept makes an Event with source assistant and
its row, Dismiss puts one away; the memo is written on the chronology with
its citations and its event marks, says what it was written from, says when
the chronology changed since, exports to Word with the Chronology, and
counts nowhere until written; a numbered label does not show under a tile;
and the settings, the templates and the documents exist.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from core import (
    assistant,
    chronology,
    engine,
    incident_assistant,
    incidents,
    prompts,
    settings_store,
    tasks,
)
from core.assistant import DigestPart, PromptTemplate
from core.audit import Row
from core.cases import Case
from core.chronology import Event
from core.jobs import Segment, Transcript
from core.models import LoginSession, User
from core.recordings import Batch, MediaState, Recording
from tests.test_prepare import engine_answering

APP = Path(__file__).resolve().parent.parent
ROOT = APP.parent
PASSWORD = "a-long-enough-password"


@pytest.fixture(autouse=True)
def its_own_disk(tmp_path, settings):
    settings.DATA_DIR = tmp_path
    settings.SCRATCH_DIR = tmp_path / "scratch"
    settings.UPLOADS_DIR = tmp_path / "uploads"


@pytest.fixture(autouse=True)
def switched_on(db):
    settings_store.set_to("folder_management", True)
    settings_store.set_to("incidents", True)
    settings_store.set_to("assistant_available", True)
    settings_store.set_to("moments_available", True)
    settings_store.set_to("digests_available", True)


@pytest.fixture(autouse=True)
def no_tasks(monkeypatch):
    monkeypatch.setattr(tasks.propose_incident_events, "defer", lambda **f: None)
    monkeypatch.setattr(tasks.write_incident_memo, "defer", lambda **f: None)


@pytest.fixture
def person(db):
    return User.objects.create_local_admin("asker", PASSWORD)


@pytest.fixture
def a_case(person):
    return Case.objects.create(owner=person, name="Traffic stop")


def signed_in(client, who):
    client.force_login(who)
    LoginSession.objects.create(user=who, session_key=client.session.session_key)
    return client


def stamp(time: str, camera: str):
    return {
        "date": "06/07/2025",
        "time": time,
        "camera": camera,
        "at": 2.0,
        "checked": True,
    }


def video(person, case, title, *, seconds=600.0, stamp=None, lines=()):
    recording = Recording.objects.create(
        batch=Batch.objects.create(user=person),
        user=person,
        case=case,
        title=title,
        original_filename=f"{title}.mp4",
        media_state=MediaState.READY,
        duration_seconds=seconds,
        playback_ready=True,
        stamp=stamp,
        probe={},
    )
    recording.folder.mkdir(parents=True, exist_ok=True)
    (recording.folder / "playback.mp4").write_bytes(b"not really media")
    if lines:
        transcript = Transcript.objects.create(recording=recording, language="en")
        for start, speaker, text in lines:
            Segment.objects.create(
                transcript=transcript,
                start=start,
                end=start + 4,
                text=text,
                speaker=speaker,
                speaker_label=speaker.upper().replace(" ", "_"),
            )
    return recording


def digest(recording, text: str) -> None:
    DigestPart.objects.create(
        transcript=recording.transcript, number=1, text=text, made_at=None
    )


@pytest.fixture
def incident(person, a_case):
    # The first camera's clock reads 21:56:19 two seconds in, so the incident's
    # zero is 21:56:17; the second starts 281 s later.
    first = video(
        person,
        a_case,
        "first",
        stamp=stamp("21:56:19", "BWC2-1"),
        lines=(
            (10.0, "Speaker 1", "Step out of the vehicle for me."),
            (30.0, "Officer Hale", "Hands where I can see them."),
        ),
    )
    second = video(
        person,
        a_case,
        "second",
        stamp=stamp("22:01:00", "BWC2-2"),
        lines=((5.0, "Speaker 4", "Stay in the car, please."),),
    )
    digest(
        first,
        '1. [00:00:10]-[00:00:14] (said) Speaker 1 said "Step out of the vehicle '
        'for me."\n2. [00:00:30]-[00:00:34] (both) Officer Hale asked for hands '
        "in view; the camera showed a grey jacket.",
    )
    return incidents.make(a_case, "Stop", [first, second], by=person)


def cameras_of(incident) -> dict:
    return {one.camera_id(): one for one in incident.cameras.all()}


# The Case Chat is told the placements ------------------------------------------------


def test_the_case_chat_is_told_each_synced_camera_s_start(incident, a_case):
    settings_store.set_to("incidents_case_chat", True)
    read = list(a_case.recordings.order_by("title"))
    block = incident_assistant.placements_block(a_case, read)
    assert block.startswith(
        "Incident Stop, 2 cameras on one clock, the time of day on 06/07/2025."
    )
    assert "Recording 1 (BWC2-1) starts at 21:56:17 and runs to 22:06:17." in block
    assert "Recording 2 (BWC2-2) starts at 22:00:58 and runs to 22:10:58." in block
    assert "the incident's clock is right" in block
    # A camera the question does not read is not in the block; the switch
    # off gives nothing.
    assert "Recording 2" not in incident_assistant.placements_block(a_case, read[:1])
    settings_store.set_to("incidents_case_chat", False)
    assert incident_assistant.placements_block(a_case, read) == ""
    settings_store.set_to("incidents_case_chat", True)
    # The Case chat template carries the time-of-day sentences, and the
    # shipped wording's history knows the new text.
    assert "time of day by the cameras' clock" in prompts.CASE_CHAT
    assert (
        prompts.text_hash(prompts.CASE_CHAT)
        in prompts.SHIPPED_HISTORY["prompt:case_chat"]
    )


# The incident record -----------------------------------------------------------------


def test_the_record_merges_the_cameras_onto_the_clock_and_drops_numbered_labels(
    incident,
):
    record = incident_assistant.record_of(incident)
    lines = [line for _, _, line in record["rows"]]
    names = [name for _, name, _ in record["rows"]]
    # The Digest's lines, rewritten to the clock, their numbers dropped; the
    # second camera read from its words, the label Speaker 4 dropped.
    assert lines[0].startswith("[21:56:27]-[21:56:31] (said)")
    assert names == ["BWC2-1", "BWC2-1", "BWC2-2"]
    assert lines[2] == "[22:01:03] Stay in the car, please."
    assert record["used"] == ["BWC2-1"] and record["transcript_only"] == ["BWC2-2"]
    # A name a person set stays when a camera is read from its words.
    DigestPart.objects.all().delete()
    record = incident_assistant.record_of(incident)
    lines = [line for _, _, line in record["rows"]]
    assert lines[0] == "[21:56:27] Step out of the vehicle for me."
    assert lines[1] == "[21:56:47] Officer Hale: Hands where I can see them."
    assert record["transcript_only"] == ["BWC2-1", "BWC2-2"]
    # Nothing of it is stored.
    assert not hasattr(incident, "record")


def test_the_record_counts_from_the_first_camera_without_a_clock(person, a_case):
    first = video(person, a_case, "one", lines=((10.0, "Speaker 1", "Hello."),))
    made = incidents.make(a_case, "No clock", [first], by=person)
    camera = made.cameras.get()
    incidents.place_by_hand(camera, 0.0, by=person)
    record = incident_assistant.record_of(made)
    assert [line for _, _, line in record["rows"]] == ["[00:00:10] Hello."]
    block = incident_assistant.placements_block(a_case, [first])
    assert "no camera has a clock" in block and "starts at 00:00:00" in block


# Proposed events --------------------------------------------------------------------


def proposals(*items):
    return json.dumps(
        {
            "events": [
                {"time": time, "end": end, "text": text, "rests_on": rests}
                for time, end, text, rests in items
            ]
        }
    )


def test_proposals_are_checked_kept_replaced_and_never_offered_again_once_dismissed(
    incident, person, monkeypatch
):
    settings_store.set_to("incidents_propose", True)
    # One look per camera and no watch phrases here; the windows, the second
    # look and the search have tests of their own (Phase 7 chapter 2).
    settings_store.set_to("incidents_second_look", False)
    settings_store.set_to("incidents_watch_phrases", "")
    cams = cameras_of(incident)
    # An Event already stands at the first camera's 00:00:10.
    chronology.add(
        incident,
        {"at": "10", "text": "Asked out of the car", "camera": str(cams["BWC2-1"].pk)},
        by=person,
    )
    asked = engine_answering(
        monkeypatch,
        [
            proposals(
                # Kept: inside the camera, rests on a real line.
                (
                    "00:00:30",
                    "",
                    "Hands asked into view",
                    "Officer Hale asked for hands",
                ),
                # Dropped: within five seconds of the standing Event.
                ("00:00:12", "", "Out of the car", "Step out of the vehicle"),
                # Dropped: outside the camera.
                ("00:20:00", "", "Too late", "Officer Hale asked for hands"),
                # Dropped: rests on nothing the camera was given.
                ("00:00:31", "", "Made up", "the dog barked at the moon"),
            ),
            proposals(
                ("00:00:05", "00:00:09", "Passenger told to stay", "Stay in the car")
            ),
        ],
    )
    assert incident_assistant.ask_for_proposals(incident, by=person)
    incident_assistant.propose(incident.pk)
    incident.refresh_from_db()
    assert incident.proposals_state == "done"
    assert incident.proposals_found == 2 and incident.proposals_cameras == 2
    # One call per camera, the schema and no thinking, with the standing
    # Event named so it is not proposed again.
    assert len(asked) == 2
    assert asked[0]["schema"]["properties"]["events"]["items"]["required"] == [
        "time",
        "end",
        "text",
        "why",
        "rests_on",
    ]
    assert asked[0]["thinking"] is False
    user = asked[0]["messages"][-1]["content"]
    assert "Camera BWC2-1 starts 21:56:17 by the cameras' clock" in user
    assert "21:56:27 Asked out of the car" in user
    # The second camera is told the first camera's proposal as well.
    assert "21:56:47 Hands asked into view" in asked[1]["messages"][-1]["content"]
    pending = list(incident.events.filter(proposed=True, dismissed=False))
    assert [one.text for one in pending] == [
        "Hands asked into view",
        "Passenger told to stay",
    ]
    assert pending[0].source == "assistant" and pending[0].camera == cams["BWC2-1"]
    assert (
        pending[0].at == 30.0 and pending[0].rests_on == "Officer Hale asked for hands"
    )
    assert pending[1].until == cams["BWC2-2"].starts_at + 9.0
    words = incident_assistant.proposals_json(incident)
    assert words["pending"] == 2 and words["words"].startswith("Proposed 2 events at ")
    # The state carries them as proposed, with the words each rests on; the
    # exports and the counts leave them out.
    rows = chronology.events_json(incident)
    assert [one["proposed"] for one in rows] == [False, True, True]
    assert rows[1]["source_words"] == "Proposed from BWC2-1"
    assert rows[1]["rests_on"] == "Officer Hale asked for hands"
    assert len(chronology._rows(incident)) == 1
    assert incidents.strip_rows(incident.case)[0]["events"] == 1
    # Dismiss one; Accept the other: an Event with source assistant, its row,
    # and the person who accepted it.
    incident_assistant.dismiss(pending[1], by=person)
    incident_assistant.accept(pending[0], by=person)
    kept = incident.events.filter(proposed=False).order_by("at")
    assert [one.text for one in kept] == [
        "Asked out of the car",
        "Hands asked into view",
    ]
    assert kept[1].source == "assistant" and kept[1].added_by == person
    assert chronology.events_json(incident)[1]["source_words"] == "Assistant, BWC2-1"
    assert chronology._rows(incident)[1]["assistant"]
    events = list(Row.objects.filter(category="cases").values_list("event", flat=True))
    assert "events proposed" in events and "event dismissed" in events
    added = Row.objects.filter(event="event added").order_by("-at").first()
    assert added.details.get("source") == "assistant"
    assert Row.objects.filter(
        event="AI assistant call", details__feature="incident_events"
    ).exists()
    for row in Row.objects.all():
        assert "Hands" not in json.dumps(row.details) and "Stay" not in json.dumps(
            row.details
        )
    # A dismissed proposal is off the page and a second run offers nothing
    # within five seconds of it on that camera.
    assert all(not one["proposed"] for one in chronology.events_json(incident))
    engine_answering(
        monkeypatch,
        [
            proposals(),
            proposals(
                ("00:00:07", "", "Passenger told to stay, again", "Stay in the car")
            ),
        ],
    )
    incident.proposals_state = ""
    incident.save()
    assert incident_assistant.ask_for_proposals(incident, by=person)
    incident_assistant.propose(incident.pk)
    assert not incident.events.filter(proposed=True, dismissed=False).exists()
    assert incident.events.filter(proposed=True, dismissed=True).count() == 1


def test_proposals_need_the_switch_and_a_synced_camera_with_words(
    incident, person, a_case
):
    settings_store.set_to("incidents_propose", False)
    # With the assistant off, the watch phrases alone keep the button (Phase
    # 7 chapter 2); with none listed, nothing offers it.
    assert incident_assistant.proposals_possible(incident)
    assert incident_assistant.proposer_offered()
    settings_store.set_to("incidents_watch_phrases", "")
    assert not incident_assistant.proposals_possible(incident)
    assert not incident_assistant.proposer_offered()
    assert not incident_assistant.ask_for_proposals(incident, by=person)
    settings_store.set_to("incidents_propose", True)
    silent = video(person, a_case, "silent", stamp=stamp("21:30:00", "BWC2-9"))
    alone = incidents.make(a_case, "Silent", [silent], by=person)
    assert not incident_assistant.proposals_possible(alone)
    assert incident_assistant.proposals_json(alone)["possible"] is False


def test_the_watch_phrases_are_searched_by_the_app_itself(person, a_case):
    # The floor (Phase 7 chapter 2): with the assistant off, the run is the
    # search alone; whole words, any case, the plural, a phrase heard again
    # within the join window joining its first hit, and "begun" never "gun".
    settings_store.set_to("incidents_propose", False)
    settings_store.set_to(
        "incidents_watch_phrases", "gun\n\nGun\nstep out of the vehicle\n"
    )
    assert settings_store.watch_phrases() == ["gun", "step out of the vehicle"]
    assert incident_assistant.phrase_pattern("gun").search("two guns down")
    assert not incident_assistant.phrase_pattern("gun").search("we had begun")
    # The run folds the line first: case, curly apostrophes, spacing.
    assert incident_assistant.phrase_pattern("I can't breathe").search(
        incident_assistant._plain("He said I can\u2019t  BREATHE")
    )
    cam = video(
        person,
        a_case,
        "one",
        lines=(
            (10.0, "Speaker 1", "Step out of the vehicle, please."),
            (270.0, "Speaker 4", "I got, I got gun, I got gun."),
            (285.0, "Speaker 4", "Guns down."),
            (400.0, "Speaker 1", "We had begun the search."),
        ),
    )
    made = incidents.make(a_case, "Search", [cam], by=person)
    camera = made.cameras.get()
    incidents.place_by_hand(camera, 0.0, by=person)
    assert incident_assistant.proposals_possible(made)
    assert incident_assistant.ask_for_proposals(made, by=person)
    incident_assistant.propose(made.pk)
    made.refresh_from_db()
    assert made.proposals_state == "done"
    pending = list(made.events.filter(proposed=True).order_by("at"))
    assert [(one.at, one.source) for one in pending] == [
        (10.0, "watch"),
        (270.0, "watch"),
    ]
    assert (
        pending[0].text
        == 'Watch phrase "step out of the vehicle": Step out of the vehicle, please.'
    )
    assert (
        pending[1].text
        == 'Watch phrase "gun": I got, I got gun, I got gun. (said 2 times)'
    )
    assert pending[1].why == 'The office watches for "gun".'
    assert pending[1].rests_on == "I got, I got gun, I got gun."
    assert made.proposals_watch == 2 and made.proposals_found == 2
    words = incident_assistant.proposals_json(made)
    assert words["on"] and "2 from the watch phrases" in words["words"]
    names = {str(camera.pk): camera.camera_id()}
    assert (
        chronology.source_words(pending[1], names)
        == f"Watch phrase, {camera.camera_id()}"
    )
    # No engine call, so no AI assistant call row; the run's own row counts the hits.
    assert not Row.objects.filter(event="AI assistant call").exists()
    proposed = Row.objects.filter(event="events proposed").order_by("-at").first()
    assert (
        proposed.details.get("watch_hits") == 2 and proposed.details.get("windows") == 0
    )
    for row in Row.objects.all():
        assert "gun" not in json.dumps(row.details)
    # Accepted, the Event keeps its source and its why, and Edit keeps the source.
    incident_assistant.accept(pending[1], by=person)
    pending[1].refresh_from_db()
    chronology.change(pending[1], {"at": "270", "text": "Gun found"}, by=person)
    pending[1].refresh_from_db()
    assert pending[1].source == "watch" and pending[1].why.startswith("The office")
    rows = chronology._rows(made)
    assert rows[0]["why"] == 'The office watches for "gun".'
    assert "Why it matters" in chronology.spreadsheet(made).decode("utf-8-sig")


def test_the_run_reads_in_windows_with_a_second_look_and_says_why(
    person, a_case, monkeypatch
):
    settings_store.set_to("incidents_propose", True)
    settings_store.set_to("incidents_watch_phrases", "")
    settings_store.set_to("incidents_events_context", "A public defender's office.")
    cam = video(
        person,
        a_case,
        "long",
        seconds=1500.0,
        lines=(
            (10.0, "Speaker 1", "Step out of the vehicle, please."),
            (700.0, "Speaker 4", "I got, I got gun, I got gun."),
            (1300.0, "Speaker 1", "We had begun the search."),
        ),
    )
    made = incidents.make(a_case, "Long", [cam], by=person)
    camera = made.cameras.get()
    incidents.place_by_hand(camera, 0.0, by=person)
    monkeypatch.setattr(engine, "is_reachable", lambda: True)
    monkeypatch.setattr(engine, "address", lambda: "http://gideon-generator:8000/v1")
    asked = []

    def complete(messages, **options):
        asked.append(messages[-1]["content"])
        user = messages[-1]["content"]
        # The second window's first look finds the gun; its second look adds
        # a why-bearing item; one answer is cut short.
        if "I got gun" in user and "You proposed" not in user:
            text = json.dumps(
                {
                    "events": [
                        {
                            "time": "00:11:40",
                            "end": "",
                            "text": "An officer said he had found a gun.",
                            "why": "A weapon changes the stop.",
                            "rests_on": "I got, I got gun",
                        }
                    ]
                }
            )
            return {
                "text": text,
                "finish_reason": "length",
                "input_tokens": 1,
                "output_tokens": 1,
                "model": "m",
            }
        return {
            "text": json.dumps({"events": []}),
            "finish_reason": "stop",
            "input_tokens": 1,
            "output_tokens": 1,
            "model": "m",
        }

    monkeypatch.setattr(engine, "complete", complete)
    assert incident_assistant.ask_for_proposals(
        made, by=person, look_for="anything about the gun"
    )
    made.refresh_from_db()
    assert made.proposals_look_for == "anything about the gun"
    incident_assistant.propose(made.pk)
    made.refresh_from_db()
    # Three windows of ten minutes, two looks each.
    assert len(asked) == 6
    assert sum("You proposed" in one for one in asked) == 3
    assert all(
        "About the office and its cases: A public defender's office." in one
        for one in asked
    )
    assert all("look for: anything about the gun" in one for one in asked)
    assert "one stretch of" in asked[0]
    # The second look is told what the first proposed.
    second = [one for one in asked if "You proposed" in one and "I got gun" in one][0]
    assert "00:11:40 An officer said he had found a gun." in second
    assert made.proposals_look_for == "" and made.proposals_cut == 1
    pending = list(made.events.filter(proposed=True))
    assert len(pending) == 1 and pending[0].why == "A weapon changes the stop."
    assert pending[0].at == 700.0 and pending[0].source == "assistant"
    words = incident_assistant.proposals_json(made)["words"]
    assert words.startswith("Proposed 1 event at ") and "1 answer cut short" in words
    row = chronology.events_json(made)[0]
    assert row["why"] == "A weapon changes the stop." and row["proposed"]
    call = Row.objects.filter(event="AI assistant call").order_by("-at").first()
    assert call.details.get("windows") == 3 and call.details.get("second_looks") == 3
    assert call.details.get("look_for") is True and call.details.get("cut_short") == 1
    assert "anything about the gun" not in json.dumps(call.details)
    # With the second look off, one look a window.
    settings_store.set_to("incidents_second_look", False)
    asked.clear()
    made.proposals_state = ""
    made.save()
    assert incident_assistant.ask_for_proposals(made, by=person)
    incident_assistant.propose(made.pk)
    assert len(asked) == 3 and not any("You proposed" in one for one in asked)


def test_the_windows_split_a_record_by_its_times():
    lines = ["[00:00:05] a", "[00:09:59] b", "no time", "[00:10:00] c", "[00:31:00] d"]
    assert incident_assistant._windows(lines, 600) == [
        ["[00:00:05] a", "[00:09:59] b", "no time"],
        ["[00:10:00] c"],
        ["[00:31:00] d"],
    ]


def test_a_heading_without_its_colon_is_still_a_heading():
    """v1.90.0: the model drops the colon, sets a part in bold, or writes it
    in capitals; the page and the Word export still read a heading."""
    given = (
        "Summary\nOne paragraph.\n\n**Breakdown**\n**Arrival 20:24 to 20:27**\n"
        "Text here.\nPOINTS FOR COUNSEL\n1. A point.\nNone noticed.\n"
        "**It was a long night.**\n"
    )
    assert prompts.with_heading_colons(given) == (
        "Summary:\nOne paragraph.\n\nBreakdown:\nArrival 20:24 to 20:27:\n"
        "Text here.\nPoints for counsel:\n1. A point.\nNone noticed.\n"
        "**It was a long night.**"
    )


# The Incident memo ------------------------------------------------------------------


MEMO = (
    "Summary:\n"
    "At about 21:56 on 6 July 2025 a driver was asked out of the car [21:56:27] "
    "(Event 1) and told to keep hands in view [21:56:47] (Event 2).\n\n"
    "The cameras:\n"
    "BWC2-1 ran from [21:56:17] to [22:06:17]; BWC2-2 from [22:00:58].\n\n"
    "What happened:\n"
    "The passenger was told to stay in the car on BWC2-2 [22:01:03]. Nothing "
    "was seen at [23:59:00].\n\n"
    "Unclear parts:\n"
    "None.\n"
)
# What the app keeps: the model's The cameras lines replaced by its own
# (v1.89.0), each camera's kind, its span as citations, and how it was read.
KEPT = MEMO.replace(
    "BWC2-1 ran from [21:56:17] to [22:06:17]; BWC2-2 from [22:00:58].\n\n",
    "BWC2-1: camera, [21:56:17] to [22:06:17].\n"
    "BWC2-2: camera, [22:00:58] to [22:10:58]; no picture read.\n\n",
)
# The facts sheet (v1.91.0): the first pass's answer, out of order, with one
# quote the record holds word for word, one it does not, and one time off
# the record.
SHEET = json.dumps(
    {
        "people": [
            {
                "who": "The driver",
                "how": "the person asked out of the car",
                "where": "21:56:27 on BWC2-1",
                "cameras": ["BWC2-1"],
            }
        ],
        "timeline": [
            {
                "at": "22:01:03",
                "camera": "BWC2-2",
                "what": "The passenger was told to stay in the car.",
                "quote": "Stay in the car please",
                "said_by": "the officer wearing BWC2-2",
                "event": 0,
            },
            {
                "at": "21:56:27",
                "camera": "BWC2-1",
                "what": "The driver was asked out of the car.",
                "quote": "Step out of the vehicle for me.",
                "said_by": "an officer",
                "event": 1,
            },
            {
                "at": "22:10:00",
                "camera": "BWC2-2",
                "what": "The cameras stopped.",
                "quote": "",
                "said_by": "",
                "event": 0,
            },
        ],
        "rights": [],
        "questions_before_rights": [
            {
                "at": "21:57:00",
                "camera": "BWC2-1",
                "asked_by": "Speaker 4",
                "to": "Unidentified speaker",
                "question": "Whose car is it?",
                "answer": "Speaker 2 said it was a rental.",
            }
        ],
        "searches_and_force": [],
        "statements": [
            {
                "at": "23:59:00",
                "camera": "BWC2-1",
                "who": "the driver",
                "quote": "Not mine",
                "prompted": True,
            }
        ],
        "gaps": [],
        "outcome": "Both went on their way.",
    }
)


def test_the_memo_is_written_on_the_chronology_with_citations_and_marks(
    incident, person, monkeypatch, client
):
    settings_store.set_to("incidents_memo", True)
    cams = cameras_of(incident)
    first = chronology.add(
        incident, {"at": "10", "text": "Asked out of the car"}, by=person
    )
    second = chronology.add(
        incident,
        {
            "at": "30",
            "text": "Hands where I can see them.",
            "source": "words",
            "camera": str(cams["BWC2-1"].pk),
        },
        by=person,
    )
    asked = engine_answering(monkeypatch, [SHEET, MEMO])
    memo = incident_assistant.ask_for_memo(incident, by=person)
    assert memo is not None and memo.state == "queued"
    incident_assistant.write_memo(memo.pk)
    memo.refresh_from_db()
    assert memo.state == "done" and memo.text == KEPT.strip()
    # Two passes (v1.91.0, ADR 0016). The first, the facts sheet, is given
    # the cameras, the chronology numbered, the record on the clock, the
    # fixed incident rules with the sheet's template, and answers as data.
    assert len(asked) == 2
    system = asked[0]["messages"][0]["content"]
    user = asked[0]["messages"][-1]["content"]
    assert prompts.MEMO_SHEET in system and prompts.INCIDENT_RULES in system
    assert prompts.MEMO_SHEET_FORMAT in system and prompts.INCIDENT_MEMO not in system
    assert asked[0]["schema"] == prompts.memo_sheet_schema()
    assert asked[0]["max_completion_tokens"] == 12000
    # The second, the memo, is given the sheet and never the record.
    writer = asked[1]["messages"]
    assert prompts.INCIDENT_MEMO in writer[0]["content"]
    assert prompts.NARRATIVE_RULES in writer[0]["content"]
    given = writer[-1]["content"]
    assert "The facts sheet, drawn from the record" in given
    assert "Timeline, in order:" in given and "Outcome:" in given
    assert "BWC2-2: [22:01:03] Stay in the car, please." not in given
    assert "Event 1, 21:56:27, Asked out of the car" not in given
    assert asked[1]["max_completion_tokens"] == 4000
    # The app's checks on the sheet: in time order, the quotes looked for in
    # the record, a time off the record marked, the end reached.
    checks = memo.sheet["checks"]
    assert checks["moments"] == 3 and checks["reaches_the_end"] is True
    assert checks["quotes_not_found"] == 1 and checks["times_outside"] == 1
    assert checks["last_moment"] == "22:10:00"
    timeline = memo.sheet["timeline"]
    assert [one["at"] for one in timeline] == ["21:56:27", "22:01:03", "22:10:00"]
    assert timeline[1]["verbatim"] is True
    statement = memo.sheet["statements"][0]
    assert statement["time_outside"] is True and statement["verbatim"] is False
    assert not memo.sheet_cut_short
    sheet_text = prompts.sheet_text(memo.sheet)
    assert "- [21:56:27] BWC2-1: The driver was asked out of the car." in sheet_text
    assert '"Step out of the vehicle for me." (Event 1)' in sheet_text
    assert (
        "(time not on the record) BWC2-1: the driver (answering a question): "
        in sheet_text
    )
    assert "Not mine (not found word for word in the record)" in sheet_text
    # The machine's labels are scrubbed (v1.91.1): a voice on the camera.
    asked_before = memo.sheet["questions_before_rights"][0]
    assert asked_before["asked_by"] == "a voice on BWC2-1"
    assert asked_before["to"] == "a voice on BWC2-1"
    assert asked_before["answer"] == "a voice said it was a rental."
    assert asked_before["question"] == "Whose car is it?"
    assert "Speaker" not in sheet_text and "Unidentified" not in sheet_text
    assert (
        "- [21:57:00] BWC2-1: a voice on BWC2-1 asked a voice on BWC2-1: "
        '"Whose car is it?"; answer: "a voice said it was a rental."'
    ) in sheet_text
    words = incident_assistant.sheet_words(memo.sheet)
    assert words.startswith("3 moments, the last at 22:10:00; the timeline reaches")
    assert "3 quotes checked against the record, 1 not found word for word." in words
    assert "timeline" in prompts.MEMO_SHEET_FORMAT.split('"gaps"')[1]
    assert "1 time not on the record." in words
    assert (
        "Event 1, 21:56:27, Asked out of the car; seen on BWC2-1; added by a person."
        in user
    )
    assert (
        "Event 2, 21:56:47, Hands where I can see them.; seen on BWC2-1; "
        "from the words on BWC2-1." in user
    )
    assert "BWC2-2: [22:01:03] Stay in the car, please." in user
    assert "BWC2-1, 21:56:17 to 22:06:17, from its clock, checked" in user
    # The citations inside the span, the marks by the numbering at writing;
    # the sheet's times cite too (v1.91.0).
    assert memo.citations == {
        "[21:56:27]": 10.0,
        "[21:56:47]": 30.0,
        "[21:56:17]": 0.0,
        "[22:06:17]": 600.0,
        "[22:00:58]": 281.0,
        "[22:01:03]": 286.0,
        # The app's cameras line ends BWC2-2 at its end (v1.89.0).
        "[22:10:58]": 881.0,
        "[22:10:00]": 823.0,
        "[21:57:00]": 43.0,
    }
    assert memo.event_numbers == {"1": str(first.pk), "2": str(second.pk)}
    assert memo.cameras_used == ["BWC2-1"]
    assert memo.cameras_transcript_only == ["BWC2-2"]
    assert memo.events_count == 2 and memo.record_lines == 3
    words = incident_assistant.written_words(memo)
    assert words.startswith("Written ")
    assert (
        "the transcripts and the vision of 1 camera and the words alone of 1" in words
    )
    assert "on the chronology's 2 events." in words
    assert incident_assistant.stale_words(memo, incident) == ""
    assert incident_assistant.memo_line(incident).startswith("memo written ")
    state = incident_assistant.memo_json(incident)
    assert state["state"] == "done" and state["event_numbers"]["2"] == str(second.pk)
    assert state["notice"] and state["busy"] is False
    # The tab has the sheet, its words, and Rewrite from the sheet on offer.
    assert state["sheet_text"].startswith("People:")
    assert state["sheet"]["checks"]["moments"] == 3
    assert state["sheet_words"] and state["rewrite_possible"] is True
    # The row: metadata only, the calls and the moments counted.
    row = Row.objects.get(event="AI assistant call", details__feature="incident_memo")
    assert row.details["events"] == 2 and "Asked" not in json.dumps(row.details)
    assert row.details["calls"] == 2 and row.details["moments"] == 3
    assert row.details["from_sheet"] is False and "Not mine" not in json.dumps(
        row.details
    )
    # The chronology changes: the tab says so, and so does the case page.
    chronology.add(incident, {"at": "40", "text": "Tow truck"}, by=person)
    assert incident_assistant.stale_words(memo, incident) == (
        "The chronology has changed since this memo was written: 1 event added. "
        "Regenerate to write it on them."
    )
    assert incident_assistant.memo_line(incident).endswith(", 1 event newer")
    # A stale memo cannot be rewritten from its sheet; the page says use Regenerate.
    assert incident_assistant.memo_json(incident)["rewrite_possible"] is False
    assert incident_assistant.ask_for_rewrite(incident, by=person) is None
    chronology.change(first, {"at": "11", "text": "Asked out"}, by=person)
    chronology.remove(Event.objects.get(text="Tow truck"), by=person)
    assert "an event changed" in incident_assistant.stale_words(memo, incident)
    incidents.place_by_hand(cams["BWC2-2"], 290.0, by=person)
    assert "a camera was synced or joined" in incident_assistant.stale_words(
        memo, incident
    )
    # The Word export carries the memo and the Chronology, and writes its row.
    signed_in(client, person)
    answer = client.post(
        f"/case/{incident.case_id}/incident/{incident.pk}/export/memo", {}
    )
    assert answer.status_code == 200
    assert answer["Content-Disposition"].endswith('"Incident memo - Stop.docx"')
    assert answer.content[:2] == b"PK"
    assert Row.objects.filter(event="incident memo exported").exists()
    # Every page carries the marking, the title and the page of pages (v1.90.0).
    import io as _io

    from docx import Document

    exported = Document(_io.BytesIO(answer.content))
    head = " ".join(p.text for p in exported.sections[0].header.paragraphs)
    foot = exported.sections[0].footer._element.xml
    assert "Privileged and confidential. Attorney work product." in head
    assert "Traffic stop, Stop - Incident memo" in head
    assert "PAGE" in foot and "NUMPAGES" in foot
    # The facts sheet is printed after the memo, before the Chronology (v1.91.0).
    body = "\n".join(p.text for p in exported.paragraphs)
    assert "Facts sheet" in body and "What the memo was written from" in body
    assert "The driver was asked out of the car." in body
    assert body.index("Facts sheet") < body.index("Chronology")

    # Regenerate resets the one row; a failed engine says so.
    def failing(messages, **options):
        from core import engine

        raise engine.Problem(engine.TIMEOUT, "took too long")

    monkeypatch.setattr("core.engine.complete", failing)
    again = incident_assistant.ask_for_memo(incident, by=person)
    assert again.pk == memo.pk and again.state == "queued" and again.text == ""
    incident_assistant.write_memo(again.pk)
    again.refresh_from_db()
    assert again.state == "failed" and again.reason_class == "llm_timeout"
    assert incident_assistant.memo_json(incident)["reason_words"]
    assert incident_assistant.memo_line(incident) == "no memo yet"


def test_a_memo_cut_at_the_cap_goes_on_from_where_it_stopped(
    incident, person, monkeypatch
):
    """v1.90.1: the twelve-camera memo hit its 4,000-token cap five parts
    short. A cut memo is continued, up to three calls, each shown the memo
    so far and told to go on; the pieces are stitched, the memo is whole,
    and the audit row counts the calls."""
    settings_store.set_to("incidents_memo", True)
    cams = cameras_of(incident)
    chronology.add(incident, {"at": "10", "text": "Asked out of the car"}, by=person)
    chronology.add(
        incident,
        {
            "at": "30",
            "text": "Hands where I can see them.",
            "source": "words",
            "camera": str(cams["BWC2-1"].pk),
        },
        by=person,
    )
    monkeypatch.setattr(engine, "is_reachable", lambda: True)
    monkeypatch.setattr(engine, "address", lambda: "http://gideon-generator:8000/v1")
    head, tail = MEMO.split("What happened:")
    pieces = [(SHEET, "stop"), (head + "What happened:", "length"), (tail, "stop")]
    asked = []

    def complete(messages, **options):
        asked.append(messages)
        text, finish = pieces.pop(0)
        return {
            "text": text,
            "finish_reason": finish,
            "input_tokens": 10,
            "output_tokens": 5,
            "model": "the-model",
        }

    monkeypatch.setattr(engine, "complete", complete)
    memo = incident_assistant.ask_for_memo(incident, by=person)
    incident_assistant.write_memo(memo.pk)
    memo.refresh_from_db()
    assert memo.state == "done" and not memo.cut_short
    assert memo.text == KEPT.strip()
    # The second call: the same system and input, then the memo so far as the
    # assistant's turn, then the ask to go on.
    assert len(asked) == 3
    assert asked[2][0] == asked[1][0] and asked[2][1] == asked[1][1]
    assert asked[2][2]["role"] == "assistant"
    assert asked[2][2]["content"].endswith("What happened:")
    assert asked[2][3] == {"role": "user", "content": prompts.CONTINUE_MEMO}
    said = Row.objects.filter(event="AI assistant call").order_by("-at").first()
    assert said.details["calls"] == 3 and said.details["cut_short"] is False
    assert said.details["input_tokens"] == 30 and said.details["output_tokens"] == 15


def test_rewrite_from_the_sheet_runs_the_second_pass_alone(
    incident, person, monkeypatch, client
):
    """v1.91.0: a written memo keeps its facts sheet, and Rewrite from the
    sheet writes the memo again from it without reading the cameras: one
    call, the sheet unchanged, the row saying so. A stale memo refuses."""
    settings_store.set_to("incidents_memo", True)
    chronology.add(incident, {"at": "10", "text": "Asked out of the car"}, by=person)
    engine_answering(monkeypatch, [SHEET, MEMO])
    memo = incident_assistant.ask_for_memo(incident, by=person)
    incident_assistant.write_memo(memo.pk)
    memo.refresh_from_db()
    assert memo.state == "done" and memo.sheet["checks"]["moments"] == 3
    sheet_before = json.dumps(memo.sheet, sort_keys=True)
    asked = engine_answering(monkeypatch, [MEMO])
    again = incident_assistant.ask_for_rewrite(incident, by=person)
    assert again is not None and again.pk == memo.pk
    assert again.state == "queued" and again.text == "" and again.sheet
    incident_assistant.write_memo(again.pk, from_sheet=True)
    again.refresh_from_db()
    assert again.state == "done" and again.text == KEPT.strip()
    assert json.dumps(again.sheet, sort_keys=True) == sheet_before
    assert len(asked) == 1
    assert prompts.INCIDENT_MEMO in asked[0]["messages"][0]["content"]
    assert "Timeline, in order:" in asked[0]["messages"][-1]["content"]
    row = Row.objects.filter(event="AI assistant call").order_by("-at").first()
    assert row.details["from_sheet"] is True and row.details["calls"] == 1
    assert again.cameras_used == ["BWC2-1"]
    assert again.cameras_transcript_only == ["BWC2-2"]
    # From the page: the action, and its refusal once the chronology moved.
    signed_in(client, person)
    act = f"/case/{incident.case_id}/incident/{incident.pk}/act"
    engine_answering(monkeypatch, [MEMO])
    got = client.post(act, {"action": "memo_rewrite"})
    assert got.status_code == 200 and got.json()["said"] == (
        "Rewriting the memo from its facts sheet."
    )
    chronology.add(incident, {"at": "40", "text": "Tow truck"}, by=person)
    got = client.post(act, {"action": "memo_rewrite"})
    assert got.status_code == 400 and "use Regenerate" in got.json()["error"]


def test_the_sheet_is_checked_and_a_bad_one_is_a_bad_output(incident, person):
    """The app's checks, on their own: an unreadable answer is the engine's
    bad output; a time before the incident is marked, never dropped; the
    timeline that stops early is said so."""
    from core import engine

    with pytest.raises(engine.Problem) as caught:
        incident_assistant._parse_sheet("not json at all {")
    assert caught.value.reason == engine.BAD_OUTPUT
    with pytest.raises(engine.Problem):
        incident_assistant._parse_sheet("[1, 2]")
    sheet = incident_assistant._parse_sheet(
        json.dumps(
            {
                "people": "not a list",
                "timeline": [
                    {
                        "at": "21:56:30",
                        "camera": "BWC2-1",
                        "what": "Early.",
                        "quote": "",
                    },
                    "not an entry",
                ],
                "outcome": 7,
            }
        )
    )
    assert sheet["people"] == [] and len(sheet["timeline"]) == 1
    record = incident_assistant.record_of(incident)
    checked = incident_assistant.check_sheet(incident, sheet, record)
    assert checked["checks"]["reaches_the_end"] is False
    assert checked["checks"]["moments"] == 1 and checked["checks"]["times_outside"] == 0
    assert "stops before the last camera does" in incident_assistant.sheet_words(
        checked
    )
    assert incident_assistant.sheet_words({}) == ""
    assert prompts.sheet_text({}) == "" and prompts.sheet_text("x") == ""


def test_a_continuation_is_stitched_at_a_sentence_end():
    """v1.91.1: the office's memo read "Aldridgeanswered" where a cut fell at a
    word boundary. The cut text is trimmed back to its last whole sentence
    and the model is told to begin with the sentence that was cut."""
    stitch = prompts.stitch_continuation
    assert prompts.cut_back("He said no. Then Aldridge") == "He said no."
    assert prompts.cut_back('He said no. Then Aldridge answered, "Yes') == (
        "He said no."
    )
    assert prompts.cut_back("Was it his? (It was.) Then") == "Was it his? (It was.)"
    assert prompts.cut_back("no sentence end at all") == "no sentence end at all"
    assert stitch("He said no. Then Aldridge", "Then Aldridge answered no.") == (
        "He said no. Then Aldridge answered no."
    )
    assert stitch("Summary:", "\n\nThe stop began at [20:24:36].") == (
        "Summary:\nThe stop began at [20:24:36]."
    )
    assert stitch("the officer left.", "Points for counsel:\n1. The stop.") == (
        "the officer left.\nPoints for counsel:\n1. The stop."
    )
    assert stitch("", "Whole.") == "Whole." and stitch("Whole.", "") == "Whole."
    assert "cut back to the end of its last whole sentence" in prompts.CONTINUE_MEMO


def test_the_memo_needs_a_synced_camera_and_reads_a_long_camera_by_words_alone(
    incident, person, monkeypatch
):
    settings_store.set_to("incidents_memo", True)
    possible, why = incident_assistant.memo_possible(incident)
    assert possible and why == ""
    # One sitting (Phase 8 chapter 9): the first camera's Digest is made long.
    # When the whole does not fit the window, that camera is read by its words
    # alone (its transcript in place of its Digest) and the memo says so; when
    # even that does not fit, the memo fails as too long, since a camera
    # without a Digest has nothing more to give up.
    first = cameras_of(incident)["BWC2-1"].recording.transcript
    first.digest_parts.all().delete()
    DigestPart.objects.create(
        transcript=first,
        number=1,
        text="\n".join(
            f"{n + 1}. [00:00:{10 + n:02d}]-[00:00:{14 + n:02d}] (both) A long "
            "description of the roadside and the two officers standing by the car."
            for n in range(40)
        ),
        made_at=None,
    )
    asked = engine_answering(monkeypatch, [SHEET, MEMO, SHEET, MEMO])
    # The sheet's call reads the record (v1.91.0); the memo's reads the sheet.
    system_cost = prompts.tokens(
        prompts.system_message(
            PromptTemplate.named(PromptTemplate.GROUND_RULES).text,
            PromptTemplate.named(PromptTemplate.MEMO_SHEET).text,
            prompts.MEMO_SHEET_FORMAT + "\n\n" + prompts.INCIDENT_RULES,
        )
    )
    settings_store.set_to("incidents_memo_answer_tokens", 500)
    words_alone = incident_assistant.record_of(incident, words_alone={"BWC2-1"})
    assert words_alone["words_alone"] == ["BWC2-1"] and not words_alone["used"]
    event_lines, _ = incident_assistant._chronology_lines(incident)
    alone_cost = prompts.tokens(
        prompts.incident_memo_input(
            incident_assistant._cameras_line(incident),
            event_lines,
            incident_assistant._record_text(incident, words_alone, set()),
        )
    )
    settings_store.set_to("engine_window_tokens", 4096)
    settings_store.set_to(
        "incidents_memo_sheet_answer_tokens", 4096 - system_cost - alone_cost - 5
    )
    memo = incident_assistant.ask_for_memo(incident, by=person)
    incident_assistant.write_memo(memo.pk)
    memo.refresh_from_db()
    assert memo.state == "done"
    assert memo.cameras_words_alone == ["BWC2-1"]
    assert memo.cameras_not_read == []
    sent = asked[0]["messages"][-1]["content"]
    assert "Step out of the vehicle" in sent and "grey jacket" not in sent
    assert "1 camera by words alone" in incident_assistant.written_words(memo)
    said = Row.objects.filter(event="AI assistant call").order_by("-at").first()
    assert said.details["words_alone"] == 1
    assert said.details["window_source"] == "setting"
    # Pinned, the camera stays in with its Digest, and the whole no longer fits.
    camera = cameras_of(incident)["BWC2-1"]
    incidents.pin(camera, True, by=person)
    memo = incident_assistant.ask_for_memo(incident, by=person)
    incident_assistant.write_memo(memo.pk)
    memo.refresh_from_db()
    assert memo.state == "failed" and memo.reason_class == "llm_too_long"
    incidents.pin(camera, False, by=person)
    assert Row.objects.filter(event="Camera pinned").count() == 1
    assert Row.objects.filter(event="Camera unpinned").count() == 1
    settings_store.set_to("incidents_memo_sheet_answer_tokens", 4096 - system_cost - 5)
    memo = incident_assistant.ask_for_memo(incident, by=person)
    incident_assistant.write_memo(memo.pk)
    memo.refresh_from_db()
    assert memo.state == "failed" and memo.reason_class == "llm_too_long"
    # No synced camera: no memo.
    settings_store.set_to("engine_window_tokens", 131072)
    for camera in incident.cameras.all():
        camera.placed = incidents.GUESS
        camera.save()
    possible, why = incident_assistant.memo_possible(incident)
    assert not possible and why.startswith("No camera is synced yet")
    assert incident_assistant.ask_for_memo(incident, by=person) is None
    settings_store.set_to("incidents_memo", False)
    state = incident_assistant.memo_json(incident)
    assert state["on"] is False and state["state"] == ""


# The page ---------------------------------------------------------------------------


def test_the_page_and_the_act_endpoint_carry_chapter_3(incident, person, client):
    settings_store.set_to("incidents_memo", True)
    settings_store.set_to("incidents_propose", True)
    signed_in(client, person)
    url = f"/case/{incident.case_id}/incident/{incident.pk}"
    page = client.get(url).content.decode()
    assert 'data-panel="memo"' in page
    assert '<span class="w">Memo</span></button>' in page
    assert 'id="export-memo"' in page
    assert "exportMemo:" in page
    state = client.get(url + "/state").json()
    assert state["memo"]["on"] and state["memo"]["state"] == ""
    assert state["memo"]["before"] == "2 cameras synced, 0 events on the chronology."
    assert state["proposals"]["on"] and state["proposals"]["possible"]
    # Propose events and Write the memo queue their runs and say so.
    got = client.post(url + "/act", {"action": "propose"}).json()
    assert got["ok"] and got["state"]["proposals"]["busy"]
    got = client.post(url + "/act", {"action": "memo"}).json()
    assert got["ok"] and got["state"]["memo"]["busy"]
    got = client.post(url + "/act", {"action": "memo_cancel"}).json()
    assert got["ok"] and got["state"]["memo"]["state"] == ""
    # Accept and Dismiss from the page, and Accept all.
    cams = cameras_of(incident)
    for text in ("One", "Two", "Three"):
        Event.objects.create(
            incident=incident,
            at=20.0,
            text=text,
            source="assistant",
            camera=cams["BWC2-1"],
            proposed=True,
            rests_on="a line",
        )
    one, two, three = Event.objects.filter(proposed=True).order_by("added")
    # While the run is going nothing is accepted or dismissed: the proposals
    # are still arriving, and one accepted now would be proposed again by
    # the cameras read after it (v1.63.2). The buttons show greyed.
    type(incident).objects.filter(pk=incident.pk).update(proposals_state="running")
    for held in (
        {"action": "event_accept", "event": str(one.pk)},
        {"action": "event_dismiss", "event": str(one.pk)},
        {"action": "event_accept_all"},
    ):
        answer = client.post(url + "/act", held)
        assert answer.status_code == 409
        assert answer.json()["error"] == (
            "The assistant is still proposing. Wait for it to finish."
        )
    assert Event.objects.filter(proposed=True).count() == 3
    type(incident).objects.filter(pk=incident.pk).update(proposals_state="done")
    got = client.post(
        url + "/act", {"action": "event_accept", "event": str(one.pk)}
    ).json()
    assert {e["text"]: e["proposed"] for e in got["state"]["events"]} == {
        "One": False,
        "Two": True,
        "Three": True,
    }
    got = client.post(
        url + "/act", {"action": "event_dismiss", "event": str(two.pk)}
    ).json()
    assert sorted(e["text"] for e in got["state"]["events"]) == ["One", "Three"]
    got = client.post(url + "/act", {"action": "event_accept_all"}).json()
    assert got["said"] == "1 event added."
    assert all(not e["proposed"] for e in got["state"]["events"])
    # An Event that is not a proposal cannot be accepted.
    assert (
        client.post(
            url + "/act", {"action": "event_accept", "event": str(one.pk)}
        ).status_code
        == 404
    )
    # With the memo off the tab is gone; with proposals off, the button.
    settings_store.set_to("incidents_memo", False)
    settings_store.set_to("incidents_propose", False)
    # The watch phrases alone would keep the button (Phase 7 chapter 2).
    settings_store.set_to("incidents_watch_phrases", "")
    page = client.get(url).content.decode()
    assert 'data-panel="memo"' not in page
    state = client.get(url + "/state").json()
    assert state["proposals"]["on"] is False and state["memo"]["on"] is False
    assert client.post(url + "/act", {"action": "propose"}).status_code == 400
    assert client.post(url + "/act", {"action": "memo"}).status_code == 400
    assert client.post(url + "/export/memo").status_code == 404
    # The case page says whether an incident has a memo.
    settings_store.set_to("incidents_memo", True)
    case_page = client.get(f"/case/{incident.case_id}?tab=incidents").content.decode()
    assert "no memo yet" in case_page


def test_a_numbered_label_does_not_show_under_a_tile(incident, person, client):
    signed_in(client, person)
    cams = cameras_of(incident)
    got = client.get(f"/incident-camera/{cams['BWC2-1'].pk}/lines").json()
    assert [one["speaker"] for one in got["segments"]] == ["", "Officer Hale"]
    assert got["segments"][0]["text"] == "Step out of the vehicle for me."


# The settings, the templates, the documents ----------------------------------------


def test_the_settings_templates_and_documents_exist(db, client):
    for key in (
        "incidents_case_chat",
        "incidents_propose",
        "incidents_events_answer_tokens",
        "incidents_events_time_seconds",
        "incidents_memo",
        "incidents_memo_answer_tokens",
        "incidents_memo_time_seconds",
    ):
        assert settings_store.definition(key).page == "incidents"
        assert settings_store.definition(key).group == "The assistant"
    assert settings_store.time_limit_seconds("incident_events") == 180
    assert settings_store.time_limit_seconds("incident_memo") == 600
    assert settings_store.incident_events_answer_cap() == 2000
    assert settings_store.incident_memo_answer_cap() == 4000
    for key, text, name in (
        (PromptTemplate.INCIDENT_EVENTS, prompts.INCIDENT_EVENTS, "Proposed events"),
        (PromptTemplate.INCIDENT_MEMO, prompts.INCIDENT_MEMO, "Incident memo"),
        (PromptTemplate.MEMO_SHEET, prompts.MEMO_SHEET, "Memo facts sheet"),
    ):
        row = PromptTemplate.named(key)
        assert row.text == text and not row.behind and row.name == name
        assert prompts.text_hash(text) in prompts.SHIPPED_HISTORY[f"prompt:{key}"]
    admin = User.objects.create_local_admin("adm", PASSWORD)
    signed_in(client, admin)
    templates = client.get("/panel/templates").content.decode()
    assert "Proposed events" in templates and "Incident memo" in templates
    assert "Memo facts sheet" in templates
    assert settings_store.definition("incidents_memo_sheet_answer_tokens").page == (
        "incidents"
    )
    assert settings_store.definition("documents_compare_left_out_most").page == (
        "documents"
    )
    panel = client.get("/panel/settings/incidents").content.decode()
    for name in (
        "Case chat knows the incidents",
        "Assistant proposes events",
        "Incident memo answer cap",
    ):
        assert name in panel
    glossary = (ROOT / "CONTEXT.md").read_text(encoding="utf-8")
    assert "**Incident memo**" in glossary
    guide = (ROOT / "docs" / "user-guide.md").read_text(encoding="utf-8")
    assert "Propose events" in guide and "Memo to Word" in guide
    assert "Speaker labels on an incident" in guide
    admin_guide = (ROOT / "docs" / "admin-guide.md").read_text(encoding="utf-8")
    assert "Case chat knows the incidents" in admin_guide
    catalogue = (ROOT / "docs" / "spec" / "ADMIN-SETTINGS-CATALOGUE.md").read_text(
        encoding="utf-8"
    )
    assert "Proposed events time limit" in catalogue and "v1.61.0" in catalogue
    changelog = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
    assert "## v1.61.0" in changelog
    assert assistant.PromptTemplate.DEFAULTS[PromptTemplate.INCIDENT_MEMO][0] == (
        "Incident memo"
    )
