"""v1.99.0: a note where nothing was said.

The rules checked here: a note written at a moment goes on the line being
spoken then, and is kept at the moment itself when nothing is being said,
never on a line a silence away; a second note at the same moment is the
same note; cleared, it is removed; it is an Event on the Chronology of a
synced camera at the exact moment, changed and removed from there as a
line's note is; it shows on the recording's page, the Notes tab, Search,
Find and the exports with notes, and never in a plain export, the captions
or what the assistant reads as the transcript; Process again leaves it be;
the audit row never holds the words; and only whoever may write a note on
the recording may write one.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from core import (
    case_search,
    chronology,
    exports,
    incidents,
    notes,
    prompts,
    settings_store,
)
from core.audit import Row
from core.cases import Case
from core.jobs import Job, JobState, Segment, Transcript
from core.models import LoginSession, User
from core.notes import MomentNote
from core.recordings import Batch, MediaState, Recording
from docx import Document

PASSWORD = "a-long-enough-password"
APP = Path(__file__).resolve().parent.parent
ROOT = APP.parent
WORDS = "The gun is lifted here, before anyone speaks of it."


@pytest.fixture(autouse=True)
def its_own_disk(tmp_path, settings):
    settings.DATA_DIR = tmp_path
    settings.SCRATCH_DIR = tmp_path / "scratch"
    settings.UPLOADS_DIR = tmp_path / "uploads"


@pytest.fixture(autouse=True)
def switched_on(db):
    settings_store.set_to("folder_management", True)
    settings_store.set_to("incidents", True)
    settings_store.set_to("assistant_available", True)


@pytest.fixture
def person(db):
    return User.objects.create_local_admin("asker", PASSWORD)


@pytest.fixture
def a_case(person):
    return Case.objects.create(owner=person, name="Traffic stop")


def signed_in(client, who):
    client.force_login(who)
    LoginSession.objects.create(user=who, session_key=client.session.session_key)
    return client


@pytest.fixture
def camera(person, a_case):
    """Nothing is said from 2:17 to 4:04."""
    recording = Recording.objects.create(
        batch=Batch.objects.create(user=person),
        user=person,
        case=a_case,
        title="BWC2-1",
        original_filename="BWC2-1.mp4",
        media_state=MediaState.READY,
        duration_seconds=600.0,
        playback_ready=True,
        stamp={
            "date": "06/07/2025",
            "time": "21:56:19",
            "camera": "BWC2-1",
            "at": 2.0,
            "checked": True,
        },
        probe={},
    )
    recording.folder.mkdir(parents=True, exist_ok=True)
    (recording.folder / "playback.mp4").write_bytes(b"not really media")
    transcript = Transcript.objects.create(recording=recording, language="en")
    for start, end, text in (
        (133.0, 135.0, "Hands where I can see them."),
        (136.0, 137.0, "Stay where you are."),
        (244.0, 245.0, "Bag it."),
        (246.0, 248.0, "Copy."),
    ):
        Segment.objects.create(
            transcript=transcript,
            start=start,
            end=end,
            text=text,
            speaker="Speaker 1",
            speaker_label="SPEAKER_1",
        )
    return recording


def line(camera, start):
    return Segment.objects.get(transcript=camera.transcript, start=start)


def write(client, camera, **sent):
    return client.post(
        f"/recording/{camera.pk}/note-at",
        data=json.dumps(sent),
        content_type="application/json",
    )


def test_the_line_being_spoken_and_the_silence(camera):
    transcript = camera.transcript
    # Inside a line, and within three seconds of its start or its end.
    assert notes.spoken_at(transcript, 136.5) == line(camera, 136.0)
    assert notes.spoken_at(transcript, 139.9) == line(camera, 136.0)
    assert notes.spoken_at(transcript, 241.5) == line(camera, 244.0)
    # Between two lines a moment falls inside neither: the nearer one.
    assert notes.spoken_at(transcript, 135.4) == line(camera, 133.0)
    # In the silence, nothing.
    assert notes.spoken_at(transcript, 210.0) is None
    assert notes.spoken_at(transcript, 140.1) is None
    assert notes.spoken_at(None, 10.0) is None
    assert notes.silence_at(transcript, 210.0) == {
        "from": 137.0,
        "to": 244.0,
        "seconds": 107,
    }
    # A short gap is not said, and a moment inside a line is in no silence.
    assert notes.silence_at(transcript, 135.5) is None
    assert notes.silence_at(transcript, 136.5) is None
    # After the last line, to the recording's end.
    assert notes.silence_at(transcript, 400.0)["to"] == 600.0


def test_the_preview_is_told_where_the_note_goes(camera, person, client):
    signed_in(client, person)
    got = client.get(f"/recording/{camera.pk}/around", {"t": "210"}).json()
    assert got["note_on"] is None
    assert got["note_at_url"] == f"/recording/{camera.pk}/note-at"
    assert got["silence"]["words"] == "Nothing is said for 1 minute 47 seconds."
    assert got["moment_notes"] == []
    got = client.get(f"/recording/{camera.pk}/around", {"t": "136.5"}).json()
    assert got["note_on"] == line(camera, 136.0).pk and got["silence"] is None


def test_a_note_in_a_silence_is_kept_at_the_moment(camera, person, client):
    signed_in(client, person)
    said = write(client, camera, at=210.0, note=WORDS).json()
    assert said["saved"] is True and said["on"] == "moment"
    kept = MomentNote.objects.get()
    assert kept.recording == camera and kept.at == 210.0 and kept.note == WORDS
    assert kept.note_by == person and kept.note_changed is not None
    assert said["moment"]["id"] == kept.pk and said["moment"]["clock"] == "00:03:30"
    # No line took it: the one a silence away least of all.
    assert not Segment.objects.exclude(note="").exists()
    row = Row.objects.filter(event="Note added").get()
    assert row.object_type == "moment" and row.object_label == "210.0"
    for one in Row.objects.all():
        assert "gun" not in json.dumps(one.details)
    # The Preview and the recording's page are told of it.
    got = client.get(f"/recording/{camera.pk}/around", {"t": "210"}).json()
    assert [one["id"] for one in got["moment_notes"]] == [kept.pk]
    page = client.get(f"/recording/{camera.pk}/segments").json()
    assert page["moment_notes"] == [
        {
            "id": kept.pk,
            "at": 210.0,
            "clock": "00:03:30",
            "note": WORDS,
            "note_by": person.shown_name,
            "note_on": notes.date_of(kept.note_changed),
        }
    ]
    assert len(page["segments"]) == 4
    # Written again at the same moment it is the same note, changed.
    said = write(client, camera, at=210.4, note="Changed.").json()
    assert said["moment"]["id"] == kept.pk and MomentNote.objects.count() == 1
    assert Row.objects.filter(event="Note changed").count() == 1
    # By its id, and cleared, it is removed.
    said = write(client, camera, id=kept.pk, note="  ").json()
    assert said["saved"] is True and said["moment"]["id"] is None
    assert not MomentNote.objects.exists()
    assert Row.objects.filter(event="Note removed").count() == 1
    # An empty note is never kept.
    assert write(client, camera, at=300.0, note="").json()["moment"]["id"] is None
    assert not MomentNote.objects.exists()


def test_a_note_where_a_line_is_spoken_goes_on_the_line(camera, person, client):
    signed_in(client, person)
    said = write(client, camera, at=136.5, note="He had stopped moving.").json()
    assert said["on"] == "line" and said["line"] == line(camera, 136.0).pk
    assert said["note"] == "He had stopped moving."
    assert line(camera, 136.0).note == "He had stopped moving."
    assert not MomentNote.objects.exists()


def test_what_is_refused(camera, person, client):
    signed_in(client, person)
    assert write(client, camera, note="No moment.").status_code == 400
    assert write(client, camera, at="soon", note="x").status_code == 400
    assert write(client, camera, at=-4, note="x").status_code == 400
    assert write(client, camera, at=9000, note="x").status_code == 400
    assert write(client, camera, id=12345, note="x").status_code == 404
    assert client.get(f"/recording/{camera.pk}/note-at").status_code == 405
    again = Batch.objects.create(user=person, is_reprocessing=True)
    Job.objects.create(recording=camera, batch=again, state=JobState.RUNNING)
    held = write(client, camera, at=210.0, note="x")
    assert held.status_code == 409
    assert held.json()["error"].startswith("This transcript is being replaced")
    other = User.objects.create(username="nobody", display_name="nobody")
    signed_in(client, other)
    assert write(client, camera, at=210.0, note="x").status_code == 404
    assert not MomentNote.objects.exists()


def test_it_is_an_event_at_the_exact_moment_on_a_synced_camera(
    camera, person, a_case, client
):
    incident = incidents.make(a_case, "Stop", [camera], by=person)
    placed = incident.cameras.get()
    incidents.place_by_hand(placed, 50.0, by=person)
    signed_in(client, person)
    write(client, camera, at=210.0, note=WORDS)
    kept = MomentNote.objects.get()
    event = chronology.Event.objects.get(incident=incident)
    assert event.moment_note == kept and event.segment is None
    assert event.at == 260.0 and event.text == WORDS and event.source == "note"
    assert event.is_note() and event.anchor() == kept
    assert event.cameras == [str(placed.pk)]
    row = chronology.events_json(incident)[0]
    assert row["at_a_moment"] is True and row["segment_id"] == ""
    assert row["recording_id"] == str(camera.pk)
    assert row["line_words"] == "nothing was said here"
    assert row["line_url"] == f"/recording/{camera.pk}?t=210.0&note=1"
    assert row["source_words"] == f"Note by {person.shown_name}, {placed.camera_id()}"
    # The camera moved, the event moves with it.
    incidents.place_by_hand(placed, 80.0, by=person)
    event.refresh_from_db()
    assert event.at == 290.0
    # Changed on the chronology, the note follows; the time is the note's.
    chronology.change(event, {"at": "5", "text": "Changed there."}, by=person)
    kept.refresh_from_db()
    event.refresh_from_db()
    assert kept.note == "Changed there." and event.text == "Changed there."
    assert event.at == 290.0
    changed = Row.objects.filter(event="Note changed").get()
    assert changed.details.get("where") == "chronology"
    # The assistant is told to respect it, and it is found once, as a note.
    assert notes.any_on_chronology(incident) is True
    found = case_search.find(incident, "changed")
    assert [(one["kind"], one["at"]) for one in found["hits"]] == [("note", 290.0)]
    # Removed on the chronology, the note goes.
    chronology.remove(event, by=person)
    assert not MomentNote.objects.exists()
    assert not chronology.Event.objects.filter(incident=incident).exists()


def test_a_camera_leaving_takes_the_event_and_leaves_the_note(camera, person, a_case):
    notes.note_at(camera, 210.0, WORDS, by=person)
    incident = incidents.make(a_case, "Stop", [camera], by=person)
    placed = incident.cameras.get()
    # Written before the recording joined, it appears when the camera is synced.
    assert chronology.Event.objects.filter(incident=incident).count() == 1
    incidents.place_by_hand(placed, 50.0, by=person)
    assert chronology.Event.objects.get(incident=incident).at == 260.0
    notes.drop_note_events(camera, incident)
    assert not chronology.Event.objects.filter(incident=incident).exists()
    assert MomentNote.objects.count() == 1


def test_it_shows_with_the_notes_and_never_as_a_line(camera, person, a_case, client):
    notes.note_at(camera, 210.0, WORDS, by=person)
    notes.set_note(line(camera, 136.0), "He had stopped moving.", by=person)
    # The Notes tab and its counts.
    listed = notes.of_case(a_case)
    assert listed["total"] == 2 and listed["lines"] == 2 and listed["moments"] == 1
    assert notes.count(a_case) == 2
    kept = [one for one in listed["rows"] if one["kind"] == "moment"][0]
    assert kept["when"] == "00:03:30" and kept["rests_on"] == "nothing was said here"
    assert kept["url"] == f"/recording/{camera.pk}?t=210.0&note=1"
    signed_in(client, person)
    page = client.get(f"/case/{a_case.pk}?tab=notes").content.decode()
    assert "at a moment</span>" in page and "nothing was said here" in page
    document = Document(io.BytesIO(notes.word(a_case, "asker")))
    told = "\n".join(one.text for one in document.paragraphs)
    assert (
        "2 on lines of transcripts (1 of them at a moment where nothing was said)"
        in (told)
    )
    # Search finds it as a note, with a Preview.
    got = case_search.search(a_case, "lifted")
    assert {one["key"]: one["count"] for one in got["kinds"]} == {"notes": 1}
    hit = got["groups"][0]["hits"][0]
    assert hit["when"] == "00:03:30" and "<mark>lifted</mark>" in str(hit["under"])
    assert hit["preview"] == {"recording": str(camera.pk), "at": 210.0}
    # The exports with notes print it at its own time, between the lines.
    text = exports.plain_text(camera, with_notes=True).splitlines()
    at = [n for n, one in enumerate(text) if "nothing was said here" in one]
    assert len(at) == 1
    assert text[at[0]].strip().startswith("[00:03:30] Note (")
    assert text[at[0]].strip().endswith("; nothing was said here): " + WORDS)
    assert "Stay where you are." in text[at[0] - 2] and "Bag it." in text[at[0] + 1]
    word = Document(io.BytesIO(exports.word(camera, "asker", with_notes=True)))
    paragraphs = [one.text for one in word.paragraphs]
    assert any(one.startswith("[00:03:30] Note (") for one in paragraphs)
    assert "With the office's notes: 2." in paragraphs
    # A plain export, the captions and the assistant's reading never carry it.
    assert WORDS not in exports.plain_text(camera)
    plain = Document(io.BytesIO(exports.word(camera, "asker")))
    assert not any(WORDS in one.text for one in plain.paragraphs)
    assert WORDS not in exports.srt(camera)
    assert WORDS not in prompts.render(prompts.lines_of(camera.transcript))
    # The case chat is told it as the office's note, and that nothing was said.
    block = notes.recording_block(camera)
    assert "[00:03:30] (nothing was said here) " in block and WORDS in block
    assert block.index("00:02:16") < block.index("00:03:30")
    assert notes.any_on([camera]) is True


def test_process_again_leaves_it_be(camera, person, a_case):
    incident = incidents.make(a_case, "Stop", [camera], by=person)
    incidents.place_by_hand(incident.cameras.get(), 50.0, by=person)
    notes.note_at(camera, 210.0, WORDS, by=person)
    kept = notes.remember(camera)
    assert kept == []
    camera.transcript.delete()
    fresh = Transcript.objects.create(recording=camera, language="en")
    Segment.objects.create(
        transcript=fresh, start=136.0, end=137.0, text="Stay.", speaker="Speaker 1"
    )
    notes.carry(kept, fresh)
    assert MomentNote.objects.get().note == WORDS
    event = chronology.Event.objects.get(incident=incident)
    assert event.at == 260.0 and event.text == WORDS


def test_the_pages_and_the_words(camera, person, client):
    script = (APP / "static" / "preview.js").read_text(encoding="utf-8")
    assert (
        "Nothing is being said at this moment, so the note is kept at the moment"
        in (script)
    )
    assert "told.note_at_url" in script and "data-preview-moment" in script
    page = (APP / "static" / "viewer.js").read_text(encoding="utf-8")
    assert "body.moment_notes" in page and "nothing was said here" in page
    assert "—" not in script + page
    glossary = (ROOT / "CONTEXT.md").read_text(encoding="utf-8")
    guide = (ROOT / "docs" / "user-guide.md").read_text(encoding="utf-8")
    assert "kept at the moment itself" in glossary
    assert "the note is kept at the moment itself" in guide
