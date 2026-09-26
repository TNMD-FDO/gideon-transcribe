"""The Timeline report (Phase 8 chapter 14).

The rules checked here: the Export menu offers the report when the setting
is on and the chronology has an event; the document carries the cover, the
cameras, the band, the spells and one entry per event with its still, the
words heard around it on its camera, what the camera showed, the note and
the source; stills follow their switch and their limit; the export writes
the chronology's audit row with the kind report; the words are in the
glossary, the guide and the specification.
"""

from __future__ import annotations

# ruff: noqa: F811 - the fixtures are imported by name, as pytest wants them
import io
from pathlib import Path

import pytest
from core import chronology, media, settings_store
from core.audit import Row
from tests.test_incident_assistant import (  # noqa: F401 - fixtures
    a_case,
    incident,
    its_own_disk,
    no_tasks,
    person,
    signed_in,
    switched_on,
)

ROOT = Path(__file__).resolve().parent.parent.parent


def fake_frames(monkeypatch):
    """A frame cutter that draws a small grey picture instead of calling ffmpeg."""
    cut = []

    def grab(source, folder, times, *, height, timeout=0):
        from PIL import Image

        out = []
        for number, at in enumerate(times, start=1):
            target = folder / f"frame-{number}.jpg"
            Image.new("RGB", (32, 18), "grey").save(target)
            out.append(target)
            cut.append((Path(source).name, round(at, 1), height))
        return out

    monkeypatch.setattr(media, "grab_frames", grab)
    return cut


def with_events(incident, person):
    first = incident.cameras.get(recording__title="first")
    second = incident.cameras.get(recording__title="second")
    chronology.add(
        incident,
        {
            "at": "30",
            "text": "Hands asked for",
            "camera": str(first.pk),
            "source": "words",
        },
        by=person,
    )
    chronology.add(
        incident,
        {"at": "286", "text": "Told to stay in the car", "camera": str(second.pk)},
        by=person,
    )
    return first, second


def read_document(content: bytes):
    from docx import Document

    document = Document(io.BytesIO(content))
    words = "\n".join(one.text for one in document.paragraphs)
    cells = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return document, words, cells


@pytest.mark.django_db
def test_the_report_prints_the_chronology_with_stills_and_words(
    client, incident, person, monkeypatch
):
    cut = fake_frames(monkeypatch)
    first, second = with_events(incident, person)
    signed_in(client, person)
    base = f"/case/{incident.case_id}/incident/{incident.pk}"
    page = client.get(base).content.decode()
    assert f'href="{base}/export/report"' in page and "Timeline report (Word)" in page

    answer = client.get(f"{base}/export/report")
    assert answer.status_code == 200
    assert answer["Content-Disposition"].endswith(
        'filename="Timeline report - Stop.docx"'
    )
    document, words, cells = read_document(answer.content)
    assert "Timeline report: Stop" in words
    assert "Each bar is one camera's recording on the clock." in words
    # The cameras' table, and a spell per stretch: 30 s and 286 s are 256 s
    # apart, so one spell of two events.
    assert "BWC2-1" in cells and "BWC2-2" in cells
    assert "21:56:47 to 22:01:03" in words and "2 events" in words
    # Each entry: the line, the camera, the words heard around it, the camera's
    # own description from the Digest, and the source.
    assert "Hands asked for" in cells and "Told to stay in the car" in cells
    assert "Said on BWC2-1" in cells and "Hands where I can see them." in cells
    assert "Camera: " in cells and "grey jacket" in cells
    assert "Said on BWC2-2" in cells and "Stay in the car, please." in cells
    assert "Added by asker" in cells
    # The band, then one still per event, cut at the setting's height from
    # each event's own camera at its moment.
    assert len(document.inline_shapes) == 3
    # Cut a few at a time, so in no fixed order.
    assert sorted(cut) == [("playback.mp4", 5.0, 360), ("playback.mp4", 30.0, 360)]
    # The audit row is the chronology's, with the kind report.
    row = Row.objects.get(event="chronology exported")
    assert row.details["kind"] == "report" and row.object_label == "Stop"


@pytest.mark.django_db
def test_stills_follow_their_switch_and_their_limit(
    client, incident, person, monkeypatch
):
    cut = fake_frames(monkeypatch)
    with_events(incident, person)
    signed_in(client, person)
    base = f"/case/{incident.case_id}/incident/{incident.pk}"
    settings_store.set_to("incidents_report_stills_most", 1)
    document, words, _ = read_document(client.get(f"{base}/export/report").content)
    assert len(document.inline_shapes) == 2 and len(cut) == 1
    assert "1 event prints without a still" in words
    settings_store.set_to("incidents_report_stills", False)
    document, words, _ = read_document(client.get(f"{base}/export/report").content)
    assert len(document.inline_shapes) == 1 and len(cut) == 1
    assert "without a still" not in words
    # No words around the event when the window is 0.
    settings_store.set_to("incidents_report_words_seconds", 0)
    _, _, cells = read_document(client.get(f"{base}/export/report").content)
    assert "Said on" not in cells and "Camera: " not in cells


@pytest.mark.django_db
def test_the_report_is_greyed_without_events_and_gone_when_off(
    client, incident, person, monkeypatch
):
    signed_in(client, person)
    base = f"/case/{incident.case_id}/incident/{incident.pk}"
    page = client.get(base).content.decode()
    assert "Timeline report: add an event first" in page
    assert f'href="{base}/export/report"' not in page
    # Without an event the document still opens and says so.
    _, words, _ = read_document(client.get(f"{base}/export/report").content)
    assert "No events on the chronology yet." in words
    settings_store.set_to("incidents_report", False)
    page = client.get(base).content.decode()
    assert "Timeline report" not in page
    assert client.get(f"{base}/export/report").status_code == 404


def test_the_words_are_in_the_glossary_the_guide_and_the_spec():
    glossary = (ROOT / "CONTEXT.md").read_text(encoding="utf-8")
    assert "**Timeline report**:" in glossary
    guide = (ROOT / "docs" / "user-guide.md").read_text(encoding="utf-8")
    assert "**The timeline report.**" in guide and "Timeline report (Word)" in guide
    spec = (ROOT / "docs" / "spec" / "SPEC-PHASE-8.md").read_text(encoding="utf-8")
    assert "## 14. The Timeline report" in spec
    assert "## 15. Deferred and ruled out" in spec
    catalogue = (ROOT / "docs" / "spec" / "ADMIN-SETTINGS-CATALOGUE.md").read_text(
        encoding="utf-8"
    )
    for name in (
        "| Timeline report |",
        "| Stills in the report |",
        "| Still height |",
        "| Most stills in a report |",
        "| Words around an event |",
    ):
        assert name in catalogue, name
