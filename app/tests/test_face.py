"""The face: the sign-in page, the office logo and name, and the exports' head.

The logo is a row in the database, served to a page nobody has signed in to
yet, recognised by its bytes, refused when it is not an image or too large,
and stamped on a Word export only when the Admin's toggle says so.
"""

from __future__ import annotations

import io
import zlib

import pytest
from core import branding, exports, settings_store
from core.audit import Row
from core.jobs import Segment, Transcript
from core.models import LoginSession, User
from core.recordings import Batch, MediaState, Recording

PASSWORD = "a-long-enough-password"


def a_png(width: int = 2, height: int = 2) -> bytes:
    """A real, tiny PNG, made here so no image file lives in the repository."""

    def chunk(kind: bytes, body: bytes) -> bytes:
        crc = zlib.crc32(kind + body) & 0xFFFFFFFF
        return len(body).to_bytes(4, "big") + kind + body + crc.to_bytes(4, "big")

    raw = b"".join(b"\x00" + b"\x10\x5e\x86" * width for _ in range(height))
    header = (
        width.to_bytes(4, "big") + height.to_bytes(4, "big") + b"\x08\x02\x00\x00\x00"
    )
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def test_a_logo_is_known_by_its_first_bytes():
    assert branding.kind_of(a_png()[:12]) == "image/png"
    assert branding.kind_of(b"\xff\xd8\xff\xe0rest") == "image/jpeg"
    assert branding.kind_of(b"<svg xmlns") is None
    assert branding.kind_of(b"GIF89a") is None


@pytest.fixture
def admin(db):
    return User.objects.create_local_admin("the-admin", PASSWORD)


def signed_in(client, who):
    client.force_login(who)
    LoginSession.objects.create(user=who, session_key=client.session.session_key)
    return client


@pytest.mark.django_db
def test_the_sign_in_page_wears_the_apps_own_face_until_a_logo_is_uploaded(
    admin, client
):
    page = client.get("/sign-in").content.decode()
    assert 'class="mark big"' in page and "Gideon Transcribe" in page
    assert "/branding/logo" not in page
    assert "<title>Sign in &middot; Gideon Transcribe</title>" in page
    assert client.get("/branding/logo").status_code == 404

    signed_in(client, admin)
    settings_store.set_to("office_name", "Office of the Federal Public Defender")
    sent = io.BytesIO(a_png())
    sent.name = "seal.png"
    answer = client.post("/panel/appearance/logo", {"logo": sent})
    assert answer.status_code == 302 and answer["Location"].endswith("/appearance")
    row = Row.objects.get(category="admin", event="Logo uploaded")
    assert row.details["kind"] == "image/png" and row.details["bytes"] > 0

    client.logout()
    page = client.get("/sign-in").content.decode()
    assert '<img src="/branding/logo"' in page
    assert "Office of the Federal Public Defender" in page
    served = client.get("/branding/logo")
    assert served.status_code == 200 and served["Content-Type"] == "image/png"
    assert served.content == a_png()
    assert served["Cache-Control"] == "public, max-age=300"
    again = client.get("/branding/logo", HTTP_IF_NONE_MATCH=served["ETag"])
    assert again.status_code == 304


@pytest.mark.django_db
def test_an_upload_that_is_not_a_small_image_is_refused_and_said(admin, client):
    signed_in(client, admin)
    text = io.BytesIO(b"<svg xmlns='http://www.w3.org/2000/svg'></svg>")
    text.name = "logo.svg"
    client.post("/panel/appearance/logo", {"logo": text})
    page = client.get("/panel/settings/appearance").content.decode()
    assert "not a PNG or JPEG image" in page
    assert branding.current() is None

    big = io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"\x00" * (branding.LARGEST + 1))
    big.name = "big.png"
    client.post("/panel/appearance/logo", {"logo": big})
    page = client.get("/panel/settings/appearance").content.decode()
    assert "larger than 2 MB" in page
    assert not Row.objects.filter(event="Logo uploaded").exists()


@pytest.mark.django_db
def test_the_logo_can_be_replaced_and_removed(admin, client):
    signed_in(client, admin)
    first = io.BytesIO(a_png(2, 2))
    first.name = "one.png"
    client.post("/panel/appearance/logo", {"logo": first})
    second = io.BytesIO(a_png(3, 1))
    second.name = "two.png"
    client.post("/panel/appearance/logo", {"logo": second})
    assert branding.Branding.objects.count() == 1
    assert client.get("/branding/logo").content == a_png(3, 1)
    page = client.get("/panel/settings/appearance").content.decode()
    assert "Replace logo" in page and "Remove logo" in page

    client.post("/panel/appearance/logo", {"action": "remove"})
    assert branding.current() is None
    assert Row.objects.filter(category="admin", event="Logo removed").count() == 1
    assert client.get("/branding/logo").status_code == 404


@pytest.mark.django_db
def test_the_export_carries_the_logo_only_when_the_toggle_says_so(
    admin, settings, tmp_path
):
    from docx import Document

    settings.DATA_DIR = tmp_path
    recording = Recording.objects.create(
        batch=Batch.objects.create(user=admin),
        user=admin,
        title="Interview",
        original_filename="interview.m4a",
        media_state=MediaState.READY,
        duration_seconds=60,
    )
    transcript = Transcript.objects.create(recording=recording, language="en")
    Segment.objects.create(
        transcript=transcript, start=0, end=3, text="Hello.", speaker="A"
    )
    branding.upload(a_png(), by=admin)
    settings_store.set_to("office_name", "The Office")

    plain = Document(io.BytesIO(exports.word(recording, "the-admin")))
    assert len(plain.inline_shapes) == 0
    assert plain.paragraphs[0].text == "The Office"

    settings_store.set_to("logo_on_exports", True)
    stamped = Document(io.BytesIO(exports.word(recording, "the-admin")))
    assert len(stamped.inline_shapes) == 1
    assert stamped.paragraphs[1].text == "The Office"


def test_the_appearance_page_and_its_two_rows_exist():
    assert settings_store.APPEARANCE in dict(settings_store.PAGES)
    assert settings_store.DEFINITIONS["logo_on_exports"].default is False
    assert settings_store.DEFINITIONS["office_name"].page == settings_store.APPEARANCE
