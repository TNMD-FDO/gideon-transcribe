"""The Case Chat: a Chat whose ground is every Transcript in a Case.

The packing and Citation rules run without a database. The rest runs against a
fake engine: what one question sends (the People line, one header per
Transcript, the template), what it stores, how a Citation opens the right
Recording later, the parts and the combining call, the hours ceiling, the
Recordings skipped, the tab's gates, the export, and the audit row that never
carries a word of the question.
"""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from core import case_chat, engine, exports, people, prompts, settings_store
from core.assistant import PromptTemplate
from core.audit import Row
from core.case_chat import CaseChat, CaseChatTurn
from core.cases import Case
from core.jobs import Segment, Transcript
from core.models import LoginSession, User
from core.recordings import Batch, MediaState, Recording

PASSWORD = "a-long-enough-password"


# Without a database ------------------------------------------------------------------


def test_whole_transcripts_are_packed_in_order_and_never_cut():
    big = "x" * (case_chat.READING_TOKENS * prompts.CHARS_PER_TOKEN)
    small = "y" * 400
    assert case_chat.pack([(1, small), (2, small), (3, small)]) == [[1, 2, 3]]
    # A Transcript larger than a Reading is alone; order is kept, so the small
    # one after it starts a Reading of its own rather than joining an earlier one.
    assert case_chat.pack([(1, small), (2, big), (3, small)]) == [[1], [2], [3]]


def test_a_citation_names_a_recording_and_a_real_line():
    starts = {1: {"recording": "r1", "starts": {765: 765.4}}}
    found = case_chat.citations(
        "At [Recording 1, 00:12:45] yes; [Recording 1, 00:12:46] no; "
        "[Recording 2, 00:12:45] neither.",
        starts,
    )
    assert found == {"[Recording 1, 00:12:45]": {"recording": "r1", "seconds": 765.4}}


def test_a_time_inside_the_recording_links_where_no_line_starts():
    """v1.82.0: what the camera showed has times of its own in the Digest, so
    a cited second inside the recording links to that second; past the end,
    or a number not read, stays plain text."""
    starts = {1: {"recording": "r1", "starts": {765: 765.4}, "length": 900.0}}
    found = case_chat.citations(
        "Said at [Recording 1, 00:12:45]; seen at [Recording 1, 00:12:46]; "
        "past the end [Recording 1, 00:15:01]; not read [Recording 2, 00:00:10].",
        starts,
    )
    assert found == {
        "[Recording 1, 00:12:45]": {"recording": "r1", "seconds": 765.4},
        "[Recording 1, 00:12:46]": {"recording": "r1", "seconds": 766.0},
    }


def test_a_short_time_after_the_recordings_name_takes_the_full_shape():
    """v1.82.0: the engine names the recording once and writes the time alone;
    the nearest name before it in the paragraph is the recording, a full
    citation counts as a name, and a paragraph with no name is left alone."""
    named = case_chat.with_recordings_named
    assert (
        named("In Recording 15, the camera shows a door [00:03:40]. Later [00:04:21].")
        == "In Recording 15, the camera shows a door [Recording 15, 00:03:40]. "
        "Later [Recording 15, 00:04:21]."
    )
    assert (
        named("Recording 1 says so [00:00:05]; Recording 2 too [00:00:09].")
        == "Recording 1 says so [Recording 1, 00:00:05]; "
        "Recording 2 too [Recording 2, 00:00:09]."
    )
    assert (
        named("Blue [Recording 3, 00:12:45], and again [00:12:50].")
        == "Blue [Recording 3, 00:12:45], and again [Recording 3, 00:12:50]."
    )
    # A new paragraph starts afresh; "Recordings 4 and 5" names none.
    assert named("Recording 4 here.\nAlone [00:00:01].") == (
        "Recording 4 here.\nAlone [00:00:01]."
    )
    assert named("Recordings 4 and 5 agree [00:00:01].") == (
        "Recordings 4 and 5 agree [00:00:01]."
    )
    assert named("No citation at all.") == "No citation at all."


def test_the_not_read_line_counts_each_reason():
    assert case_chat.not_read_line([]) == ""
    assert (
        case_chat.not_read_line(
            [{"why": "still transcribing"}, {"why": "still transcribing"}]
        )
        == "2 recordings were not read: still transcribing."
    )


# Against a fake engine ----------------------------------------------------------------


@pytest.fixture(autouse=True)
def its_own_disk(tmp_path, settings):
    settings.DATA_DIR = tmp_path
    settings.SCRATCH_DIR = tmp_path / "scratch"
    settings.UPLOADS_DIR = tmp_path / "uploads"
    return tmp_path


@pytest.fixture
def owner(db):
    return User.objects.create_local_admin("case-owner", PASSWORD)


@pytest.fixture
def a_case(owner):
    settings_store.set_to("folder_management", True)
    settings_store.set_to("assistant_available", True)
    return Case.objects.create(owner=owner, name="Ramirez")


def a_recording(owner, case, title, *, hours=0.25, transcript=True, translate=False):
    recording = Recording.objects.create(
        batch=Batch.objects.create(user=owner),
        user=owner,
        case=case,
        title=title,
        original_filename=f"{title}.m4a",
        media_state=MediaState.READY,
        duration_seconds=hours * 3600,
        diarize=True,
        recording_type="Jail call",
    )
    if transcript:
        made = Transcript.objects.create(
            recording=recording,
            language="es" if translate else "en",
            task_run="translate" if translate else "transcribe",
        )
        for start, speaker, text in (
            (0.0, "Speaker 1", f"This is {title}, good morning."),
            (765.0, "Speaker 2", "The car was blue."),
        ):
            Segment.objects.create(
                transcript=made,
                start=start,
                end=start + 5,
                text=text,
                speaker=speaker,
                speaker_label=speaker.upper().replace(" ", "_"),
            )
    return recording


def signed_in(client, who):
    client.force_login(who)
    LoginSession.objects.create(user=who, session_key=client.session.session_key)
    return client


def reachable(monkeypatch, answers):
    """An engine that is up and answers each call with the next text."""
    monkeypatch.setattr(engine, "is_reachable", lambda: True)
    monkeypatch.setattr(engine, "address", lambda: "http://gideon-generator:8000/v1")
    asked = []
    queue = list(answers)

    def complete(messages, **options):
        asked.append({"messages": messages, **options})
        return {
            "text": queue.pop(0) if len(queue) > 1 else queue[0],
            "finish_reason": "stop",
            "input_tokens": 100,
            "output_tokens": 10,
            "model": "the-model",
        }

    monkeypatch.setattr(engine, "complete", complete)
    return asked


def swallow_defer(monkeypatch):
    from core import tasks

    monkeypatch.setattr(tasks.answer_case_turn, "defer", lambda **fields: None)


@pytest.mark.django_db
def test_one_reading_sends_the_people_the_headers_and_the_template(
    owner, a_case, client, monkeypatch
):
    swallow_defer(monkeypatch)
    first = a_recording(owner, a_case, "Jail call 1")
    second = a_recording(owner, a_case, "Jail call 2")
    people.join_or_create(a_case, "Maria Lopez", by=owner, how="test", role="Client")
    asked = reachable(
        monkeypatch,
        ["The car was blue [Recording 2, 00:12:45], twice [Recording 1, 00:12:45]."],
    )
    signed_in(client, owner)

    made = client.post(f"/case/{a_case.pk}/chats")
    chat_id = made.json()["id"]
    answer = client.post(
        f"/case-chat/{chat_id}/ask",
        data=json.dumps({"question": "What colour was the car?"}),
        content_type="application/json",
    )
    assert answer.status_code == 200
    case_chat.answer_case_turn(answer.json()["id"])

    assert len(asked) == 1
    call = asked[0]
    assert call["max_completion_tokens"] == case_chat.PART_CAP
    system = call["messages"][0]["content"]
    assert prompts.CASE_CHAT in system and prompts.CASE_CHAT_FORMAT in system
    user = call["messages"][-1]["content"]
    assert user.startswith("People in this case: Maria Lopez (Client)")
    assert "Recording 1 of 2: Jail call 1, Jail call, uploaded" in user
    assert "Recording 2 of 2: Jail call 2" in user
    assert "machine transcription in English, speakers separated automatically" in user
    assert "[1] [00:00:00] Speaker 1: This is Jail call 1" in user
    assert user.rstrip().endswith("What colour was the car?")

    turn = CaseChatTurn.objects.get()
    assert turn.state == "done" and turn.parts == 0
    assert [one["title"] for one in turn.readings] == ["Jail call 1", "Jail call 2"]
    assert turn.citations["[Recording 2, 00:12:45]"] == {
        "recording": str(second.pk),
        "seconds": 765.0,
    }

    state = client.get(f"/case/{a_case.pk}/chat").json()
    chat = state["chats"][0]
    assert chat["name"] == "What colour was the car?"
    cited = chat["turns"][0]["citations"]["[Recording 1, 00:12:45]"]
    assert cited["title"] == "Jail call 1" and cited["clock"] == "00:12:45"
    assert cited["href"] == f"/recording/{first.pk}?t=765.0"
    # The line the pill points to, for its hover; and when things happened.
    assert cited["line"] == "Speaker 2: The car was blue."
    assert chat["started"] and chat["turns"][0]["asked_at"]
    assert chat["turns"][0]["answered_at"] and chat["turns"][0]["parts"] == 0

    row = Row.objects.get(category="llm", event="AI assistant call")
    assert row.details["feature"] == "case_chat_turn"
    assert row.details["transcripts_read"] == 2 and row.details["readings"] == 1
    assert row.details["templates"] == "ground-rules v1; case-chat v1"
    assert row.object_label == "Ramirez"
    assert "car" not in json.dumps(row.details)


@pytest.mark.django_db
def test_a_case_read_in_parts_asks_each_part_then_combines(
    owner, a_case, client, monkeypatch
):
    swallow_defer(monkeypatch)
    for n in range(3):
        a_recording(owner, a_case, f"Call {n + 1}")
    # Each Transcript fills a Reading on its own, so three parts and one combining call.
    monkeypatch.setattr(settings_store, "reading_tokens", lambda: 10)
    # The Readings run two at a time, so the fake answers by what it was
    # asked, not by the order the calls arrive in.
    by_content = {
        "Recording 1 of 3": "Part one says blue [Recording 1, 00:12:45].",
        "Recording 2 of 3": "Part two says nothing.",
        "Recording 3 of 3": "Part three says blue [Recording 3, 00:12:45].",
        prompts.COMBINING: (
            "Blue, in two calls [Recording 1, 00:12:45] [Recording 3, 00:12:45]."
        ),
    }
    asked = reachable(monkeypatch, [""])

    def by_the_question(messages, **options):
        asked.append({"messages": messages, **options})
        user = messages[-1]["content"]
        text = next(answer for key, answer in by_content.items() if key in user)
        return {
            "text": text,
            "finish_reason": "stop",
            "input_tokens": 100,
            "output_tokens": 10,
            "model": "the-model",
        }

    monkeypatch.setattr(engine, "complete", by_the_question)
    signed_in(client, owner)
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="Colour?")
    case_chat.answer_case_turn(turn.pk)

    assert len(asked) == 4
    combining = asked[-1]
    assert combining["max_completion_tokens"] == case_chat.COMBINED_CAP
    assert (
        combining["messages"][0]["content"] == PromptTemplate.named("ground_rules").text
    )
    user = combining["messages"][-1]["content"]
    assert user.startswith(prompts.COMBINING)
    assert "Part 1 of 3:\nPart one says blue" in user and "Part 3 of 3:" in user
    # The combine joins the parts in time order and never dresses coverage up
    # as disagreement (v1.87.0).
    assert "not a disagreement" in prompts.COMBINING
    # Each part was told which recordings it holds (v1.87.0).
    parts_asked = [one["messages"][-1]["content"] for one in asked[:3]]
    assert all("This is part " in one for one in parts_asked)
    assert any(
        one.startswith(
            "This is part 2 of 3 of the case's recordings: it holds Recording 2 of 3."
        )
        for one in parts_asked
    )
    turn.refresh_from_db()
    assert turn.parts == 3 and turn.state == "done"
    # Each part read was counted as it came back, for the page's wait line.
    assert turn.parts_done == 3
    assert turn.answer.startswith("Blue, in two calls")
    assert set(turn.citations) == {"[Recording 1, 00:12:45]", "[Recording 3, 00:12:45]"}
    row = Row.objects.get(category="llm", event="AI assistant call")
    assert row.details["readings"] == 3 and row.details["input_tokens"] == 400


def a_synced_incident(owner, a_case, *, digest_lines=2):
    """A call and an incident of two synced cameras: the first with a Digest,
    the second with its words alone; the second starts 281 s after the first."""
    from core import incidents
    from core.assistant import DigestPart
    from tests.test_incident_assistant import stamp, video

    settings_store.set_to("incidents", True)
    settings_store.set_to("moments_available", True)
    settings_store.set_to("digests_available", True)
    a_recording(owner, a_case, "Jail call 1")
    first = video(
        owner,
        a_case,
        "first",
        stamp=stamp("21:56:19", "BWC2-1"),
        lines=(
            (10.0, "Speaker 1", "Step out of the vehicle for me."),
            (30.0, "Officer Hale", "Hands where I can see them."),
        ),
    )
    second = video(
        owner,
        a_case,
        "second",
        stamp=stamp("22:01:00", "BWC2-2"),
        lines=((5.0, "Speaker 4", "Stay in the car, please."),),
    )
    lines = [
        '1. [00:00:10]-[00:00:14] (said) Speaker 1 said "Step out of the vehicle '
        'for me."',
        "2. [00:00:30]-[00:00:34] (both) Officer Hale asked for hands in view; "
        "the camera showed a grey jacket.",
    ]
    for n in range(2, digest_lines):
        lines.append(
            f"{n + 1}. [00:01:{n:02d}]-[00:01:{n + 3:02d}] (seen) A long description "
            "of the roadside and the two officers standing by the car, at length."
        )
    DigestPart.objects.create(
        transcript=first.transcript, number=1, text="\n".join(lines), made_at=None
    )
    return incidents.make(a_case, "Stop", [first, second], by=owner)


@pytest.mark.django_db
def test_an_incidents_cameras_are_read_as_one_record(
    owner, a_case, client, monkeypatch
):
    """v1.87.0: the synced cameras of an incident are one item of a Reading,
    the incident's record in time order by the cameras' clock, each line
    cited as its recording and its own time, as the memo and the incident
    page's Gideon read them; a camera's own transcript and Digest are not
    sent as well."""
    swallow_defer(monkeypatch)
    a_synced_incident(owner, a_case)
    asked = reachable(
        monkeypatch,
        ["The stop began [Recording 2, 00:00:10]; then [Recording 3, 00:00:05]."],
    )
    signed_in(client, owner)
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="What happened?")
    case_chat.answer_case_turn(turn.pk)

    assert len(asked) == 1
    user = asked[0]["messages"][-1]["content"]
    # Every recording keeps its header; the incident's two cameras are one item.
    for n in (1, 2, 3):
        assert f"Recording {n} of 3" in user
    assert "The record of the incident Stop: its 2 synced cameras" in user
    assert "each camera's Digest" in user
    # The first camera's Digest lines, cited as the recording and its own
    # time, and never its transcript lines as well.
    assert "[Recording 2, 00:00:10]-[00:00:14] (said) Speaker 1 said" in user
    assert "[Recording 2, 00:00:10] Step out" not in user
    # The second camera has no Digest, so its words, without its numbered label.
    assert "[Recording 3, 00:00:05] Stay in the car, please." in user
    # In time order by the clock: the second camera starts 281 s later.
    assert user.index("[Recording 2, 00:00:30]") < user.index("[Recording 3, 00:00:05]")
    turn.refresh_from_db()
    assert turn.state == "done" and turn.parts == 0
    assert set(turn.citations) == {"[Recording 2, 00:00:10]", "[Recording 3, 00:00:05]"}
    assert turn.readings[1]["incident"] == "Stop" and turn.readings[1]["digest"] is True
    assert turn.readings[2]["incident"] == "Stop"
    # The second camera has no Digest, so the answer says it was read from
    # its words; the first, read from its Digest in the record, is not named.
    assert turn.answer.startswith("Recording 3 was read from the transcript alone.")


@pytest.mark.django_db
def test_a_cameras_words_go_alone_when_the_record_would_not_fit_a_reading(
    owner, a_case, client, monkeypatch
):
    """v1.87.0: when the incident's record is larger than a Reading, the
    longest unpinned camera is read by its words alone, as the Sitting does,
    and the answer says so; the cameras are never sent separately."""
    swallow_defer(monkeypatch)
    a_synced_incident(owner, a_case, digest_lines=40)
    # The Digest alone is over this; the words of both cameras are well under.
    monkeypatch.setattr(settings_store, "reading_tokens", lambda: 600)
    asked = reachable(monkeypatch, ["Hands [Recording 2, 00:00:30]."])
    signed_in(client, owner)
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="Hands?")
    case_chat.answer_case_turn(turn.pk)

    assert len(asked) == 1
    user = asked[0]["messages"][-1]["content"]
    assert "the words alone of BWC2-1" in user
    assert "[Recording 2, 00:00:30] Officer Hale: Hands where I can see them." in user
    assert "(seen) A long description" not in user
    turn.refresh_from_db()
    assert turn.state == "done" and turn.parts == 0
    assert (
        turn.readings[1]["digest"] is False and turn.readings[2]["incident"] == "Stop"
    )
    # The first camera's words went alone; the second never had a Digest.
    assert turn.answer.startswith(
        "Recording 2 and Recording 3 were read from the transcript alone."
    )
    assert turn.citations == {
        "[Recording 2, 00:00:30]": {
            "recording": turn.readings[1]["recording"],
            "seconds": 30.0,
        }
    }


def test_a_calls_time_limit_grows_with_what_it_sends(monkeypatch):
    """v1.87.0: one second for every 500 tokens sent, the text at four
    characters a token, on top of the feature's setting; a picture counts
    nothing."""
    messages = [
        {"role": "system", "content": "x" * 4000},
        {
            "role": "user",
            "content": [
                {"type": "text", "text": "y" * 2000},
                {"type": "image_url", "image_url": {"url": "data:..."}},
            ],
        },
    ]
    assert engine.reading_allowance(messages) == 3.0
    given = []

    class Completions:
        def create(self, **request):
            class Choice:
                finish_reason = "stop"

                class message:
                    content = "ready"

            class Answer:
                choices = [Choice()]
                usage = None
                model = "the-model"

            return Answer()

    class Client:
        class chat:
            completions = Completions()

    monkeypatch.setattr(
        engine, "client", lambda timeout: given.append(timeout) or Client()
    )
    monkeypatch.setattr(engine, "model_name", lambda: "the-model")
    engine.complete(
        messages,
        max_completion_tokens=8,
        temperature=0,
        top_p=1,
        thinking=False,
        timeout=120,
    )
    assert given == [123.0]


@pytest.mark.django_db
def test_recordings_not_ready_are_skipped_and_named(owner, a_case, monkeypatch):
    a_recording(owner, a_case, "Ready call")
    a_recording(owner, a_case, "Queued call", transcript=False)
    reachable(monkeypatch, ["Nothing about that."])
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="Anything?")
    case_chat.answer_case_turn(turn.pk)
    turn.refresh_from_db()
    assert turn.answer.startswith("1 recording was not read: still transcribing.")
    assert turn.skipped == [
        {
            "recording": str(Recording.objects.get(title="Queued call").pk),
            "title": "Queued call",
            "why": "still transcribing",
        }
    ]


@pytest.mark.django_db
def test_the_hours_ceiling_refuses_with_its_figures(owner, a_case, monkeypatch):
    a_recording(owner, a_case, "Long one", hours=4)
    a_recording(owner, a_case, "Long two", hours=4)
    settings_store.set_to("case_chat_hours", 6)
    asked = reachable(monkeypatch, ["never"])
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="Anything?")
    case_chat.answer_case_turn(turn.pk)
    turn.refresh_from_db()
    assert turn.state == "failed" and turn.reason_class == "llm_case_too_large"
    assert case_chat.what_to_say(turn) == (
        "This case is too large for one question (8 hours of recordings; the "
        "limit is 6)."
    )
    assert asked == []
    row = Row.objects.get(category="llm", event="AI assistant call")
    assert row.reason_class == "llm_case_too_large"


@pytest.mark.django_db
def test_a_transcript_too_long_even_alone_names_itself(owner, a_case, monkeypatch):
    a_recording(owner, a_case, "Dense call")
    monkeypatch.setattr(engine, "window", lambda: 100)
    reachable(monkeypatch, ["never"])
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    turn = CaseChatTurn.objects.create(chat=chat, number=1, question="Anything?")
    case_chat.answer_case_turn(turn.pk)
    turn.refresh_from_db()
    assert turn.reason_class == "llm_too_long"
    assert "Dense call" in case_chat.what_to_say(turn)


@pytest.mark.django_db
def test_the_tab_and_its_endpoints_follow_the_switches(owner, a_case, client):
    a_recording(owner, a_case, "A call")
    signed_in(client, owner)
    page = client.get(f"/case/{a_case.pk}?tab=chat").content.decode()
    # Since v1.68.0 Gideon is the floating panel, not a tab (Phase 8 chapter 1).
    assert 'id="gideon-open"' in page and "?tab=chat" not in page
    assert page.index("chat-ui.js") < page.index("case-chat.js")
    # The viewer of a recording in the case offers the case's Chat tab.
    recording = Recording.objects.get(title="A call")
    viewer = client.get(f"/recording/{recording.pk}").content.decode()
    assert f'caseChatUrl: "/case/{a_case.pk}?tab=chat"' in viewer
    assert client.get(f"/case/{a_case.pk}/chat").status_code == 200
    assert client.get(f"/case/{a_case.pk}/chat").json()["starters"] == []
    settings_store.set_to("case_chat_starters", "Summarise this case")
    assert client.get(f"/case/{a_case.pk}/chat").json()["starters"] == [
        "Summarise this case"
    ]

    settings_store.set_to("chat_across_cases", False)
    page = client.get(f"/case/{a_case.pk}?tab=chat").content.decode()
    assert 'id="gideon-open"' not in page and "?tab=chat" not in page
    assert client.get(f"/case/{a_case.pk}/chat").status_code == 404
    assert client.post(f"/case/{a_case.pk}/chats").status_code == 404

    settings_store.set_to("chat_across_cases", True)
    settings_store.set_to("assistant_available", False)
    assert client.get(f"/case/{a_case.pk}/chat").status_code == 404
    # The viewer's own Chat toggle has no say over the Case Chat.
    settings_store.set_to("assistant_available", True)
    settings_store.set_to("chat_available", False)
    assert client.get(f"/case/{a_case.pk}/chat").status_code == 200


@pytest.mark.django_db
def test_a_question_needs_a_transcript_and_moves_the_clock(
    owner, a_case, client, monkeypatch
):
    swallow_defer(monkeypatch)
    signed_in(client, owner)
    chat_id = client.post(f"/case/{a_case.pk}/chats").json()["id"]
    refused = client.post(
        f"/case-chat/{chat_id}/ask",
        data=json.dumps({"question": "Anything?"}),
        content_type="application/json",
    )
    assert refused.status_code == 409
    a_recording(owner, a_case, "A call")
    before = Case.objects.get(pk=a_case.pk).last_activity
    asked = client.post(
        f"/case-chat/{chat_id}/ask",
        data=json.dumps({"question": "Anything?"}),
        content_type="application/json",
    )
    assert asked.status_code == 200
    assert Case.objects.get(pk=a_case.pk).last_activity >= before


@pytest.mark.django_db
def test_delete_writes_a_row_and_the_case_counts_its_chats(owner, a_case, client):
    signed_in(client, owner)
    chat = CaseChat.objects.create(case=a_case, asked_by=owner)
    assert client.get(f"/case/{a_case.pk}/what-would-go").json()["chats"] == 1
    assert client.post(f"/case-chat/{chat.pk}/delete").json()["ok"] is True
    assert not CaseChat.objects.exists()
    row = Row.objects.get(category="edits", event="Chat deleted")
    assert row.details["kind"] == "case chat" and row.object_label == "Ramirez"


@pytest.mark.django_db
def test_a_recording_that_left_leaves_its_citation_marked(owner, a_case, client):
    signed_in(client, owner)
    gone = a_recording(owner, a_case, "Moved out")
    chat = CaseChat.objects.create(case=a_case, asked_by=owner, name="Colour?")
    CaseChatTurn.objects.create(
        chat=chat,
        number=1,
        question="Colour?",
        answer="Blue [Recording 1, 00:12:45].",
        citations={
            "[Recording 1, 00:12:45]": {"recording": str(gone.pk), "seconds": 765.0}
        },
        readings=[case_chat.reading_of(gone, gone.transcript)],
        state="done",
    )
    gone.case = None
    gone.save(update_fields=["case"])
    state = client.get(f"/case/{a_case.pk}/chat").json()
    assert state["chats"][0]["turns"][0]["citations"] == {
        "[Recording 1, 00:12:45]": {"removed": True}
    }


@pytest.mark.django_db
def test_a_picture_moments_pill_shows_the_digests_line(
    owner, a_case, client, monkeypatch
):
    """v1.82.0: a citation to a second where nothing was said carries the
    Digest's line marked seen for that span on the pill's hover; a line
    marked said is not read as one."""
    from core import assistant
    from core.assistant import DigestPart

    signed_in(client, owner)
    recording = a_recording(owner, a_case, "Body camera")
    DigestPart.objects.create(
        transcript=recording.transcript,
        number=1,
        span_start=0.0,
        span_end=900.0,
        signature="x",
        text=(
            "1. [00:00:00]-[00:00:30] (said) A greeting.\n"
            "2. [00:06:30]-[00:06:40] (seen) A black handgun on the ground "
            "next to the front tire.\n"
            "3. [00:12:45] (both) The car was blue."
        ),
    )
    monkeypatch.setattr(assistant, "digests_on", lambda: True)
    here = str(recording.pk)
    chat = CaseChat.objects.create(case=a_case, asked_by=owner, name="Gun?")
    CaseChatTurn.objects.create(
        chat=chat,
        number=1,
        question="Gun?",
        answer=(
            "Seen [Recording 1, 00:06:35]; said [Recording 1, 00:12:45]; "
            "nothing [Recording 1, 00:00:10]."
        ),
        citations={
            "[Recording 1, 00:06:35]": {"recording": here, "seconds": 395.0},
            "[Recording 1, 00:12:45]": {"recording": here, "seconds": 765.0},
            "[Recording 1, 00:00:10]": {"recording": here, "seconds": 10.0},
        },
        readings=[case_chat.reading_of(recording, recording.transcript)],
        state="done",
    )
    cited = client.get(f"/case/{a_case.pk}/chat").json()["chats"][0]["turns"][0][
        "citations"
    ]
    seen = cited["[Recording 1, 00:06:35]"]
    assert seen["line"] == "Seen: A black handgun on the ground next to the front tire."
    assert seen["href"] == f"/recording/{recording.pk}?t=395.0"
    assert seen["clock"] == "00:06:35"
    assert cited["[Recording 1, 00:12:45]"]["line"] == "Speaker 2: The car was blue."
    assert cited["[Recording 1, 00:00:10]"]["line"] == ""


@pytest.mark.django_db
def test_process_again_is_said_and_the_export_lists_what_was_read(
    owner, a_case, client
):
    signed_in(client, owner)
    recording = a_recording(owner, a_case, "Interview", translate=True)
    chat = CaseChat.objects.create(case=a_case, asked_by=owner, name="Colour?")
    old = dict(case_chat.reading_of(recording, recording.transcript))
    old["transcript_created"] = "2000-01-01T00:00:00+00:00"
    CaseChatTurn.objects.create(
        chat=chat,
        number=1,
        question="Colour?",
        answer="Blue [Recording 1, 00:12:45].",
        citations={
            "[Recording 1, 00:12:45]": {
                "recording": str(recording.pk),
                "seconds": 765.0,
            }
        },
        readings=[old],
        state="done",
        model="the-model",
    )
    state = client.get(f"/case/{a_case.pk}/chat").json()
    assert state["chats"][0]["earlier"].startswith(
        "Earlier answers used a previous transcript of Interview"
    )

    from docx import Document

    exported = client.get(f"/case-chat/{chat.pk}/export")
    assert exported.status_code == 200
    assert exported["Content-Disposition"].startswith(
        'attachment; filename="Ramirez - case chat '
    )
    document = Document(BytesIO(exported.content))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    # The export is titled with the chat's name and the first question (v1.67.0).
    assert "Gideon: Colour?" in text and "Question 1." in text
    assert "Interview [00:12:45]" in text
    assert "Interview" in cells and "Jail call" in cells
    row = Row.objects.get(category="exports", event="export made")
    assert row.details["kind"] == "case chat"
    assert exports.case_chat_name(chat).startswith("Ramirez - case chat ")


def test_the_settings_and_the_template_exist():
    assert settings_store.DEFINITIONS["chat_across_cases"].default is True
    assert (
        settings_store.DEFINITIONS["chat_across_cases"].needs == "assistant_available"
    )
    hours = settings_store.DEFINITIONS["case_chat_hours"]
    assert (hours.default, hours.least, hours.most) == (120, 6, 600)
    starters = settings_store.DEFINITIONS["case_chat_starters"]
    assert starters.default == "" and starters.page == settings_store.TEMPLATES
    assert settings_store.TEMPLATES not in dict(settings_store.PAGES)
    assert PromptTemplate.DEFAULTS["case_chat"] == ("Case chat", prompts.CASE_CHAT)
    assert engine.WHAT_TO_SAY["llm_case_too_large"]


def test_the_answer_says_which_videos_were_read_from_the_transcript_alone(db):
    from core import case_chat, settings_store

    settings_store.set_to("assistant_available", True)
    settings_store.set_to("moments_available", True)
    settings_store.set_to("digests_available", True)
    readings = [
        {"video": True, "digest": True},
        {"video": True, "digest": False},
        {"video": False, "digest": False},
        {"video": True, "digest": False},
    ]
    assert (
        case_chat.transcript_alone_line(readings)
        == "Recording 2 and Recording 4 were read from the transcript alone."
    )
    assert (
        case_chat.transcript_alone_line(readings[:2])
        == "Recording 2 was read from the transcript alone."
    )
    assert case_chat.transcript_alone_line(readings[:1]) == ""
    # Digests off: every video reads from the transcript, so nothing is said.
    settings_store.set_to("digests_available", False)
    assert case_chat.transcript_alone_line(readings) == ""
